import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from docx import Document

print("=== KIỂM TRA FILE EXCEL ===")
wb = openpyxl.load_workbook(r'D:\UNIGO\Thời khóa biểu giáo viên\Thời khóa biểu - Đậu Đình Nguyên.xlsx', data_only=True)
ws = wb['TKB Đậu Đình Nguyên']
for r in range(5, 18):
    vals = [ws.cell(r, c).value for c in range(1, 9)]
    row_text = " | ".join(str(v).replace('\n', ' ') if v is not None else "---" for v in vals)
    print(f"Row {r:2d}: {row_text}")

print("\n=== KIỂM TRA FILE DOCX ===")
doc = Document(r'D:\UNIGO\Thời khóa biểu giáo viên\Thời khóa biểu - Đậu Đình Nguyên.docx')
tbl = doc.tables[0]
for ri, r in enumerate(tbl.rows):
    cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
    print(f"Row {ri:2d}: {' | '.join(cells)}")
