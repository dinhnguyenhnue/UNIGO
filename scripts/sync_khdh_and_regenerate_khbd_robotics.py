# -*- coding: utf-8 -*-
"""
sync_khdh_and_regenerate_khbd_robotics.py
=========================================
1. Đồng bộ Kế hoạch dạy học (KHDH) môn Robotics Lớp 1 -> 8 chuẩn 35 tiết/năm:
   - Đúng 4 mốc tuần kiểm tra đánh giá định kỳ: Tuần 9/10, Tuần 18/19, Tuần 27/28, Tuần 33/34, Tổng kết Tuần 35.
   - Cập nhật file KHDH tổng hợp, TH, THCS và từng lớp (Lớp 1-8).

2. Tạo lại toàn bộ KHBD Robotics Lớp 1 -> 8 chuẩn 100% theo AGENTS.md & LBG:
   - Lớp 1-4: Tuần 01 đến Tuần 35 (mỗi tuần 1 file, không có Tuần 36).
   - Lớp 5-8: Các tuần LẺ (Tuần 01, 03, 05, 07, 09, 11, 13, 15, 17, 19, 21, 23, 25, 27, 29, 31, 33, 35).
   - Đúng ngày soạn (Thứ 7 tuần trước), ngày dạy theo TKB.
   - Table 0: NO BORDER. Bảng tiến trình: CÓ VIỀN. Bảng chữ ký: NO BORDER.
   - Font Times New Roman 13pt toàn bộ.
"""

import os
import re
import sys
import shutil
from datetime import date, timedelta
import docx
from docx import Document
from docx.shared import Pt, Cm, Emu, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

# ─── PATH CONFIG ───────────────────────────────────────────────────────────
KHDH_MASTER = r"D:\UNIGO\Hệ thống mẫu văn bản\Nguyên đã làm\Kế hoạch dạy học môn Robotics 2026-2027.docx"
KHDH_TH     = r"D:\UNIGO\Hệ thống mẫu văn bản\Nguyên đã làm\Kế hoạch dạy học môn Robotics (TH) - 2026 - 2027.docx"
KHDH_THCS   = r"D:\UNIGO\Hệ thống mẫu văn bản\Nguyên đã làm\Kế hoạch dạy học môn Robotics (THCS) - 2026 - 2027.docx"
KHDH_DIR    = r"D:\UNIGO\Hệ thống mẫu văn bản\Nguyên đã làm\Kế hoạch dạy học Robotics từng lớp"
TPL_DOC     = r"D:\UNIGO\Hệ thống mẫu văn bản\PL4-Khung kế hoạch bài dạy (THCS).docx"
OUT_BASE    = r"D:\UNIGO\KHBD_Robotics"

TUAN_01_START = date(2026, 8, 3)

ROB_SCHEDULE = {
    '1': (3, '1A1'),  # Thứ Năm sáng (T2)
    '2': (2, '2A1'),  # Thứ Tư chiều (T4)
    '3': (2, '3A1'),  # Thứ Tư sáng (T2)
    '4': (0, '4C1'),  # Thứ Hai sáng (T3)
    '5': (1, '5C1'),  # Thứ Ba sáng (T3-T4)
    '6': (4, '6A1'),  # Thứ Sáu sáng (T4-T5)
    '7': (1, '7A1'),  # Thứ Ba chiều (T3-T4)
    '8': (4, '8A1'),  # Thứ Sáu chiều (T1-T2)
}

FONT_NAME = "Times New Roman"
FONT_SIZE_PT = 13

INDENT_0 = 0
INDENT_1 = 180340       # ~0.5cm
INDENT_2 = 360045       # ~1.0cm
INDENT_BULLET = 540000   # ~1.5cm
INDENT_TH_1 = 457200    # ~1.27cm
INDENT_TH_2 = 450215    # ~1.25cm


# ─── MASTER LESSON LISTS (35 TIẾT CHUẨN) ──────────────────────────────────
# Lớp 1: 35 tiết
LESSONS_L1 = [
    ("Tiết 0: Định hướng môn học", "Giới thiệu tổng quan môn học Robotics, an toàn sử dụng thiết bị và nội quy phòng học."),
    ("Bài 1. Tập thể dục nào!", "Nhận biết tác dụng của vận động; lắp robot thể dục đơn giản."),
    ("Bài 2. Chú cún dễ thương", "Nhận biết đặc điểm loài chó; lắp mô hình chú cún dễ thương."),
    ("Bài 3. Tăng cường sức khỏe", "Tìm hiểu ưu điểm của đi bộ; lắp robot bước đi tăng cường sức khỏe."),
    ("Bài 4. Chú ốc sên chậm chạp", "Tìm hiểu cách di chuyển và cấu tạo vỏ ốc; lắp robot ốc sên."),
    ("Bài 5. Xe cảnh sát tuần tra", "Tìm hiểu phương tiện cảnh sát; lắp xe cảnh sát tuần tra an ninh."),
    ("Bài 6. Khám phá xe cứu hoả", "Tìm hiểu bộ phận xe cứu hỏa và trang bị; lắp xe cứu hoả cơ bản."),
    ("Bài 7. Người giao hàng đã đến", "Tìm hiểu hình thức giao hàng, phân loại hàng; lắp robot xe giao hàng."),
    ("Ôn tập Đánh giá định kỳ 1", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 1."),
    ("Đánh giá định kỳ 1", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 1."),
    ("Bài 8. Thế giới khủng long", "Tìm hiểu nguồn gốc khủng long; lắp mô hình khủng long cử động."),
    ("Bài 9. Hồ bơi mùa hè", "Tìm hiểu an toàn bơi lội; lắp robot mô phỏng người bơi lội."),
    ("Bài 10. Khám phá đại dương", "Tìm hiểu đặc điểm cá mập; lắp mô hình cá mập săn mồi."),
    ("Bài 11. Chú cua cứng cáp", "Tìm hiểu nguyên lý cua bò ngang; lắp mô hình chú cua bò ngang."),
    ("Bài 12. Tôi có thể di chuyển đến bất cứ đâu", "Tìm hiểu các loại phương tiện; lắp robot di chuyển đa năng."),
    ("Bài 13. Hoạt động mùa hè", "Tìm hiểu hoạt động ngày hè; lắp mô hình người chèo thuyền kayak."),
    ("Bài 14. Bắt đầu chuyến hành trình cùng tàu hoả", "Tìm hiểu hệ thống tàu hỏa; lắp mô hình đoàn tàu chuyển động."),
    ("Ôn tập Đánh giá định kỳ 2", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 2."),
    ("Đánh giá định kỳ 2", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 2."),
    ("Bài 15. Phía trên bầu trời", "Tìm hiểu cấu tạo máy bay trực thăng; lắp robot trực thăng cứu hộ."),
    ("Bài 16. Khám phá vũ trụ rộng lớn", "Tìm hiểu phi hành gia và tàu vũ trụ; lắp mô hình tàu thám hiểm vũ trụ."),
    ("Bài 17. Máy bắn đá khổng lồ", "Tìm hiểu cơ cấu đòn bẩy; lắp mô hình máy bắn đá cổ đại."),
    ("Bài 18. Trò chơi dân gian", "Tìm hiểu trò chơi dân gian truyền thống; lắp mô hình trò chơi thú vị."),
    ("Bài 19. Khám phá trò chơi truyền thống các nước", "Tìm hiểu nét đẹp văn hóa truyền thống; lắp robot trò chơi dân gian."),
    ("Bài 20. Đấu vật thú vị", "Tìm hiểu môn đấu vật cổ truyền; lắp hai võ sĩ robot thi đấu."),
    ("Bài 21. Sóc nhỏ dễ thương", "Tìm hiểu tập tính sóc nhỏ nhặt hạt; lắp robot sóc chuyền cành."),
    ("Ôn tập Đánh giá định kỳ 3", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 3."),
    ("Đánh giá định kỳ 3", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 3."),
    ("Bài 22. Chú hươu tuyệt đẹp", "Tìm hiểu tập tính loài hươu cao cổ; lắp mô hình chú hươu tuyệt đẹp."),
    ("Bài 23. Chú rùa thông minh", "Tìm hiểu cấu tạo mai rùa; lắp robot chú rùa bò chậm rãi."),
    ("Bài 24. Đôi chân mạnh mẽ của chuột túi", "Tìm hiểu cơ chế chuột túi nhảy cao; lắp robot chuột túi bật nhảy."),
    ("Bài 25. Chú sư tử dũng mãnh", "Tìm hiểu tập tính sư tử chúa sơn lâm; lắp robot sư tử gầm vang."),
    ("Ôn tập Đánh giá định kỳ 4", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 4."),
    ("Đánh giá định kỳ 4", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 4."),
    ("Tổng kết môn học & Triển lãm Robotics", "Tổng kết đánh giá kết quả học tập và trưng bày sản phẩm Robotics."),
]

# Lớp 2, 3, 4: 35 tiết
LESSONS_L2_3_4 = [
    ("Tiết 0: Định hướng môn học", "Giới thiệu tổng quan môn học Robotics, an toàn sử dụng thiết bị và nội quy phòng học."),
    ("Bài 1. Hãy nấu những món ăn ngon", "Tìm hiểu văn hóa ẩm thực và nghề đầu bếp; lắp mô hình robot nhà bếp khuấy đảo."),
    ("Bài 2. Hàm răng trắng sáng", "Tìm hiểu cấu tạo hàm răng và vệ sinh răng miệng; lắp mô hình hàm răng cử động."),
    ("Bài 3. Sự phản xạ ánh sáng", "Tìm hiểu nguyên lý phản xạ ánh sáng; lắp mô hình kính tiềm vọng quang học."),
    ("Bài 4. Đàn gà con", "Tìm hiểu tập tính đàn gà; lắp mô hình gà mẹ dẫn gà con mổ thóc."),
    ("Bài 5. Những bạn nhỏ lễ phép", "Tìm hiểu hành vi giao tiếp lịch sự; lắp robot cúi chào lễ phép."),
    ("Bài 6. Động vật thân mềm", "Tìm hiểu cấu tạo động vật thân mềm; lắp mô hình bạch tuộc bơi lội."),
    ("Bài 7. Khu phố của chúng ta", "Tìm hiểu quy hoạch khu dân cư; lắp mô hình xe vệ sinh môi trường đường phố."),
    ("Ôn tập Đánh giá định kỳ 1", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 1."),
    ("Đánh giá định kỳ 1", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 1."),
    ("Bài 8. Rèn luyện sức khỏe", "Tìm hiểu tầm quan trọng của thể dục thể thao; lắp robot nâng tạ rèn luyện cơ bắp."),
    ("Bài 9. Môi trường biển", "Tìm hiểu bảo vệ hệ sinh thái biển; lắp mô hình tàu dọn rác đại dương."),
    ("Bài 10. Cùng câu cá nào!", "Tìm hiểu cơ cấu ròng rọc kéo dây; lắp mô hình cần câu cá tự động."),
    ("Bài 11. Thế giới khủng long", "Tìm hiểu các loài khủng long ăn cỏ; lắp mô hình khủng long cổ dài sải bước."),
    ("Bài 12. Người hiệp sĩ dũng cảm", "Tìm hiểu vũ khí và áo giáp hiệp sĩ thời xưa; lắp mô hình hiệp sĩ cưỡi ngựa."),
    ("Bài 13. Vận chuyển đồ vật", "Tìm hiểu nguyên lý bánh xe và trục quay; lắp mô hình xe kéo hàng giảm ma sát."),
    ("Bài 14. Sức mạnh của máy ủi", "Tìm hiểu nguyên lý máy ủi san lấp mặt bằng; lắp mô hình máy ủi công trình."),
    ("Ôn tập Đánh giá định kỳ 2", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 2."),
    ("Đánh giá định kỳ 2", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 2."),
    ("Bài 15. Máy xúc", "Tìm hiểu cơ cấu tay đòn thủy lực máy xúc; lắp mô hình máy xúc đào đất."),
    ("Bài 16. Cùng nhau đi khắp thế giới", "Tìm hiểu các phương tiện vượt đại dương; lắp mô hình thuyền buồm hai thân."),
    ("Bài 17. Nhắm và bắn!", "Tìm hiểu nguyên lý bắn đạn cơ học; lắp mô hình máy bắn mục tiêu chuẩn xác."),
    ("Bài 18. Độ đàn hồi, lực đẩy của cung tên", "Tìm hiểu thế năng đàn hồi; lắp mô hình cung nỏ cơ khí bắn tên."),
    ("Bài 19. Ba! Hai! Một! Bắn!!!", "Tìm hiểu cơ chế bệ phóng tên lửa; lắp mô hình bệ phóng tên lửa tự động."),
    ("Bài 20. Cùng nhau tham quan thành phố", "Tìm hiểu phương tiện giao thông công cộng; lắp mô hình xe buýt 2 tầng chở khách."),
    ("Bài 21. Khám phá trang phục người da đỏ", "Tìm hiểu văn hóa thổ dân; lắp mô hình vũ công múa lễ hội truyền thống."),
    ("Ôn tập Đánh giá định kỳ 3", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 3."),
    ("Đánh giá định kỳ 3", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 3."),
    ("Bài 22. Cá sấu thật ngầu!", "Tìm hiểu cơ chế đớp mồi của cá sấu; lắp mô hình hàm cá sấu đớp mồi chớp nhoáng."),
    ("Bài 23. Những nốt nhạc vui", "Tìm hiểu nguyên lý gõ tạo âm thanh; lắp mô hình nhạc công gõ đàn phiêu lãng."),
    ("Bài 24. Choo Choo! Tàu hỏa", "Tìm hiểu cơ cấu thanh truyền đầu máy hơi nước; lắp đoàn tàu hỏa bánh răng."),
    ("Bài 25. Có sáu chân thật tuyệt!!!", "Tìm hiểu cơ chế bước đi 6 chân của côn trùng; lắp robot bọ 6 chân vượt địa hình."),
    ("Ôn tập Đánh giá định kỳ 4", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 4."),
    ("Đánh giá định kỳ 4", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 4."),
    ("Tổng kết môn học & Triển lãm Robotics", "Tổng kết đánh giá kết quả học tập và trưng bày sản phẩm Robotics."),
]

# Lớp 5, 6, 7, 8: 35 tiết (Dạy 18 tuần lẻ: Tuần 1: 1 tiết, Tuần 3-35: 17 tuần x 2 tiết = 34 tiết)
LESSONS_L5_6_7_8 = [
    # Tuần 1 (1 tiết)
    ("Tiết 0: Định hướng môn học", "Giới thiệu tổng quan môn học Robotics, phương pháp học tập thực hành, an toàn sử dụng thiết bị và nội quy phòng học bộ môn."),
    # Tuần 3 (2 tiết)
    ("Bài 1. Động cơ là gì?", "Tìm hiểu động cơ Dynamixel, chốt đặc biệt; lắp và vận hành robot xe nâng."),
    ("Bài 2. Robot cố định vật", "Tìm hiểu dụng cụ cố định đồ vật; thiết kế robot xe gắp và di chuyển vật thể."),
    # Tuần 5 (2 tiết)
    ("Bài 3. Robot nhận biết âm thanh", "Tìm hiểu cảm biến âm thanh; lắp robot phản ứng với tiếng vỗ tay."),
    ("Bài 4. Bộ điều khiển từ xa", "Tìm hiểu chức năng bộ điều khiển; điều khiển robot pháo bắn mục tiêu."),
    # Tuần 7 (2 tiết)
    ("Bài 5. Động cơ Dynamixel hoạt động thế nào?", "Tìm hiểu cách hoạt động của động cơ servo Dynamixel; lắp robot khuếch đại chuyển động."),
    ("Bài 6. Phương tiện giao thông qua các thời đại", "Tìm hiểu cơ cấu chuyển động lặp; chế tạo robot đua ngựa thi đấu."),
    # Tuần 9 (2 tiết) - ĐÁNH GIÁ ĐỊNH KỲ 1
    ("Ôn tập Đánh giá định kỳ 1", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 1."),
    ("Đánh giá định kỳ 1", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 1."),
    # Tuần 11 (2 tiết)
    ("Bài 7. Cơ chế tăng giảm chiều dài tự động", "Tìm hiểu cơ cấu kẹp linh hoạt và cánh tay robot; lắp robot gỡ bom/gắp vật xa."),
    ("Bài 8. Robot nhận biết vật thể bằng cách nào?", "Tìm hiểu cảm biến hồng ngoại phát hiện vật cản; lắp robot tự động tránh chướng ngại vật."),
    # Tuần 13 (2 tiết)
    ("Bài 9. Số ngẫu nhiên", "Ứng dụng thuật toán số ngẫu nhiên trong điều khiển robot hoạt động tự động."),
    ("Bài 10. Robot hút bụi", "Tìm hiểu chức năng robot dọn dẹp; lắp và lập trình robot hút bụi tự động."),
    # Tuần 15 (2 tiết)
    ("Bài 11. Lực hấp dẫn", "Tìm hiểu trọng lực và ứng dụng; chế tạo robot Rodeo sử dụng lực hấp dẫn."),
    ("Bài 12. Domino trong Robotics", "Tìm hiểu chuyển động chính xác; lắp robot xếp khối domino tự động."),
    # Tuần 17 (2 tiết)
    ("Bài 13. Robot phục vụ đời sống", "Tìm hiểu ứng dụng robot trong công nghiệp và đời sống."),
    ("Luyện tập & Thực hành sáng tạo Robotics", "Luyện tập nâng cao, mở rộng tính năng và thử nghiệm sản phẩm."),
    # Tuần 19 (2 tiết) - ĐÁNH GIÁ ĐỊNH KỲ 2
    ("Ôn tập Đánh giá định kỳ 2", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 2."),
    ("Đánh giá định kỳ 2", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 2."),
    # Tuần 21 (2 tiết)
    ("Bài 14. Robot thám hiểm", "Tìm hiểu chức năng robot thám hiểm địa hình hiểm trở; lắp và vận hành robot bánh xích vượt chướng ngại vật."),
    ("Bài 15. Cảm biến siêu âm", "Tìm hiểu nguyên lý phát và thu sóng siêu âm đo khoảng cách; lắp robot đo cự ly tự động."),
    # Tuần 23 (2 tiết)
    ("Bài 16. Tránh vật cản tự động", "Ứng dụng cảm biến siêu âm lập trình thuật toán tự động chuyển hướng khi gặp chướng ngại vật."),
    ("Bài 17. Robot theo dõi đường line", "Tìm hiểu cảm biến dò đường (line sensor); lắp và lập trình robot bám vạch kẻ đường chính xác."),
    # Tuần 25 (2 tiết)
    ("Bài 18. Lập trình vòng lặp trong điều khiển robot", "Ứng dụng cấu trúc lặp (vòng lặp vô hạn, lặp có điều kiện) trong phần mềm điều khiển robot."),
    ("Bài 19. Khám phá cánh tay robot", "Tìm hiểu cấu tạo các khớp chuyển động; lắp ráp mô hình cánh tay robot 3 bậc tự do."),
    # Tuần 27 (2 tiết) - ĐÁNH GIÁ ĐỊNH KỲ 3
    ("Ôn tập Đánh giá định kỳ 3", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 3."),
    ("Đánh giá định kỳ 3", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 3."),
    # Tuần 29 (2 tiết)
    ("Bài 20. Robot phân loại sản phẩm", "Tìm hiểu hệ thống tự động hóa trong nhà máy; lắp robot phân loại đồ vật theo kích thước/màu sắc."),
    ("Bài 21. Cơ chế kẹp gắp vật thể", "Tìm hiểu động học ngón kẹp cơ khí; tối ưu hóa lực kẹp và độ chính xác khi di chuyển đồ vật."),
    # Tuần 31 (2 tiết)
    ("Bài 22. Dự án sáng tạo robot tự hành", "Vận dụng tổng hợp kiến thức cơ khí và cảm biến thiết kế mô hình robot tự hành đa năng."),
    ("Bài 23. Tối ưu hóa thuật toán điều khiển robot", "Tinh chỉnh thông số điều khiển, sửa lỗi lập trình và thử nghiệm robot hoàn thành nhiệm vụ thực tế."),
    # Tuần 33 (2 tiết) - ĐÁNH GIÁ ĐỊNH KỲ 4
    ("Ôn tập Đánh giá định kỳ 4", "Hệ thống hóa kiến thức chuẩn bị Đánh giá định kỳ 4."),
    ("Đánh giá định kỳ 4", "Kiểm tra đánh giá kết quả học tập thực hành Robotics định kỳ 4."),
    # Tuần 35 (2 tiết) - TỔNG KẾT NĂM HỌC
    ("Tổng kết môn học & Triển lãm Robotics (Tiết 1)", "Tổng kết đánh giá kết quả học tập và trưng bày sản phẩm Robotics sáng tạo."),
    ("Tổng kết môn học & Triển lãm Robotics (Tiết 2)", "Báo cáo dự án, trao chứng nhận và định hướng phát triển công nghệ trong năm học tiếp theo."),
]


def get_lessons_for_grade(grade):
    if grade == 1:
        return LESSONS_L1
    elif grade in [2, 3, 4]:
        return LESSONS_L2_3_4
    else:
        return LESSONS_L5_6_7_8


# ─── FORMATTING HELPERS ────────────────────────────────────────────────────
def afont(run, bold=False, italic=False, size_pt=FONT_SIZE_PT):
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rFonts.set(qn('w:cs'), FONT_NAME)
    rFonts.set(qn('w:eastAsia'), FONT_NAME)


def clean_body_preserve_sectpr(doc):
    for child in list(doc.element.body):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag != 'sectPr':
            doc.element.body.remove(child)


def add_p(doc, text="", bold=False, italic=False, first_indent=None,
          left_indent=None, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
          space_after_pt=3, space_before_pt=0, line_spacing=1.15, size_pt=FONT_SIZE_PT):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after_pt)
    pf.space_before = Pt(space_before_pt)
    if first_indent is not None:
        pf.first_line_indent = Emu(first_indent)
    if left_indent is not None:
        pf.left_indent = Emu(left_indent)
    if text:
        run = p.add_run(text)
        afont(run, bold=bold, italic=italic, size_pt=size_pt)
    return p


def add_multi_run(doc, runs_data, first_indent=None, left_indent=None,
                  alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(3)
    pf.space_before = Pt(0)
    if first_indent is not None:
        pf.first_line_indent = Emu(first_indent)
    if left_indent is not None:
        pf.left_indent = Emu(left_indent)
    for text, bold, italic in runs_data:
        run = p.add_run(text)
        afont(run, bold=bold, italic=italic)
    return p


def set_borders(table, val="single", sz="4", color="000000"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders_el = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), val)
        b.set(qn('w:sz'), sz)
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)
        borders_el.append(b)
    tblPr.append(borders_el)


def set_no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders_el = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'nil')
        borders_el.append(b)
    tblPr.append(borders_el)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, v in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def fill_cell(cell, text, bold=False, italic=False,
              align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=FONT_SIZE_PT,
              space_after=2, line_spacing=1.15):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    set_cell_margins(cell)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.line_spacing = line_spacing
            p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(line)
        afont(run, bold=bold, italic=italic, size_pt=size_pt)


def sanitize(name):
    import unicodedata
    name = unicodedata.normalize('NFD', name).encode('ascii', 'ignore').decode('utf-8')
    name = re.sub(r'[^a-zA-Z0-9_\-]', '_', name)
    name = re.sub(r'_+', '_', name).strip('_')
    return name


def compute_dates(tuan_so, day_of_week):
    week_start = TUAN_01_START + timedelta(weeks=tuan_so - 1)
    ngay_day = week_start + timedelta(days=day_of_week)
    ngay_soan = week_start - timedelta(days=2)  # Thứ 7 tuần trước
    return ngay_soan.strftime('%d/%m/%Y'), ngay_day.strftime('%d/%m/%Y')


def get_bac_nls(grade):
    if grade <= 3:
        return 1
    elif grade <= 5:
        return 2
    elif grade <= 7:
        return 3
    else:
        return 4


def get_kit_name(grade):
    if grade in [1, 2]:
        return "OLLO Kinder"
    elif grade in [3, 4]:
        return "OLLO Initiate"
    else:
        return "OLLO Excel 1"


def get_rob_competencies(grade, title, yccd="", kit_name=""):
    t_lower = (title + " " + yccd).lower()
    
    if any(k in t_lower for k in ['định hướng', 'tiết 0']):
        c1 = f"- NL1 (Nhận thức công nghệ): Nhận biết tổng quan chương trình môn học, an toàn sử dụng bộ kit {kit_name} và nội quy phòng học Robotics. (Đạt được thông qua Hoạt động 1, Hoạt động 2)"
        c2 = f"- NL2 (Sử dụng công nghệ): Tuân thủ quy tắc làm việc nhóm, bảo quản linh kiện và sử dụng thiết bị đúng hướng dẫn. (Đạt được thông qua Hoạt động 3, Hoạt động 4)"
        return [c1, c2]

    if any(k in t_lower for k in ['triển lãm', 'tổng kết', 'trình bày', 'chia sẻ', 'giới thiệu', 'giao lưu']):
        c1 = f"- NL5 (Giao tiếp công nghệ): Tự tin thuyết minh nguyên lý hoạt động và giới thiệu sản phẩm robot {title} trước tập thể. (Đạt được thông qua Hoạt động 3, Hoạt động 4)"
        c2 = f"- NL4 (Đánh giá công nghệ): Lắng nghe, nhận xét và đánh giá sản phẩm robot của nhóm bạn theo tiêu chí kỹ thuật. (Đạt được thông qua Hoạt động 4)"
        return [c1, c2]
        
    elif any(k in t_lower for k in ['thử nghiệm', 'đo', 'so sánh', 'lực', 'quán tính', 'tốc độ', 'góc', 'vận tốc', 'gia tốc', 'lực hấp dẫn', 'đàn hồi', 'ném', 'chuyển động']):
        c1 = f"- NL4 (Đánh giá công nghệ): Quan sát, đo đạc thông số và đánh giá hiệu quả vận hành của mô hình robot {title} trong các điều kiện thử nghiệm khác nhau. (Đạt được thông qua Hoạt động 3, Hoạt động 4)"
        c2 = f"- NL3 (Thiết kế kĩ thuật): Điều chỉnh kết cấu cơ khí và thông số điều khiển để tối ưu hóa khả năng hoạt động của robot. (Đạt được thông qua Hoạt động 3)"
        return [c1, c2]

    elif any(k in t_lower for k in ['sáng tạo', 'cải tiến', 'thiết kế', 'cánh tay', 'kẹp', 'gắp', 'tự động', 'domino', 'nhiệm vụ', 'hút bụi']):
        c1 = f"- NL3 (Thiết kế kĩ thuật): Lắp ráp, phối hợp các cơ cấu truyền động và sáng tạo cải tiến mô hình robot {title} thực hiện nhiệm vụ đặt ra. (Đạt được thông qua Hoạt động 3, Hoạt động 4)"
        c2 = f"- NL2 (Sử dụng công nghệ): Sử dụng thành thạo các module, khớp nối và linh kiện bộ kit {kit_name} theo đúng tiêu chuẩn an toàn. (Đạt được thông qua Hoạt động 2, Hoạt động 3)"
        return [c1, c2]

    elif any(k in t_lower for k in ['động cơ', 'cảm biến', 'âm thanh', 'hồng ngoại', 'dynamixel', 'bộ điều khiển', 'linh kiện', 'chốt', 'khung']):
        c1 = f"- NL1 (Nhận thức công nghệ): Nhận biết chính xác tên gọi, cấu tạo và vai trò của động cơ/cảm biến/bộ điều khiển trong mô hình {title}. (Đạt được thông qua Hoạt động 2)"
        c2 = f"- NL2 (Sử dụng công nghệ): Thao tác đấu nối đúng kỹ thuật, kiểm tra tín hiệu đầu vào/ra và vận hành robot an toàn. (Đạt được thông qua Hoạt động 3)"
        return [c1, c2]

    elif any(k in t_lower for k in ['ôn tập', 'đánh giá định kỳ']):
        c1 = f"- NL1 (Nhận thức công nghệ): Hệ thống hóa kiến thức về cấu tạo, nguyên lý hoạt động của các mô hình robot đã học. (Đạt được thông qua Hoạt động 2)"
        c2 = f"- NL2 (Sử dụng công nghệ): Thực hành thành thạo các thao tác lắp ráp và vận hành robot theo yêu cầu kiểm tra. (Đạt được thông qua Hoạt động 3)"
        return [c1, c2]

    else:
        c1 = f"- NL1 (Nhận thức công nghệ): Nhận biết mô hình {title} mô phỏng sự vật/hiện tượng thực tế và hiểu nguyên lý hoạt động cơ bản. (Đạt được thông qua Hoạt động 2)"
        c2 = f"- NL2 (Sử dụng công nghệ): Thực hiện lắp ráp đúng quy trình từng bước mô hình {title} từ bộ kit {kit_name}, vận hành chạy thử và kiểm tra hoạt động. (Đạt được thông qua Hoạt động 3)"
        return [c1, c2]


def yccd_to_noun(item):
    item = item.strip().strip('-').strip()
    if not item:
        return ''
    verb_patterns = [
        r'^Nhận biết được\s+', r'^Phân biệt được\s+', r'^Nêu được\s+',
        r'^Giải thích được\s+', r'^Biết\s+', r'^Hiểu được\s+',
        r'^Trình bày được\s+', r'^Mô tả được\s+', r'^Vận dụng được\s+',
        r'^Thực hiện được\s+', r'^Sử dụng được\s+', r'^Xác định được\s+',
        r'^Liệt kê được\s+', r'^So sánh được\s+', r'^Phân tích được\s+',
        r'^Đánh giá được\s+', r'^Tạo được\s+', r'^Viết được\s+',
        r'^Lập được\s+', r'^Thiết kế được\s+', r'^Lắp ráp được\s+',
        r'^Kể tên được\s+', r'^Nhận diện được\s+', r'^Tìm hiểu\s+',
        r'^Chế tạo\s+', r'^Lắp\s+',
    ]
    for pat in verb_patterns:
        item = re.sub(pat, '', item, count=1)
    banned = ['Sự hiểu biết về ', 'Khả năng nhận diện ', 'Khả năng phân tích ',
              'Khả năng vận dụng ', 'Sự nhận biết ']
    for b in banned:
        item = item.replace(b, '')
    if item:
        item = item[0].upper() + item[1:]
    item = item.rstrip('.')
    return item


def parse_yccd_bullets(yccd_raw):
    if not yccd_raw or len(yccd_raw.strip()) < 3:
        return []
    items = re.split(r'\s*[;–\-]\s+|\.\s+', yccd_raw.strip())
    result = []
    for item in items:
        item = item.strip()
        if not item:
            continue
        noun = yccd_to_noun(item)
        if noun and len(noun) > 3:
            result.append(noun)
    return result


def parse_yccd_for_th(yccd_raw):
    if not yccd_raw or len(yccd_raw.strip()) < 3:
        return []
    text = yccd_raw.strip()
    if ';' in text:
        parts = text.split(';')
    elif ' - ' in text:
        parts = re.split(r'\s*-\s+', text)
    else:
        parts = [text]

    cleaned = []
    for it in parts:
        it = it.lstrip('-').strip()
        if it:
            it = it[0].upper() + it[1:]
            if not it.endswith('.'):
                it += '.'
            cleaned.append(it)
    return cleaned


def format_th_title(raw_title, tiet_ppct):
    t = raw_title.strip()
    m = re.match(r'^(?:Bài|BÀI)\s*(\d+)[\.:\s]*(.*)$', t)
    if m:
        bai_num, bai_name = m.group(1), m.group(2).strip()
        return f"BÀI {bai_num}. {bai_name.upper()} (Tiết: {tiet_ppct} theo PPCT)"
    elif 'tiết 0' in t.lower() or 'định hướng' in t.lower():
        return f"TIẾT 0: ĐỊNH HƯỚNG MÔN HỌC (Tiết: {tiet_ppct} theo PPCT)"
    elif 'ôn tập' in t.lower():
        return f"{t.upper()} (Tiết: {tiet_ppct} theo PPCT)"
    elif 'đánh giá' in t.lower():
        return f"{t.upper()} (Tiết: {tiet_ppct} theo PPCT)"
    elif 'tổng kết' in t.lower():
        return f"{t.upper()} (Tiết: {tiet_ppct} theo PPCT)"
    else:
        return f"{t.upper()} (Tiết: {tiet_ppct} theo PPCT)"


# ─── BUILD DOCUMENT (TIỂU HỌC & THCS) ─────────────────────────────────────
def build_khbd_th(grade, ten_lop, title, tiet_ppct, yccd_raw, tuan_so, day_of_week, kit_name):
    doc = Document(TPL_DOC)
    clean_body_preserve_sectpr(doc)
    ngay_soan, ngay_day = compute_dates(tuan_so, day_of_week)
    bac_nls = get_bac_nls(grade)
    ls = 1.15

    # Table 0: Info table 3x2 NO BORDER
    tbl_info = doc.add_table(rows=3, cols=2)
    set_no_borders(tbl_info)
    tbl_info.rows[0].cells[0].width = Cm(8.5)
    tbl_info.rows[0].cells[1].width = Cm(7.5)

    fill_cell(tbl_info.rows[0].cells[0], 'Trường: Tiểu học và THCS UNIGO', bold=True)
    fill_cell(tbl_info.rows[0].cells[1], f'Ngày soạn: {ngay_soan}', bold=True)
    fill_cell(tbl_info.rows[1].cells[0], 'GV: Đậu Đình Nguyên')
    fill_cell(tbl_info.rows[1].cells[1], f'Ngày dạy: {ngay_day}')
    fill_cell(tbl_info.rows[2].cells[0], 'Tổ: Tổ chuyên môn Tiểu học')
    fill_cell(tbl_info.rows[2].cells[1], f'Lớp: {ten_lop}')

    add_p(doc, '', line_spacing=ls)

    # Title block for TH (3 lines)
    add_p(doc, 'KẾ HOẠCH DẠY HỌC MÔN ROBOTICS', bold=True, size_pt=14,
          alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=ls)
    add_p(doc, f'CHỦ ĐIỂM: BỘ THIẾT BỊ: {kit_name.upper()}', bold=True, size_pt=13,
          alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=ls)
    th_title = format_th_title(title, tiet_ppct)
    add_p(doc, th_title, bold=True, size_pt=13,
          alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=ls)

    add_p(doc, '', line_spacing=ls)

    # I. YÊU CẦU CẦN ĐẠT
    add_p(doc, 'I. YÊU CẦU CẦN ĐẠT:', bold=True, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '- Sau bài học này em sẽ:', bold=True, first_indent=INDENT_TH_1, line_spacing=ls)
    yccd_items = parse_yccd_for_th(yccd_raw)
    if not yccd_items:
        yccd_items = [f"Hiểu và thực hành tốt nội dung {title}."]
    for item in yccd_items:
        add_p(doc, f'+ {item}', first_indent=INDENT_TH_2, line_spacing=ls)

    # 1. Năng lực
    add_p(doc, '1. Năng lực:', bold=True, first_indent=INDENT_TH_1, line_spacing=ls)
    add_p(doc, '1.1. Năng lực đặc thù (Robotics):', bold=True,
          first_indent=INDENT_TH_2, line_spacing=ls)
    for c in get_rob_competencies(grade, title, yccd_raw, kit_name):
        add_p(doc, c, left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    add_p(doc, '1.2. Năng lực số (Thông tư 02/2025 – CV 3456):', bold=True,
          first_indent=INDENT_TH_2, line_spacing=ls)
    add_p(doc, f'- Miền III. Sáng tạo nội dung số – Thành tố 3.2. Sáng tạo sản phẩm số – Bậc {bac_nls}: Lắp ráp và lập trình mô hình robot hoàn chỉnh theo mục tiêu bài học. (Đạt được thông qua Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, f'- Miền IV. An toàn – Thành tố 4.3. Bảo vệ sức khỏe và an toàn – Bậc {bac_nls}: Tuân thủ quy tắc an toàn khi sử dụng linh kiện điện tử và bộ pin nguồn. (Đạt được thông qua Hoạt động 1, Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    add_p(doc, '1.3. Năng lực chung:', bold=True, first_indent=INDENT_TH_2, line_spacing=ls)
    add_p(doc, '- Tự chủ và tự học: Tự giác hoàn thành phần việc được phân công trong nhóm thực hành chế tạo robot. (Đạt được thông qua Hoạt động 2, Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '- Giao tiếp và hợp tác: Tích cực trao đổi, phối hợp nhịp nhàng với bạn học khi lắp ghép các bộ phận robot. (Đạt được thông qua Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    # 2. Phẩm chất
    add_p(doc, '2. Phẩm chất:', bold=True, first_indent=INDENT_TH_1, line_spacing=ls)
    add_p(doc, '- Chăm chỉ: Kiên trì, tỉ mỉ trong từng bước lắp ghép và sửa chữa mô hình robot. (Đạt được thông qua Hoạt động 2, Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '- Trách nhiệm: Giữ gìn vệ sinh chung, bảo quản và sắp xếp gọn gàng bộ linh kiện Robotics sau buổi học. (Đạt được thông qua Hoạt động 3, Hoạt động 4)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    # II. ĐỒ DÙNG DẠY HỌC
    add_p(doc, 'II. ĐỒ DÙNG DẠY HỌC:', bold=True, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '1. Giáo viên:', bold=True, first_indent=INDENT_TH_1, line_spacing=ls)
    add_p(doc, f'- Bộ Kit Robotics {kit_name}, máy tính giáo viên, máy chiếu bài giảng trực quan.',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, f'- Phiếu học tập thực hành và mô hình robot mẫu cho bài {title}.',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '2. Học sinh:', bold=True, first_indent=INDENT_TH_1, line_spacing=ls)
    add_p(doc, f'- Bộ Kit Robotics {kit_name} theo nhóm, tài liệu hướng dẫn học tập.',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    # III. CÁC HOẠT ĐỘNG DẠY HỌC CHỦ YẾU
    add_p(doc, 'III. CÁC HOẠT ĐỘNG DẠY HỌC CHỦ YẾU:', bold=True,
          first_indent=INDENT_0, line_spacing=ls)

    activities_th = [
        ('1. Hoạt động 1. Khởi động (5-7 phút)',
         f'Tạo hứng thú và kết nối vào bài học {title}.',
         'Quan sát hình ảnh/video thực tế và thảo luận.',
         'Câu trả lời của HS về chủ đề bài học.',
         '- GV trình chiếu clip/hình ảnh thực tế sinh động liên quan đến chủ đề.\n- Đặt câu hỏi dẫn dắt: "Em có biết bạn robot này giúp ích gì cho con người không?"\n- HS chú ý quan sát, thảo luận cặp đôi và xung phong trả lời.\n- GV nhận xét, tuyên dương và dẫn dắt vào bài mới.',
         'HS hào hứng, tập trung vào bài học.\nNhận diện được chủ đề bài học.'),

        ('2. Hoạt động 2. Khám phá (12-15 phút)',
         f'Tìm hiểu cấu tạo và danh mục linh kiện cần dùng cho {title}.',
         'Tìm hiểu các chi tiết khung, chốt nối, động cơ và bánh xe.',
         'Khay linh kiện đầy đủ, sắp xếp ngăn nắp theo hướng dẫn.',
         '- GV phát phiếu hướng dẫn và giới thiệu từng loại linh kiện chính.\n- Hướng dẫn HS cách chọn đúng chốt nối và cách phân biệt các khớp xoay.\n- HS làm việc nhóm, chọn đúng các chi tiết cần thiết ra khay chứa.\n- GV đi từng bàn kiểm tra, hỗ trợ nhóm còn lúng túng.',
         'HS gọi đúng tên các linh kiện cơ bản.\nChuẩn bị đầy đủ linh kiện lắp ráp.'),

        ('3. Hoạt động 3. Luyện tập – Thực hành (12-15 phút)',
         f'Lắp ráp và vận hành mô hình {title} theo đúng quy trình.',
         'Tiến hành lắp ghép mô hình robot theo các bước hướng dẫn.',
         f'Mô hình robot {title} hoàn thiện, hoạt động đúng yêu cầu.',
         '- GV nêu yêu cầu thực hành và lưu ý an toàn khi lắp ráp.\n- HS phân công nhiệm vụ: bạn xem sơ đồ, bạn lắp khung, bạn thử pin.\n- Các nhóm hoàn thiện mô hình robot và bật công tắc chạy thử.\n- GV quan sát, khích lệ và chấm điểm thi đua giữa các nhóm.',
         'Mô hình robot hoàn thành chắc chắn.\nRobot hoạt động trơn tru, chính xác.'),

        ('4. Hoạt động 4. Vận dụng – Sáng tạo (5-8 phút)',
         'Rút ra bài học thực tiễn, sáng tạo mở rộng và thu dọn đồ dùng.',
         'Chia sẻ cảm nhận, nêu ý tưởng cải tiến và dọn dẹp vệ sinh.',
         'Ý tưởng sáng tạo mới và bộ Kit được thu dọn gọn gàng.',
         '- GV đặt câu hỏi: "Em có thể gắn thêm chi tiết gì để bạn robot đẹp hơn?"\n- Đại diện 1-2 nhóm giơ robot lên giới thiệu và biểu diễn trước lớp.\n- GV tổng kết tiết học, dặn dò HS cẩn thận tháo dỡ hoặc bảo quản mô hình.\n- HS thu gom linh kiện vào đúng ngăn trong hộp Kit, xếp ghế ngay ngắn.',
         'HS tự tin chia sẻ ý tưởng trước lớp.\nHộp đồ dùng ngăn nắp, phòng học sạch sẽ.')
    ]

    for idx, (act_title, muc_tieu, noi_dung, san_pham, gv_hs_col, ket_qua_col) in enumerate(activities_th, 1):
        add_p(doc, act_title, bold=True, first_indent=INDENT_TH_1, line_spacing=ls)
        add_multi_run(doc, [
            ('a) Mục tiêu: ', True, False), (muc_tieu, False, True)
        ], first_indent=INDENT_TH_1, line_spacing=ls)
        add_multi_run(doc, [
            ('b) Nội dung: ', True, False), (noi_dung, False, True)
        ], first_indent=INDENT_TH_1, line_spacing=ls)
        add_multi_run(doc, [
            ('c) Sản phẩm: ', True, False), (san_pham, False, True)
        ], first_indent=INDENT_TH_1, line_spacing=ls)
        add_multi_run(doc, [
            ('d) Tổ chức thực hiện:', True, False),
        ], first_indent=INDENT_TH_1, line_spacing=ls)

        tbl = doc.add_table(rows=1, cols=2)
        set_borders(tbl)
        for cell in tbl.columns[0].cells:
            cell.width = Cm(9.0)
        for cell in tbl.columns[1].cells:
            cell.width = Cm(7.0)

        fill_cell(tbl.rows[0].cells[0], 'HOẠT ĐỘNG CỦA GV – HS', bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(tbl.rows[0].cells[1], 'KẾT QUẢ CẦN ĐẠT', bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

        row = tbl.add_row()
        fill_cell(row.cells[0], gv_hs_col)
        fill_cell(row.cells[1], ket_qua_col)

        add_p(doc, '', line_spacing=ls)

    # Rút kinh nghiệm & Bảng chữ ký
    add_p(doc, 'RÚT KINH NGHIỆM SAU BÀI DẠY:', bold=True,
          first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '...........................................................................................................................', line_spacing=ls)
    add_p(doc, '...........................................................................................................................', line_spacing=ls)
    add_p(doc, '', line_spacing=ls)

    tbl_sign = doc.add_table(rows=3, cols=3)
    set_no_borders(tbl_sign)
    for i, txt in enumerate(['DUYỆT CỦA BGH', 'DUYỆT CỦA TỔ CM', 'NGƯỜI SOẠN']):
        fill_cell(tbl_sign.rows[0].cells[i], txt, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    for i in range(3):
        fill_cell(tbl_sign.rows[1].cells[i], '(Ký, ghi rõ họ tên)',
                  italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i in range(3):
        fill_cell(tbl_sign.rows[2].cells[i], '\n\n\n',
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(tbl_sign.rows[2].cells[2], '\n\n\nĐậu Đình Nguyên', bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER)

    return doc


def build_khbd_thcs(grade, ten_lop, title, tiet_ppct, yccd_raw, tuan_so, day_of_week, kit_name):
    doc = Document(TPL_DOC)
    clean_body_preserve_sectpr(doc)
    ngay_soan, ngay_day = compute_dates(tuan_so, day_of_week)
    bac_nls = get_bac_nls(grade)
    ls = 1.15

    # Table 0: Info table 3x2 NO BORDER
    tbl_info = doc.add_table(rows=3, cols=2)
    set_no_borders(tbl_info)
    tbl_info.rows[0].cells[0].width = Cm(8.5)
    tbl_info.rows[0].cells[1].width = Cm(7.5)

    fill_cell(tbl_info.rows[0].cells[0], 'Trường: Tiểu học và THCS UNIGO', bold=True)
    fill_cell(tbl_info.rows[0].cells[1], f'Ngày soạn: {ngay_soan}', bold=True)
    fill_cell(tbl_info.rows[1].cells[0], 'GV: Đậu Đình Nguyên')
    fill_cell(tbl_info.rows[1].cells[1], f'Ngày dạy: {ngay_day}')
    fill_cell(tbl_info.rows[2].cells[0], 'Tổ: Tổ chuyên môn THCS')
    fill_cell(tbl_info.rows[2].cells[1], f'Lớp: {ten_lop}')

    add_p(doc, '', line_spacing=ls)

    # Title block for THCS (4 lines centered)
    clean_title = re.sub(r'^(?:Bài|BÀI)\s*\d+[\.:\s]*', '', title).strip()
    if 'tiết 0' in title.lower() or 'định hướng' in title.lower():
        title_line = "TÊN BÀI DẠY: TIẾT 0: ĐỊNH HƯỚNG MÔN HỌC"
    elif 'ôn tập' in title.lower() or 'đánh giá' in title.lower() or 'tổng kết' in title.lower() or 'luyện tập' in title.lower():
        title_line = f"TÊN BÀI DẠY: {title.upper()}"
    else:
        m = re.match(r'^(?:Bài|BÀI)\s*(\d+)', title)
        b_num = m.group(1) if m else str(tiet_ppct)
        title_line = f"TÊN BÀI DẠY: BÀI {b_num}. {clean_title.upper()}"

    add_p(doc, title_line, bold=True, size_pt=13,
          alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=ls)
    add_p(doc, 'Môn học: Robotics', bold=True, italic=True, size_pt=13,
          alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=ls)
    add_p(doc, 'Thời lượng: 1 tiết (45 phút)', bold=True, italic=True, size_pt=13,
          alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=ls)
    add_p(doc, f'Tiết theo PPCT: {tiet_ppct}', bold=True, size_pt=13,
          alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=ls)

    add_p(doc, '', line_spacing=ls)

    # I. Mục tiêu
    add_p(doc, 'I. Mục tiêu', bold=True, first_indent=INDENT_0, line_spacing=ls)

    # 1. Kiến thức
    add_p(doc, '1. Kiến thức:', bold=True, first_indent=INDENT_1, line_spacing=ls)
    kt_items = parse_yccd_bullets(yccd_raw)
    if not kt_items:
        kt_items = [f"Kiến thức cấu tạo và nguyên lý hoạt động của {title}.",
                    f"Quy trình lắp ráp và vận hành mô hình robot {title} an toàn."]
    for item in kt_items:
        add_p(doc, f'- {item}', left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    # 2. Năng lực
    add_p(doc, '2. Năng lực:', bold=True, first_indent=INDENT_1, line_spacing=ls)
    add_p(doc, '2.1. Năng lực đặc thù (Robotics):', bold=True,
          first_indent=INDENT_2, line_spacing=ls)
    for c in get_rob_competencies(grade, title, yccd_raw, kit_name):
        add_p(doc, c, left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    add_p(doc, '2.2. Năng lực số (Thông tư 02/2025 – CV 3456):', bold=True,
          first_indent=INDENT_2, line_spacing=ls)
    add_p(doc, f'- Miền III. Sáng tạo nội dung số – Thành tố 3.2. Sáng tạo sản phẩm số – Bậc {bac_nls}: Thiết kế, xây dựng và lập trình mô hình robot hoạt động chính xác. (Đạt được thông qua Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, f'- Miền IV. An toàn – Thành tố 4.3. Bảo vệ sức khỏe và môi trường số – Bậc {bac_nls}: Tuân thủ quy định sử dụng an toàn thiết bị điện tử, nguồn cấp và động cơ servo. (Đạt được thông qua Hoạt động 1, Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    add_p(doc, '2.3. Năng lực chung:', bold=True, first_indent=INDENT_2, line_spacing=ls)
    add_p(doc, '- Tự chủ và tự học: Chủ động đọc hiểu tài liệu hướng dẫn, sơ đồ lắp ráp và giải quyết vấn đề kỹ thuật phát sinh. (Đạt được thông qua Hoạt động 2, Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '- Giao tiếp và hợp tác: Phân công vai trò khoa học trong nhóm, thảo luận giải pháp kỹ thuật và hỗ trợ đồng đội. (Đạt được thông qua Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    # 3. Phẩm chất
    add_p(doc, '3. Phẩm chất:', bold=True, first_indent=INDENT_1, line_spacing=ls)
    add_p(doc, '- Chăm chỉ: Tác phong công nghiệp, tư duy khoa học và kiên trì trong thực hành chế tạo robot. (Đạt được thông qua Hoạt động 2, Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '- Trung thực: Trung thực trong báo cáo kết quả thử nghiệm và đánh giá sản phẩm của nhóm. (Đạt được thông qua Hoạt động 3)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '- Trách nhiệm: Quản lý, bảo vệ thiết bị công nghệ và giữ gìn vệ sinh phòng thực hành bộ môn. (Đạt được thông qua Hoạt động 3, Hoạt động 4)',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    # II. Thiết bị dạy học và học liệu
    add_p(doc, 'II. Thiết bị dạy học và học liệu:', bold=True,
          first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '1. Thiết bị:', bold=True, first_indent=INDENT_1, line_spacing=ls)
    add_p(doc, f'- Bộ Kit Robotics {kit_name}, máy tính giáo viên và máy tính nhóm học sinh.',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '- Phần mềm nạp lập trình, máy chiếu bài giảng 3D.',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '2. Học liệu:', bold=True, first_indent=INDENT_1, line_spacing=ls)
    add_p(doc, f'- Phiếu học tập thực hành, sơ đồ nguyên lý cơ khí và thuật toán điều khiển {title}.',
          left_indent=INDENT_BULLET, first_indent=INDENT_0, line_spacing=ls)

    # III. Tiến trình dạy học
    add_p(doc, 'III. Tiến trình dạy học', bold=True, first_indent=INDENT_0, line_spacing=ls)

    buoc1 = 'Bước 1: Chuyển giao nhiệm vụ học tập'
    buoc2 = 'Bước 2: Học sinh tiếp nhận nhiệm vụ học tập'
    buoc3 = 'Bước 3: Báo cáo kết quả hoạt động'
    buoc4 = 'Bước 4: Đánh giá kết quả thực hiện nhiệm vụ'
    buoc4_last = 'Bước 4: Giáo viên nhắc nhở nhiệm vụ về nhà'

    activities_thcs = [
        ('1. Hoạt động 1. Khởi động (Xác định vấn đề/nhiệm vụ học tập/Mở đầu) (7 phút)',
         f'Kích hoạt tư duy kỹ thuật và kết nối tình huống vào bài học {title}.',
         f'GV chiếu hình ảnh/video thực tế về {title} và đặt câu hỏi phân tích cơ chế hoạt động.',
         'Câu trả lời của HS xác định được vấn đề kỹ thuật cần giải quyết.',
         f'{buoc1}:\n- GV trình chiếu clip/hình ảnh thực tế liên quan đến {title}.\n- Yêu cầu HS quan sát và phân tích cơ chế chuyển động.\n{buoc2}:\n- HS tập trung quan sát màn chiếu, lắng nghe câu hỏi dẫn dắt.\n- HS thảo luận nhanh theo nhóm bàn, đối chiếu hiện tượng.\n{buoc3}:\n- 2 đại diện HS trả lời câu hỏi khởi động.\n- Nêu giả thuyết và nguyên lý vận hành ban đầu.\n{buoc4}:\n- GV nhận xét, chuẩn hóa kiến thức và dẫn dắt vào bài mới.',
         f'HS nhận diện được vấn đề kỹ thuật của bài học.\nHS sẵn sàng bước vào hoạt động chính.'),

        ('2. Hoạt động 2. Hình thành kiến thức mới/giải quyết vấn đề (18 phút)',
         f'Nghiên cứu sơ đồ thiết kế và tìm hiểu linh kiện cho mô hình {title}.',
         'GV yêu cầu HS tìm hiểu linh kiện khung, chốt, động cơ và cơ cấu truyền động.',
         'Sơ đồ khối nguyên lý cơ khí và danh mục linh kiện chính xác.',
         f'{buoc1}:\n- GV phát phiếu hướng dẫn và trình chiếu sơ đồ 2D/3D của {title}.\n- Hướng dẫn nhận biết chiều lắp chốt và cổng cắm động cơ.\n{buoc2}:\n- HS tiếp nhận phiếu, quan sát các góc ghép linh kiện.\n- HS đối chiếu danh mục, nhặt đúng số lượng chốt và khung nối.\n{buoc3}:\n- GV kiểm tra khay linh kiện từng nhóm tại bàn.\n- Đại diện nhóm giơ khay linh kiện để kiểm tra.\n{buoc4}:\n- GV chốt quy trình các bước lắp ráp mô hình robot.',
         f'HS nắm vững quy trình các bước lắp ráp.\nDanh mục linh kiện đầy đủ, chuẩn xác.'),

        ('3. Hoạt động 3. Luyện tập (12 phút)',
         f'Thực hành lắp ráp, lập trình và chạy thử mô hình robot {title}.',
         'GV yêu cầu các nhóm tiến hành lắp ráp phần cứng và thử nghiệm vận hành.',
         f'Mô hình robot {title} hoàn thiện, hoạt động chính xác theo yêu cầu.',
         f'{buoc1}:\n- GV giao nhiệm vụ cho từng nhóm lắp ráp mô hình {title}.\n{buoc2}:\n- HS phân công nhiệm vụ (xem sơ đồ, lắp khung, chọn chốt).\n- HS tiến hành lắp ráp cẩn thận từng chi tiết theo sơ đồ.\n{buoc3}:\n- GV yêu cầu các nhóm cấp nguồn/nạp code và cho robot vận hành thử.\n- HS đặt robot lên bàn thử nghiệm, quan sát hoạt động.\n{buoc4}:\n- GV đánh giá sản phẩm thực hành, chấm điểm tiêu chí kỹ thuật.\n- HS tinh chỉnh lại chốt nối nếu robot bị kẹt hoặc di chuyển lệch.',
         f'Mô hình robot {title} hoàn thiện.\nRobot vận hành chính xác, trơn tru.'),

        ('4. Hoạt động 4. Vận dụng (8 phút)',
         'Vận dụng sáng tạo, nâng cấp tính năng robot và dọn dẹp phòng học.',
         'GV đặt yêu cầu cải tiến tối ưu thuật toán hoặc thêm chi tiết trang trí.',
         'Báo cáo ý tưởng cải tiến mô hình và bộ Kit được sắp xếp ngăn nắp.',
         f'{buoc1}:\n- GV đặt câu hỏi: "Em có thể cải tiến cấu trúc nào để {title} hoạt động tốt hơn?"\n{buoc2}:\n- HS suy nghĩ ý tưởng sáng tạo mở rộng.\n- HS nhẹ nhàng tháo chốt, phân loại linh kiện về đúng ngăn trong hộp Kit.\n{buoc3}:\n- 1-2 HS trình bày ý tưởng nâng cấp robot trước lớp.\n{buoc4_last}:\n- GV dặn dò nhiệm vụ chuẩn bị cho bài học tiếp theo.\n- HS thu dọn bàn học, nộp lại hộp Kit Robotics ngăn nắp.',
         f'HS đề xuất được ý tưởng cải tiến mô hình robot.\nBộ Kit Robotics được phân loại và bảo quản tốt.')
    ]

    for idx, (act_title, muc_tieu, noi_dung, san_pham, gv_hs_col, ket_qua_col) in enumerate(activities_thcs, 1):
        add_p(doc, act_title, bold=True, first_indent=INDENT_1, line_spacing=ls)
        add_multi_run(doc, [
            ('a) Mục tiêu: ', True, False), (muc_tieu, False, True)
        ], first_indent=INDENT_1, line_spacing=ls)
        add_multi_run(doc, [
            ('b) Nội dung: ', True, False), (noi_dung, False, True)
        ], first_indent=INDENT_1, line_spacing=ls)
        add_multi_run(doc, [
            ('c) Sản phẩm: ', True, False), (san_pham, False, True)
        ], first_indent=INDENT_1, line_spacing=ls)
        add_multi_run(doc, [
            ('d) Tổ chức thực hiện:', True, False),
        ], first_indent=INDENT_1, line_spacing=ls)

        tbl = doc.add_table(rows=1, cols=2)
        set_borders(tbl)
        for cell in tbl.columns[0].cells:
            cell.width = Cm(9.0)
        for cell in tbl.columns[1].cells:
            cell.width = Cm(7.0)

        fill_cell(tbl.rows[0].cells[0], 'HOẠT ĐỘNG CỦA GV – HS', bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(tbl.rows[0].cells[1], 'KẾT QUẢ CẦN ĐẠT', bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

        row = tbl.add_row()
        fill_cell(row.cells[0], gv_hs_col)
        fill_cell(row.cells[1], ket_qua_col)

        add_p(doc, '', line_spacing=ls)

    # Rút kinh nghiệm & Bảng chữ ký
    add_p(doc, 'RÚT KINH NGHIỆM SAU BÀI DẠY:', bold=True,
          first_indent=INDENT_0, line_spacing=ls)
    add_p(doc, '...........................................................................................................................', line_spacing=ls)
    add_p(doc, '...........................................................................................................................', line_spacing=ls)
    add_p(doc, '', line_spacing=ls)

    tbl_sign = doc.add_table(rows=3, cols=3)
    set_no_borders(tbl_sign)
    for i, txt in enumerate(['DUYỆT CỦA BGH', 'DUYỆT CỦA TỔ CM', 'NGƯỜI SOẠN']):
        fill_cell(tbl_sign.rows[0].cells[i], txt, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    for i in range(3):
        fill_cell(tbl_sign.rows[1].cells[i], '(Ký, ghi rõ họ tên)',
                  italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i in range(3):
        fill_cell(tbl_sign.rows[2].cells[i], '\n\n\n',
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(tbl_sign.rows[2].cells[2], '\n\n\nĐậu Đình Nguyên', bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER)

    return doc


# ─── SYNC KHDH TABLES ──────────────────────────────────────────────────────
def sync_khdh_table(table, lessons, is_rotation=False):
    """
    Update a 4-col or 5-col KHDH table with 35 lessons.
    Format: STT | Bài học (1) | Số tiết (2) | [Tiết theo PPCT] | Yêu cầu cần đạt (3)
    """
    cols = len(table.columns)
    has_ppct_col = (cols >= 5)

    # Keep header row, remove existing data rows
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)

    for idx, (title, yccd) in enumerate(lessons, start=1):
        row = table.add_row()
        cells = row.cells
        
        # Col 0: STT
        fill_cell(cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        # Col 1: Bài học
        is_dg = any(k in title.lower() for k in ['định kỳ', 'danh gia dinh ky', 'ôn tập', 'tổng kết'])
        fill_cell(cells[1], title, bold=is_dg)
        # Col 2: Số tiết
        fill_cell(cells[2], "1", align=WD_ALIGN_PARAGRAPH.CENTER)
        
        if has_ppct_col:
            # Col 3: Tiết theo PPCT
            fill_cell(cells[3], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
            # Col 4: YCCD
            fill_cell(cells[4], yccd)
        else:
            # Col 3: YCCD
            fill_cell(cells[3], yccd)


def sync_all_khdh_files():
    print("\n" + "=" * 60)
    print(" 1. ĐỒNG BỘ VÀ CẬP NHẬT CÁC FILE KẾ HOẠCH DẠY HỌC ROBOTICS")
    print("=" * 60)

    # 1. Update KHDH Master
    if os.path.exists(KHDH_MASTER):
        doc = Document(KHDH_MASTER)
        # Tables 4 to 11 correspond to Grade 1 to 8
        for g in range(1, 9):
            t_idx = 3 + g
            if t_idx < len(doc.tables):
                lessons = get_lessons_for_grade(g)
                sync_khdh_table(doc.tables[t_idx], lessons, is_rotation=(g >= 5))
                print(f"  [+] KHDH Master: Đã cập nhật Bảng Lớp {g} (Table {t_idx}) -> {len(lessons)} tiết.")
        doc.save(KHDH_MASTER)
        print("  => Đã lưu KHDH Master thành công.")

    # 2. Update KHDH TH
    if os.path.exists(KHDH_TH):
        doc = Document(KHDH_TH)
        for g in range(1, 6):
            t_idx = 3 + g
            if t_idx < len(doc.tables):
                lessons = get_lessons_for_grade(g)
                sync_khdh_table(doc.tables[t_idx], lessons, is_rotation=(g >= 5))
                print(f"  [+] KHDH TH: Đã cập nhật Bảng Lớp {g} (Table {t_idx}) -> {len(lessons)} tiết.")
        doc.save(KHDH_TH)
        print("  => Đã lưu KHDH TH thành công.")

    # 3. Update KHDH THCS
    if os.path.exists(KHDH_THCS):
        doc = Document(KHDH_THCS)
        for g_idx, g in enumerate([6, 7, 8]):
            t_idx = 4 + g_idx
            if t_idx < len(doc.tables):
                lessons = get_lessons_for_grade(g)
                sync_khdh_table(doc.tables[t_idx], lessons, is_rotation=True)
                print(f"  [+] KHDH THCS: Đã cập nhật Bảng Lớp {g} (Table {t_idx}) -> {len(lessons)} tiết.")
        doc.save(KHDH_THCS)
        print("  => Đã lưu KHDH THCS thành công.")

    # 4. Update individual files in KHDH_DIR
    for g in range(1, 9):
        path = os.path.join(KHDH_DIR, f"Kế hoạch dạy học môn Robotics - Lớp {g} - 2026 - 2027.docx")
        if os.path.exists(path):
            doc = Document(path)
            # Find the main lesson table (len > 10)
            target_table = None
            for t in doc.tables:
                if len(t.rows) > 10:
                    target_table = t
                    break
            if target_table:
                lessons = get_lessons_for_grade(g)
                sync_khdh_table(target_table, lessons, is_rotation=(g >= 5))
                doc.save(path)
                print(f"  [+] KHDH Từng lớp: Lớp {g} -> Đã đồng bộ {len(lessons)} tiết.")


# ─── REGENERATE KHBD ROBOTICS ─────────────────────────────────────────────
def regenerate_all_khbd_robotics():
    print("\n" + "=" * 60)
    print(" 2. TÁI TẠO TOÀN BỘ KHBD ROBOTICS (LỚP 1 - 8) THEO TUẦN CHUẨN")
    print("=" * 60)

    if os.path.exists(OUT_BASE):
        shutil.rmtree(OUT_BASE)
        print(f"  [+] Đã xóa sạch thư mục cũ: {OUT_BASE}")
    os.makedirs(OUT_BASE, exist_ok=True)

    created_total = 0

    for grade in range(1, 9):
        g_str = str(grade)
        day_of_week, ten_lop = ROB_SCHEDULE[g_str]
        kit_name = get_kit_name(grade)
        is_rotation = (grade >= 5)
        lessons = get_lessons_for_grade(grade)

        print(f"\n---> Đang tạo KHBD Robotics Lớp {grade} ({ten_lop}, {len(lessons)} bài)...")

        for idx, (title, yccd) in enumerate(lessons, start=1):
            # Tính Tuần và Tiết PPCT chuẩn:
            if is_rotation:
                # Lớp 5, 6, 7, 8 (Rotation Tuần LẺ):
                # idx = 1 (Tiết 0) -> Tuần 1
                # idx = 2, 3 -> Tuần 3
                # idx = 4, 5 -> Tuần 5
                # idx = 6, 7 -> Tuần 7
                # idx = 8, 9 (Ôn tập, ĐGĐK 1) -> Tuần 9
                # idx = 10, 11 -> Tuần 11
                # idx = 12, 13 -> Tuần 13
                # idx = 14, 15 -> Tuần 15
                # idx = 16, 17 -> Tuần 17
                # idx = 18, 19 (Ôn tập, ĐGĐK 2) -> Tuần 19
                # idx = 20, 21 -> Tuần 21
                # idx = 22, 23 -> Tuần 23
                # idx = 24, 25 -> Tuần 25
                # idx = 26, 27 (Ôn tập, ĐGĐK 3) -> Tuần 27
                # idx = 28, 29 -> Tuần 29
                # idx = 30, 31 -> Tuần 31
                # idx = 32, 33 (Ôn tập, ĐGĐK 4) -> Tuần 33
                # idx = 34, 35 (Tổng kết năm học) -> Tuần 35
                if idx == 1:
                    tuan_so = 1
                    tiet_ppct = 0
                else:
                    tuan_so = 2 * ((idx - 2) // 2) + 3
                    tiet_ppct = idx - 1
            else:
                # Lớp 1, 2, 3, 4 (Mỗi tuần 1 tiết, Tuần 1 đến 35):
                # idx = 1 (Tiết 0) -> Tuần 1
                # idx = 2 (Bài 1) -> Tuần 2
                # ...
                # idx = 9 (Ôn tập ĐGĐK 1) -> Tuần 9
                # idx = 10 (ĐGĐK 1) -> Tuần 10
                # ...
                # idx = 18 (Ôn tập ĐGĐK 2) -> Tuần 18
                # idx = 19 (ĐGĐK 2) -> Tuần 19
                # ...
                # idx = 27 (Ôn tập ĐGĐK 3) -> Tuần 27
                # idx = 28 (ĐGĐK 3) -> Tuần 28
                # ...
                # idx = 33 (Ôn tập ĐGĐK 4) -> Tuần 33
                # idx = 34 (ĐGĐK 4) -> Tuần 34
                # idx = 35 (Tổng kết năm học) -> Tuần 35
                tuan_so = idx
                tiet_ppct = 0 if idx == 1 else (idx - 1)

            tuan_folder = f"Tuần_{tuan_so:02d}"
            out_dir = os.path.join(OUT_BASE, f"Lớp_{grade}", tuan_folder)
            os.makedirs(out_dir, exist_ok=True)

            safe_title = sanitize(title)
            filename = f"KHBD_Robotics_Lớp_{grade}_Tiet{tiet_ppct:02d}_{safe_title}.docx"
            out_file = os.path.join(out_dir, filename)

            # Build Doc
            if grade >= 6:
                doc = build_khbd_thcs(grade, ten_lop, title, tiet_ppct, yccd, tuan_so, day_of_week, kit_name)
            else:
                doc = build_khbd_th(grade, ten_lop, title, tiet_ppct, yccd, tuan_so, day_of_week, kit_name)

            try:
                doc.save(out_file)
                created_total += 1
                # print summary
                if 'đánh giá' in title.lower() or 'tiết 0' in title.lower() or 'tổng kết' in title.lower():
                    print(f"  [★] Lớp {grade} -> {tuan_folder} -> {filename}")
            except Exception as e:
                print(f"  [!] Lỗi khi lưu {out_file}: {e}")

    print("\n" + "=" * 60)
    print(f" HOÀN TẤT TẠO {created_total} FILE KHBD ROBOTICS CHUẨN ĐÚNG CÁC TUẦN KIỂM TRA ĐÁNH GIÁ!")
    print("=" * 60)


if __name__ == '__main__':
    sync_all_khdh_files()
    regenerate_all_khbd_robotics()
