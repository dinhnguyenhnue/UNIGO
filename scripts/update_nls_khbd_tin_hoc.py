# -*- coding: utf-8 -*-
"""
Cập nhật mục 2.2. Năng lực số (NLS) theo mã chuẩn CB (Chuẩn Bậc)
cho toàn bộ 318 file KHBD Tin học (Tiền TH, Lớp 1 - 8)
Tuân thủ Thông tư 02/2025 – CV 3456/BGDĐT và Quy chuẩn UNIGO (AGENTS.md).
"""
import os, sys, re, json, time
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding='utf-8')

# 1. Tải dữ liệu CV 3456
JSON_PATH = r'D:\UNIGO\.agents\skills\tao-khbd\references\cv3456_full_data.json'
with open(JSON_PATH, 'r', encoding='utf-8') as f:
    CV_DATA = json.load(f)

def get_descriptor(comp_code, grade_level, letter):
    target_key = None
    for k in CV_DATA:
        if k.startswith(comp_code + '.') or k.startswith(comp_code + ' '):
            target_key = k
            break
    if not target_key:
        raise ValueError(f'Unknown component code: {comp_code}')
    
    desc_text = CV_DATA[target_key]['descriptors'][grade_level]
    bullets = [b.strip() for b in desc_text.split('\n') if b.strip().startswith('-')]
    idx = ord(letter.lower()) - ord('a')
    if idx >= len(bullets):
        idx = len(bullets) - 1
        letter = chr(ord('a') + idx)
    
    raw = bullets[idx]
    clean = re.sub(r'^[-\s]+', '', raw).strip()
    clean = clean.rstrip(';,').strip()
    if not clean.endswith('.'):
        clean += '.'
    clean = clean[0].upper() + clean[1:]
    return clean, letter

# 2. Phân loại nội dung bài học & ánh xạ NLS
def classify_lesson(grade, filename, text_content):
    fn = filename.lower()
    tc = text_content.lower()
    
    # Bậc theo khối lớp
    if grade in ['Tiền_tiểu_học', 'Lớp_1', 'Lớp_2', 'Lớp_3']:
        lvl = 'L1-3'
        cb = 'CB1'
    elif grade in ['Lớp_4', 'Lớp_5']:
        lvl = 'L4-5'
        cb = 'CB2'
    elif grade in ['Lớp_6', 'Lớp_7']:
        lvl = 'L6-7'
        cb = 'CB3'
    else: # Lớp_8
        lvl = 'L8-9'
        cb = 'CB4'
        
    # 1. Đánh giá định kỳ & Ôn tập
    if any(k in fn for k in ['anh_gia_inh_ky', 'danh_gia_dinh_ky', 'on_tap', 'kiem_thu']):
        ctx1 = 'Tự đánh giá và hoàn thiện các kỹ năng số đã học trong kỳ'
        ctx2 = 'Kiểm tra, đánh giá tính chính xác của kết quả thực hành trên máy'
        return lvl, cb, ('5.4', 'a', ctx1, 'Hoạt động 1, Hoạt động 2'), ('1.2', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # 2. Tổng kết năm học / Ngày hội sáng tạo số / Hướng nghiệp
    if any(k in fn for k in ['tong_ket', 'ngay_hoi', 'young_innovator', 'nghe_nghiep', 'hoan_thien_san_pham']):
        if 'nghe_nghiep' in fn:
            ctx1 = 'Tìm hiểu tác động của công nghệ số và định hướng nghề nghiệp tương lai'
            ctx2 = 'Đánh giá các kỹ năng số cần thiết cho nghề nghiệp trong thời đại số'
            return lvl, cb, ('2.3', 'a', ctx1, 'Hoạt động 2'), ('5.4', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')
        else:
            ctx1 = 'Sử dụng sáng tạo các công cụ và sản phẩm số đã học trong năm'
            ctx2 = 'Xác định các kỹ năng số đã hoàn thành và hướng cải thiện tiếp theo'
            return lvl, cb, ('5.3', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('5.4', 'a', ctx2, 'Hoạt động 4')

    # 3. AI / Trí tuệ nhân tạo (dùng regex tránh nhầm 'bai_')
    if re.search(r'(^|[\s_])ai([\s_]|$)', fn) or any(k in fn for k in ['ai quanh em', 'may thong minh', 'con_nguoi_va_may_thong_minh']):
        if 'sai' in fn or 'kiem_tra' in fn:
            ctx1 = 'Nhận biết khả năng xử lý và giới hạn của hệ thống trí tuệ nhân tạo'
            ctx2 = 'Kiểm tra, phát hiện các điểm chưa chính xác trong dữ liệu do AI tạo ra'
            return lvl, cb, ('6.1', 'a', ctx1, 'Hoạt động 2'), ('6.3', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')
        else:
            ctx1 = 'Nhận biết sự hiện diện và ứng dụng cơ bản của trí tuệ nhân tạo trong đời sống'
            ctx2 = 'Tương tác và trải nghiệm với công cụ trí tuệ nhân tạo đơn giản phục vụ học tập'
            return lvl, cb, ('6.1', 'a', ctx1, 'Hoạt động 2'), ('6.2', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # 4. Tư thế ngồi học / An toàn sức khỏe
    if any(k in fn for k in ['tu_the', 'ngoi_may_tinh_an_toan']):
        ctx1 = 'Thực hiện tư thế ngồi học máy tính đúng cách để bảo vệ thị lực và sức khỏe'
        ctx2 = 'Tuân thủ quy tắc an toàn điện và bảo vệ thiết bị trong phòng thực hành'
        return lvl, cb, ('4.3', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('4.1', 'a', ctx2, 'Hoạt động 1, Hoạt động 4')

    # 5. Gõ bàn phím (rèn luyện thao tác)
    if any(k in fn for k in ['go_ban_phim', 'go_chu', 'go_ten', 'go_cau', 'nhung_phim_em_can_biet', 'ban_phim_chu_va_so']):
        ctx1 = 'Thao tác gõ bàn phím đúng cách, rèn luyện thói quen tư thế và vị trí ngón tay'
        ctx2 = 'Sử dụng bàn phím để nhập dữ liệu chữ và số chính xác vào phần mềm'
        return lvl, cb, ('4.3', 'a', ctx1, 'Hoạt động 2'), ('5.1', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # 6. Bản quyền / Đạo đức / Văn hóa mạng
    if any(k in fn for k in ['ban_quyen', 'duoc_phep', 'uoc_phep', 'dao_duc', 'van_hoa_trong_su_dung']):
        if any(k in fn for k in ['ban_quyen', 'duoc_phep', 'uoc_phep']):
            ctx1 = 'Tôn trọng quyền tác giả và tuân thủ quy định bản quyền khi sử dụng phần mềm, nội dung số'
            ctx2 = 'Tuân thủ quy tắc ứng xử chuẩn mực và bảo vệ dữ liệu khi khai thác thông tin'
            return lvl, cb, ('3.3', 'a', ctx1, 'Hoạt động 2'), ('2.5', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')
        else:
            ctx1 = 'Thực hiện văn hóa ứng xử văn minh và đạo đức trong môi trường số'
            ctx2 = 'Bảo vệ dữ liệu cá nhân và thông tin riêng tư của bản thân và người khác'
            return lvl, cb, ('2.5', 'a', ctx1, 'Hoạt động 2'), ('4.2', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # 7. An toàn thông tin / Quyền riêng tư
    if any(k in fn for k in ['an_toan_thong_tin', 'bao_ve_thong_tin', 'su_dung_cong_nghe_an_toan']):
        ctx1 = 'Nhận biết các nguy cơ mất an toàn thông tin và bảo vệ dữ liệu cá nhân trên môi trường số'
        ctx2 = 'Tuân thủ các biện pháp bảo vệ thiết bị và phòng tránh phần mềm độc hại'
        return lvl, cb, ('4.2', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('4.1', 'a', ctx2, 'Hoạt động 4')

    # 8. Mạng xã hội / Email / Giao tiếp trực tuyến
    if any(k in fn for k in ['thu_ien_tu', 'thu_dien_tu', 'mang_xa_hoi', 'ung_xu_tren_mang', 'kenh_trao_oi', 'trao_doi_thong_tin']):
        ctx1 = 'Sử dụng công nghệ số để gửi, nhận và tương tác thông tin trao đổi phục vụ học tập'
        ctx2 = 'Áp dụng các quy tắc ứng xử lịch sự, văn minh khi giao tiếp trên môi trường mạng'
        return lvl, cb, ('2.1', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('2.5', 'a', ctx2, 'Hoạt động 4')

    # 9. Hợp tác / Làm việc nhóm
    if any(k in fn for k in ['chia_se_va_lam_viec_nhom', 'lam_viec_cung_ban', 'thiet_ke_san_pham_nhom']):
        ctx1 = 'Cùng hợp tác với các bạn thông qua công cụ số để hoàn thành nhiệm vụ nhóm'
        ctx2 = 'Chia sẻ thông tin, sản phẩm số an toàn và đúng mục đích'
        return lvl, cb, ('2.4', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('2.2', 'a', ctx2, 'Hoạt động 4')

    # 10. Bảng tính (Excel)
    if any(k in fn for k in ['bang_tinh', 'tinh_toan_tu_ong', 'cong_cu_ho_tro_tinh_toan', 'sap_xep_va_loc', 'bieu_o', 'bieu_do']):
        if 'bieu_o' in fn or 'bieu_do' in fn:
            ctx1 = 'Trình bày và trực quan hóa dữ liệu bảng tính bằng các biểu đồ phù hợp'
            ctx2 = 'Quản lý, tổ chức và định dạng dữ liệu trong bảng tính'
            return lvl, cb, ('3.1', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('1.3', 'a', ctx2, 'Hoạt động 1, Hoạt động 4')
        else:
            ctx1 = 'Tổ chức, xử lý và quản lý dữ liệu hiệu quả trên phần mềm bảng tính'
            ctx2 = 'Sử dụng công thức và tính toán tự động để giải quyết bài toán thực tế'
            return lvl, cb, ('1.3', 'a', ctx1, 'Hoạt động 2'), ('5.1', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # 11. Bài trình chiếu (PowerPoint)
    if any(k in fn for k in ['trinh_chieu', 'trang_chieu', 'ban_mau_cho_bai_trinh_chieu', 'hieu_ung_chuyen_trang']):
        ctx1 = 'Thiết kế, định dạng nội dung và xây dựng bài trình chiếu hấp dẫn, khoa học'
        ctx2 = 'Tích hợp hài hòa văn bản, hình ảnh và hiệu ứng đa phương tiện vào trang chiếu'
        return lvl, cb, ('3.1', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('3.2', 'a', ctx2, 'Hoạt động 4')

    # 12. Soạn thảo văn bản (Word)
    if any(k in fn for k in ['soan_thao', 'van_ban', 'chinh_sua_van_ban', 'inh_dang_van_ban', 'dinh_dang_ki_tu', 'tim_kiem_va_thay_the', 'dau_trang_chan_trang', 'au_trang_chan_trang', 'danh_sach_dang_liet_ke', 'so_o_tu_duy', 'trinh_bay_thong_tin_o_dang_bang', 'thuc_hanh_tong_hop']):
        ctx1 = 'Soạn thảo, định dạng văn bản và trình bày nội dung số rõ ràng, thẩm mỹ'
        ctx2 = 'Tổ chức, lưu trữ và quản lý tệp văn bản trong môi trường máy tính'
        return lvl, cb, ('3.1', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('1.3', 'a', ctx2, 'Hoạt động 1, Hoạt động 4')

    # 13. Đồ họa / Kể chuyện số / Đa phương tiện / Vẽ tranh
    if any(k in fn for k in ['o_hoa', 'do_hoa', 'ke_chuyen', 'nhan_vat', 'boi_canh', 'animation', 'am_thanh', 'buc_tranh', 'hoat_hinh', 'chu_hinh_va_am_thanh', 'mau_sac', 'san_pham_au_tien', 'cau_chuyen', 'ngoi_nha_va_khu_vuon', 'nha_sang_tao_so']):
        ctx1 = 'Sáng tạo và phát triển các sản phẩm đa phương tiện bằng phần mềm số'
        ctx2 = 'Sử dụng sáng tạo các công cụ đồ họa và hiệu ứng để thể hiện ý tưởng'
        return lvl, cb, ('3.1', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('5.3', 'a', ctx2, 'Hoạt động 4')

    # 14. Lập trình / Thuật toán / Tư duy logic / Scratch / Mê cung / Các bước
    if any(k in fn for k in ['lap_trinh', 'thuat_toan', 'tuan_tu', 'cau_truc_lap', 're_nhanh', 'bien_trong_chuong_trinh', 'quy_luat', 'me_cung', 'sua_loi_uong_i', 'chuoi_lenh', 'chay_thu_va_sua_loi', 'chi_uong', 'theo_tung_buoc', 'nhieu_buoc', 'choi_voi_may_tinh', 'trai_phai', 'nghe_lenh', 'lap_lai', 'lenh_lap', 'uong_i_en', 'thuc_hien_cong_viec']):
        ctx1 = 'Thiết kế thuật toán từng bước và xây dựng khối lệnh điều khiển chương trình'
        ctx2 = 'Vận dụng tư duy logic để chạy thử, phát hiện và sửa lỗi chương trình'
        return lvl, cb, ('3.4', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('5.1', 'a', ctx2, 'Hoạt động 4')

    # 15. Robotics & Cảm biến
    if any(k in fn for k in ['robot', 'cam_bien', 'ieu_khien_robot', 'dieu_khien_robot']):
        ctx1 = 'Xác định giải pháp công nghệ và mô hình điều khiển phù hợp cho robot'
        ctx2 = 'Lập trình và vận hành robot thực hiện nhiệm vụ điều khiển theo yêu cầu'
        return lvl, cb, ('5.2', 'a', ctx1, 'Hoạt động 2'), ('3.4', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # 16. Mạng Internet / Web / Tìm kiếm thông tin / WWW
    if any(k in fn for k in ['internet', 'trang_web', 'website', 'tim_kiem_thong_tin', 'mang_may_tinh', 'mang_thong_tin_toan_cau', 'kham_pha_thong_tin', 'the_gioi_tu_nhien', 'khai_thac_thong_tin_so', 'thong_tin_trong_moi_truong_so', 'thong_tin_quanh_em']):
        ctx1 = 'Duyệt, tìm kiếm và thu thập dữ liệu, thông tin hiệu quả trên môi trường mạng'
        ctx2 = 'Đánh giá độ tin cậy và chọn lọc thông tin chính xác từ các nguồn số'
        return lvl, cb, ('1.1', 'b', ctx1, 'Hoạt động 2'), ('1.2', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # 17. Quản lý Tệp / Thư mục / Cây thư mục / Dữ liệu & Thông tin
    if any(k in fn for k in ['tep', 'thu_muc', 'cay_thu_muc', 'thong_tin_va_du_lieu', 'thong_tin_va_quyet_inh', 'xu_li_thong_tin', 'thong_tin_trong_may_tinh', 'bieu_dien_du_lieu', 'sap_xep_e_de_tim', 'dat_ten_va_luu', 'at_ten_va_luu', 'mo_lai_va_sap_xep', 'san_pham_co_to_chuc', 'thu_nhan_va_xu_ly', 'quan_ly_du_lieu']):
        ctx1 = 'Tổ chức, lưu trữ, đổi tên và sắp xếp dữ liệu, tệp tin và thư mục khoa học'
        ctx2 = 'Nhận diện, tìm kiếm và truy xuất thông tin, dữ liệu trong môi trường máy tính'
        return lvl, cb, ('1.3', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('1.1', 'c', ctx2, 'Hoạt động 1, Hoạt động 4')

    # 18. Phần cứng / Thiết bị / Chuột / Bàn phím / Hệ điều hành
    if any(k in fn for k in ['may_tinh', 'bo_phan', 'thiet_bi', 'chuot', 'ban_phim', 'phan_cung_va_phan_mem', 'phan_mem_may_tinh', 'luoc_su_cong_cu_tinh_toan', 'the_gioi_cong_nghe', 'keo_tha', 'mo_ong_va_chuyen_oi', 'chuyen_oi_ung_dung']):
        ctx1 = 'Nhận biết, kết nối và sử dụng các thành phần phần cứng và thiết bị số'
        ctx2 = 'Thực hiện thao tác sử dụng thiết bị đúng quy cách và bảo vệ an toàn cho thiết bị'
        return lvl, cb, ('5.1', 'a', ctx1, 'Hoạt động 2'), ('4.1', 'a', ctx2, 'Hoạt động 3, Hoạt động 4')

    # Mặc định an toàn
    ctx1 = 'Khai thác công cụ số và phần mềm học tập để giải quyết nhiệm vụ bài học'
    ctx2 = 'Áp dụng kỹ năng sử dụng thiết bị số an toàn, bảo vệ dữ liệu học tập'
    return lvl, cb, ('5.3', 'a', ctx1, 'Hoạt động 2, Hoạt động 3'), ('4.1', 'a', ctx2, 'Hoạt động 4')

# 3. Thiết lập định dạng đoạn văn bullet chuẩn AGENTS.md
def format_bullet_paragraph(p, text):
    p.text = text
    p.paragraph_format.left_indent = 540000 # EMU (~0.59 inch)
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is not None:
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')
            rFonts.set(qn('w:cs'), 'Times New Roman')

# 4. Hàm lưu file an toàn
def save_doc_safe(doc, path, max_retries=3):
    for attempt in range(max_retries):
        try:
            doc.save(path)
            return True
        except PermissionError:
            print(f'  [WAIT] File đang mở hoặc bị khóa: {os.path.basename(path)}. Thử lại sau 2s ({attempt+1}/{max_retries})...')
            time.sleep(2)
        except Exception as e:
            print(f'  [ERROR] Lưu file thất bại: {e}')
            return False
    return False

# 5. Xử lý 1 file KHBD
def process_khbd_file(file_path, tin_dir):
    rel = os.path.relpath(file_path, tin_dir)
    grade = rel.split(os.sep)[0]
    fname = os.path.basename(file_path)
    
    doc = Document(file_path)
    
    # Tìm 2.2
    nls_idx = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if '2.2' in t and 'Năng lực số' in t:
            nls_idx = i
            break
            
    if nls_idx is None:
        return False, f'Không tìm thấy mục 2.2 trong {rel}'
        
    # Tìm 2.3
    next_idx = None
    for j in range(nls_idx + 1, len(doc.paragraphs)):
        t = doc.paragraphs[j].text.strip()
        if t.startswith('2.3') or t.startswith('2.4') or t.startswith('3.') or t.startswith('II.'):
            next_idx = j
            break
            
    if next_idx is None:
        return False, f'Không tìm thấy mục kết thúc 2.3 trong {rel}'
        
    # Phân loại và tạo nội dung NLS
    lvl, cb, p_info, s_info = classify_lesson(grade, fname, '')
    d1, l1 = get_descriptor(p_info[0], lvl, p_info[1])
    d2, l2 = get_descriptor(s_info[0], lvl, s_info[1])
    
    # Định dạng bullet đúng chuẩn
    b1_text = f'- {p_info[0]}.{cb}{l1}: {d1} ({p_info[2]}). (Đạt được thông qua {p_info[3]})'
    b2_text = f'- {s_info[0]}.{cb}{l2}: {d2} ({s_info[2]}). (Đạt được thông qua {s_info[3]})'
    
    # Cập nhật paragraphs trong docx
    p1 = doc.paragraphs[nls_idx + 1]
    format_bullet_paragraph(p1, b1_text)
    
    old_count = next_idx - (nls_idx + 1)
    if old_count == 1:
        # Chèn thêm 1 paragraph cho bullet 2
        new_p_elm = OxmlElement('w:p')
        p1._element.addnext(new_p_elm)
        p2 = Paragraph(new_p_elm, doc)
        format_bullet_paragraph(p2, b2_text)
    elif old_count >= 2:
        p2 = doc.paragraphs[nls_idx + 2]
        format_bullet_paragraph(p2, b2_text)
        # Xóa các đoạn thừa nếu có > 2 đoạn
        for extra_idx in range(nls_idx + 3, next_idx):
            p_extra = doc.paragraphs[extra_idx]
            p_extra._element.getparent().remove(p_extra._element)
            
    # Lưu file
    if save_doc_safe(doc, file_path):
        return True, f'{p_info[0]}.{cb}{l1} + {s_info[0]}.{cb}{l2}'
    else:
        return False, f'Lỗi lưu file {rel}'

def main():
    tin_dir = r'D:\UNIGO\KHBD_Tin_học'
    khbd_files = []
    for root, dirs, files in os.walk(tin_dir):
        for f in files:
            if f.startswith('KHBD_') and f.endswith('.docx') and not f.startswith('~$'):
                khbd_files.append(os.path.join(root, f))
                
    total = len(khbd_files)
    print(f'=== BẮT ĐẦU CẬP NHẬT NLS CHO {total} FILE KHBD TIN HỌC ===\n')
    
    success = 0
    fail = 0
    grade_stats = {}
    
    for idx, f in enumerate(khbd_files, 1):
        rel = os.path.relpath(f, tin_dir)
        grade = rel.split(os.sep)[0]
        ok, msg = process_khbd_file(f, tin_dir)
        if ok:
            success += 1
            grade_stats[grade] = grade_stats.get(grade, 0) + 1
            if idx % 20 == 0 or idx == total:
                print(f'[{idx}/{total}] Đã xử lý {success} files OK: {rel} ({msg})')
        else:
            fail += 1
            print(f'[{idx}/{total}] THẤT BẠI: {rel} - {msg}')
            
    print('\n=== KẾT QUẢ CẬP NHẬT ===')
    print(f'Thành công: {success}/{total} files')
    print(f'Thất bại: {fail} files')
    print('\nThống kê theo khối lớp:')
    for g, c in sorted(grade_stats.items()):
        print(f'  {g}: {c} files')

if __name__ == '__main__':
    main()
