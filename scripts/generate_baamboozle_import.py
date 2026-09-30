# -*- coding: utf-8 -*-
"""
Tự động trích xuất bộ câu hỏi ôn tập Tuần 9 môn Tin học (Lớp 3 -> Lớp 8)
và tạo các file import chuẩn form Baamboozle:
1. Dạng Trắc nghiệm (Multiple Choice): "Câu hỏi","Đáp án ĐÚNG","Đáp án sai 1","Đáp án sai 2","Đáp án sai 3"
   (Tuân thủ quy tắc Baamboozle: The first answer in the multiple choice question must be the correct answer)
2. Dạng Cổ điển (Classic Q&A): "Câu hỏi","Đáp án đúng"
3. Dạng Cổ điển kèm phương án A,B,C,D: "Câu hỏi [A... B... C... D...]","Đáp án đúng"
Hỗ trợ:
- Trọn bộ 40 câu
- Bộ 24 câu (Game 1) & Bộ 16 câu (Game 2) tối ưu cho tài khoản Baamboozle tiêu chuẩn 24 ô
- Bản có số thứ tự "Câu X: " và Bản không có số thứ tự (phù hợp khi bật đảo ngẫu nhiên)
"""

import os
import sys
import re
import csv
import docx

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"D:\UNIGO"
OUT_DIR = os.path.join(BASE_DIR, "Baamboozle_Import_Tuan_9")
os.makedirs(OUT_DIR, exist_ok=True)

GRADES = [3, 4, 5, 6, 7, 8]

def extract_options(lines):
    combined = ' ' + ' '.join(lines)
    matches = list(re.finditer(r'(?:^|\s+)([A-D])\.\s*', combined))
    opts = {}
    for i, m in enumerate(matches):
        letter = m.group(1)
        start = m.end()
        end = matches[i+1].start() if i + 1 < len(matches) else len(combined)
        opts[letter] = combined[start:end].strip()
    return opts

def parse_docx(path):
    doc = docx.Document(path)
    
    # 1. Trích xuất bảng đáp án
    answers = {}
    for table in doc.tables:
        if len(table.rows) == 2:
            r0 = [c.text.strip() for c in table.rows[0].cells]
            r1 = [c.text.strip() for c in table.rows[1].cells]
            if any('Câu' in x for x in r0):
                for q, a in zip(r0, r1):
                    m = re.search(r'\d+', q)
                    if m:
                        answers[int(m.group())] = a.strip().upper()
                        
    # 2. Trích xuất câu hỏi và các phương án
    questions = []
    curr_q = None
    curr_text = []
    curr_opt_lines = []
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if 'PHẦN II' in txt or 'TỰ LUẬN' in txt or 'ĐÁP ÁN' in txt or 'BẢNG ĐÁP ÁN' in txt:
            break
            
        m_q = re.match(r'^Câu\s+(\d+)[:.]\s*(.*)', txt)
        if m_q:
            if curr_q is not None:
                opts = extract_options(curr_opt_lines)
                questions.append({
                    'num': curr_q,
                    'text': ' '.join(curr_text).strip(),
                    'opts': opts,
                    'ans': answers.get(curr_q, '')
                })
            curr_q = int(m_q.group(1))
            curr_text = [m_q.group(2).strip()]
            curr_opt_lines = []
            continue
            
        if curr_q is not None:
            if re.search(r'(?:^|\s+)[A-D]\.\s*', txt):
                curr_opt_lines.append(txt)
            else:
                if not curr_opt_lines:
                    curr_text.append(txt)
                    
    if curr_q is not None:
        opts = extract_options(curr_opt_lines)
        questions.append({
            'num': curr_q,
            'text': ' '.join(curr_text).strip(),
            'opts': opts,
            'ans': answers.get(curr_q, '')
        })
        
    return questions

def write_csv_rows(filepath, rows, delimiter=",", quoting=csv.QUOTE_ALL, encoding="utf-8"):
    with open(filepath, "w", encoding=encoding, newline="") as f:
        writer = csv.writer(f, delimiter=delimiter, quoting=quoting)
        for r in rows:
            writer.writerow(r)

def generate_grade_files(grade, questions, out_folder):
    os.makedirs(out_folder, exist_ok=True)
    
    def build_mc_rows(q_list, with_prefix=True):
        rows = []
        for q in q_list:
            q_text = f"Câu {q['num']}: {q['text']}" if with_prefix else q['text']
            correct_letter = q['ans']
            correct_ans_text = q['opts'].get(correct_letter, "")
            wrong_answers = [v for k, v in q['opts'].items() if k != correct_letter]
            rows.append([q_text, correct_ans_text] + wrong_answers)
        return rows

    def build_classic_rows(q_list, with_prefix=True):
        rows = []
        for q in q_list:
            q_text = f"Câu {q['num']}: {q['text']}" if with_prefix else q['text']
            correct_letter = q['ans']
            correct_ans_text = q['opts'].get(correct_letter, "")
            rows.append([q_text, correct_ans_text])
        return rows

    def build_mc_prompt_rows(q_list):
        rows = []
        for q in q_list:
            opt_str = " | ".join([f"{k}. {v}" for k, v in sorted(q['opts'].items())])
            q_full = f"Câu {q['num']}: {q['text']} [{opt_str}]"
            correct_letter = q['ans']
            ans_full = f"{correct_letter}. {q['opts'].get(correct_letter, '')}"
            rows.append([q_full, ans_full])
        return rows

    # 1. Trắc nghiệm Multiple Choice (Toàn bộ 40 câu)
    mc_all_prefix = build_mc_rows(questions, with_prefix=True)
    mc_all_noprefix = build_mc_rows(questions, with_prefix=False)
    
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Trac_nghiem_Lop_{grade}_40cau_Comma.txt"), mc_all_noprefix, delimiter=",", quoting=csv.QUOTE_ALL)
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Trac_nghiem_Lop_{grade}_40cau_CoSTT_Comma.txt"), mc_all_prefix, delimiter=",", quoting=csv.QUOTE_ALL)
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Trac_nghiem_Lop_{grade}_40cau.csv"), mc_all_noprefix, delimiter=",", quoting=csv.QUOTE_ALL, encoding="utf-8-sig")

    # 2. Bộ 24 câu (Game 1 - Khuyên dùng cho Baamboozle bản Free)
    q24 = questions[:24]
    mc_24 = build_mc_rows(q24, with_prefix=False)
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Trac_nghiem_Lop_{grade}_24cau_Game1.txt"), mc_24, delimiter=",", quoting=csv.QUOTE_ALL)

    # 3. Bộ 16 câu còn lại (Game 2 - Câu 25 đến 40)
    q16 = questions[24:]
    mc_16 = build_mc_rows(q16, with_prefix=False)
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Trac_nghiem_Lop_{grade}_16cau_Game2.txt"), mc_16, delimiter=",", quoting=csv.QUOTE_ALL)

    # 4. Dạng Cổ điển (Hỏi - Đáp 2 cột: Tile lật hiện câu hỏi, bấm Check hiện đáp án)
    classic_all = build_classic_rows(questions, with_prefix=False)
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Hoi_dap_Co_dien_Lop_{grade}_40cau.txt"), classic_all, delimiter=",", quoting=csv.QUOTE_ALL)
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Hoi_dap_Co_dien_Lop_{grade}_24cau_Game1.txt"), classic_all[:24], delimiter=",", quoting=csv.QUOTE_ALL)

    # 5. Dạng Cổ điển hiện sẵn 4 phương án A,B,C,D trong ô câu hỏi
    prompt_all = build_mc_prompt_rows(questions)
    write_csv_rows(os.path.join(out_folder, f"Baamboozle_Co_dien_Hien_4_dap_an_Lop_{grade}_24cau.txt"), prompt_all[:24], delimiter=",", quoting=csv.QUOTE_ALL)

def create_readme():
    content = """# HƯỚNG DẪN IMPORT CÂU HỎI ÔN TẬP VÀO BAAMBOOZLE

Hệ thống đã tự động trích xuất toàn bộ câu hỏi trắc nghiệm từ Đề cương ôn tập Tuần 9 môn Tin học (Lớp 3, 4, 5, 6, 7, 8) và chuẩn hóa theo đúng cú pháp Import của Baamboozle.

---

### 1. ĐỊNH DẠNG FILE
Mỗi khối lớp (từ Lớp 3 đến Lớp 8) đều có các file sẵn sàng:
1. `Baamboozle_Trac_nghiem_Lop_X_24cau_Game1.txt` (KHUYÊN DÙNG):
   - Chứa 24 câu hỏi trắc nghiệm (vừa chuẩn kích thước 1 game 24 ô của Baamboozle).
   - Định dạng: `"Câu hỏi","Đáp án ĐÚNG","Đáp án sai 1","Đáp án sai 2","Đáp án sai 3"`.
   - Baamboozle sẽ tự động xáo trộn ngẫu nhiên vị trí các đáp án khi học sinh chọn.
2. `Baamboozle_Trac_nghiem_Lop_X_16cau_Game2.txt`:
   - Chứa 16 câu còn lại (từ câu 25 đến 40) dùng để chơi tiếp hiệp 2.
3. `Baamboozle_Trac_nghiem_Lop_X_40cau_Comma.txt`:
   - Trọn bộ 40 câu hỏi trắc nghiệm một lần nhập.
4. `Baamboozle_Hoi_dap_Co_dien_Lop_X_24cau_Game1.txt`:
   - Dạng Flashcard lật ô kinh điển của Baamboozle (Cột 1: Câu hỏi, Cột 2: Câu trả lời ngắn).
5. `Baamboozle_Co_dien_Hien_4_dap_an_Lop_X_24cau.txt`:
   - Dạng lật ô hiển thị đầy đủ câu hỏi và [A... | B... | C... | D...] trên màn hình lớn.

---

### 2. CÁC BƯỚC NHẬP VÀO BAAMBOOZLE (Theo giao diện màn hình của Thầy/Cô)
1. Mở file `.txt` tương ứng của lớp cần ôn tập (ví dụ: `Baamboozle_Trac_nghiem_Lop_3_24cau_Game1.txt`).
2. Nhấn `Ctrl + A` để chọn tất cả, sau đó nhấn `Ctrl + C` để sao chép.
3. Trên trang web **Baamboozle**:
   - Nhấp vào ô văn bản lớn bên dưới dòng *"Copy and paste from ChatGPT, Quizlet Export..."*.
   - Nhấn `Ctrl + V` để dán nội dung vào.
   - Tại mục **Delimiter between question and answer**: Giữ nguyên chọn **Comma** (dấu phẩy).
   - Nhấn nút xanh **+ Import questions** ở góc dưới cùng bên trái.
4. Hoàn tất! Tất cả các câu hỏi cùng đáp án đúng và các phương án gây nhiễu sẽ được đưa vào game ngay lập tức!
"""
    with open(os.path.join(OUT_DIR, "HUONG_DAN_IMPORT_BAAMBOOZLE.md"), "w", encoding="utf-8") as f:
        f.write(content)

def main():
    print("=" * 60)
    print("BẮT ĐẦU XUẤT FILE IMPORT BAAMBOOZLE MÔN TIN HỌC TUẦN 9")
    print("=" * 60)
    
    for g in GRADES:
        doc_path = os.path.join(BASE_DIR, f"KHBD_Tin_học\\Lớp_{g}\\Tuần_09\\De_on_tap_DGDK1_Tin_hoc_Lop_{g}.docx")
        if not os.path.exists(doc_path):
            print(f"[CẢNH BÁO] Không tìm thấy file: {doc_path}")
            continue
            
        questions = parse_docx(doc_path)
        print(f"-> Lớp {g}: {len(questions)} câu hỏi")
        
        # Thư mục tập trung
        grade_out_dir = os.path.join(OUT_DIR, f"Lớp_{g}")
        generate_grade_files(g, questions, grade_out_dir)
        
        # Thư mục Tuần_09 của từng lớp
        tuan9_dir = os.path.join(BASE_DIR, f"KHBD_Tin_học\\Lớp_{g}\\Tuần_09")
        generate_grade_files(g, questions, tuan9_dir)
        
    create_readme()
    print("\n" + "=" * 60)
    print("ĐÃ XUẤT THÀNH CÔNG TẤT CẢ FILE CHO LỚP 3 ĐẾN LỚP 8!")
    print(f"Thư mục lưu trữ: {OUT_DIR}")
    print("=" * 60)

if __name__ == "__main__":
    main()
