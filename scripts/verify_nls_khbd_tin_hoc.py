# -*- coding: utf-8 -*-
"""
Verification script: Kiểm tra 100% kết quả cập nhật NLS cho 318 file KHBD Tin học
"""
import os, sys, re
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

def verify_all():
    tin_dir = r'D:\UNIGO\KHBD_Tin_học'
    khbd_files = []
    for root, dirs, files in os.walk(tin_dir):
        for f in files:
            if f.startswith('KHBD_') and f.endswith('.docx') and not f.startswith('~$'):
                khbd_files.append(os.path.join(root, f))
                
    total = len(khbd_files)
    print(f'=== BẮT ĐẦU KIỂM TRA {total} FILE KHBD TIN HỌC ===\n')
    
    cb_pattern = re.compile(r'^-\s*(\d\.\d)\.(CB[1-4])([a-z]):\s*(.+)\s*\(Đạt được thông qua (.+)\)$')
    
    valid_count = 0
    errors = []
    grade_counts = {}
    grade_expected_cb = {
        'Tiền_tiểu_học': 'CB1',
        'Lớp_1': 'CB1',
        'Lớp_2': 'CB1',
        'Lớp_3': 'CB1',
        'Lớp_4': 'CB2',
        'Lớp_5': 'CB2',
        'Lớp_6': 'CB3',
        'Lớp_7': 'CB3',
        'Lớp_8': 'CB4',
    }
    
    for f in khbd_files:
        rel = os.path.relpath(f, tin_dir)
        grade = rel.split(os.sep)[0]
        expected_cb = grade_expected_cb.get(grade)
        
        try:
            doc = Document(f)
        except Exception as e:
            errors.append(f'{rel}: Lỗi đọc file: {e}')
            continue
            
        nls_idx = None
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if '2.2' in t and 'Năng lực số' in t:
                nls_idx = i
                break
                
        if nls_idx is None:
            errors.append(f'{rel}: Không tìm thấy tiêu đề 2.2')
            continue
            
        next_idx = None
        for j in range(nls_idx + 1, len(doc.paragraphs)):
            t = doc.paragraphs[j].text.strip()
            if t.startswith('2.3') or t.startswith('2.4') or t.startswith('3.') or t.startswith('II.'):
                next_idx = j
                break
                
        if next_idx is None:
            errors.append(f'{rel}: Không tìm thấy phần kết thúc 2.3')
            continue
            
        bullets = [doc.paragraphs[k].text.strip() for k in range(nls_idx + 1, next_idx)]
        
        # Kiểm tra số lượng bullet
        if len(bullets) != 2:
            errors.append(f'{rel}: Có {len(bullets)} bullets thay vì 2')
            continue
            
        # Kiểm tra nội dung từng bullet
        file_valid = True
        for b_idx, b in enumerate(bullets):
            # Kiểm tra placeholder cũ
            if 'Sử dụng công cụ số để giải quyết vấn đề đơn giản trong bài học' in b:
                errors.append(f'{rel}: Vẫn còn placeholder cũ ở bullet {b_idx+1}')
                file_valid = False
                break
                
            m = cb_pattern.match(b)
            if not m:
                errors.append(f'{rel}: Bullet {b_idx+1} không khớp regex CB: {b[:60]}...')
                file_valid = False
                break
            else:
                code, cb, letter, content, act = m.groups()
                if cb != expected_cb:
                    errors.append(f'{rel}: Bậc {cb} không khớp với khối {grade} (kỳ vọng {expected_cb})')
                    file_valid = False
                    break
                    
        if file_valid:
            valid_count += 1
            grade_counts[grade] = grade_counts.get(grade, 0) + 1
            
    print('=== KẾT QUẢ KIỂM TRA NGHIỆM THU ===')
    print(f'Tổng số file kiểm tra: {total}')
    print(f'Hợp lệ 100% chuẩn CB: {valid_count}/{total} files')
    print(f'Số file lỗi: {len(errors)}')
    
    if errors:
        print('\nDanh sách lỗi:')
        for e in errors[:10]:
            print('  -', e)
    else:
        print('\n✅ TUYỆT VỜI! 100% file KHBD Tin học đều đạt chuẩn CB, đúng Bậc và không còn placeholder cũ!')
        print('\nThống kê file hợp lệ theo khối:')
        for g, c in sorted(grade_counts.items()):
            print(f'  {g}: {c} files (Chuẩn {grade_expected_cb[g]})')

if __name__ == '__main__':
    verify_all()
