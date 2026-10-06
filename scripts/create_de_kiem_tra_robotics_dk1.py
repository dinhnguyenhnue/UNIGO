# -*- coding: utf-8 -*-
"""
Tạo đề kiểm tra Robotics - Đánh giá định kỳ 1 cho tất cả các lớp 1-8.
Format UNIGO chuẩn, giữ nguyên header/footer từ template.

Cấu trúc đề:
- Thông tin đề (Trường, Lớp, Thời gian, Môn)
- Phần Thực hành trên máy: Lắp ghép bài kiểm tra số 1
- Tiêu chí đánh giá
- Chỗ trống cho GV chèn hình ảnh mô hình
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import copy
from docx import Document
from docx.shared import Pt, Cm, Inches, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ============================================================
# DỮ LIỆU ĐỀ KIỂM TRA THEO TỪNG KHỐI LỚP
# ============================================================

# Lớp 1-4: Bộ Kit OLLO (Bài 1-7 trước ĐK1)
# Lớp 5-8: Bộ Kit OLLO Excel 1 (Bài 1-6 trước ĐK1)

EXAM_DATA = {
    1: {
        "kit": "OLLO",
        "lop_label": "1A1",
        "to": "Tổ chuyên môn Tiểu học",
        "bai_da_hoc": [
            "Bài 1. Tập thể dục nào",
            "Bài 2. Chú cún dễ thương", 
            "Bài 3. Tăng cường sức khỏe",
            "Bài 4. Chú ốc sên chậm chạp",
            "Bài 5. Xe cảnh sát tuần tra",
            "Bài 6. Khám phá xe cứu hỏa",
            "Bài 7. Người giao hàng đa năng",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 7 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 7).",
            "Gọi tên đúng các linh kiện đã sử dụng trong mô hình.",
            "Vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "35 phút",
    },
    2: {
        "kit": "OLLO",
        "lop_label": "2A1",
        "to": "Tổ chuyên môn Tiểu học",
        "bai_da_hoc": [
            "Bài 1. Hãy nấu những món ăn ngon",
            "Bài 2. Hàm răng trắng sáng",
            "Bài 3. Sự phản xạ ánh sáng",
            "Bài 4. Đàn gà con",
            "Bài 5. Những bạn nhỏ lễ phép",
            "Bài 6. Động vật thân mềm",
            "Bài 7. Khu phố của chúng ta",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 7 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 7).",
            "Gọi tên đúng các linh kiện đã sử dụng trong mô hình.",
            "Vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "35 phút",
    },
    3: {
        "kit": "OLLO",
        "lop_label": "3A1",
        "to": "Tổ chuyên môn Tiểu học",
        "bai_da_hoc": [
            "Bài 1. Hãy nấu những món ăn ngon",
            "Bài 2. Hàm răng trắng sáng",
            "Bài 3. Sự phản xạ ánh sáng",
            "Bài 4. Đàn gà con",
            "Bài 5. Những bạn nhỏ lễ phép",
            "Bài 6. Động vật thân mềm",
            "Bài 7. Khu phố của chúng ta",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 7 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 7).",
            "Gọi tên đúng các linh kiện đã sử dụng trong mô hình.",
            "Vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "35 phút",
    },
    4: {
        "kit": "OLLO",
        "lop_label": "4C1",
        "to": "Tổ chuyên môn Tiểu học",
        "bai_da_hoc": [
            "Bài 1. Hãy nấu những món ăn ngon",
            "Bài 2. Hàm răng trắng sáng",
            "Bài 3. Sự phản xạ ánh sáng",
            "Bài 4. Đàn gà con",
            "Bài 5. Những bạn nhỏ lễ phép",
            "Bài 6. Động vật thân mềm",
            "Bài 7. Khu phố của chúng ta",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 7 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 7).",
            "Gọi tên đúng các linh kiện đã sử dụng trong mô hình.",
            "Vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "35 phút",
    },
    5: {
        "kit": "OLLO Excel 1",
        "lop_label": "5C1",
        "to": "Tổ chuyên môn Tiểu học",
        "bai_da_hoc": [
            "Bài 1. Động cơ là gì",
            "Bài 2. Robot cơ – Đỉnh vật",
            "Bài 3. Robot nhận biết âm thanh",
            "Bài 4. Bộ điều khiển từ xa",
            "Bài 5. Động cơ Dynamixel hoạt động thế nào",
            "Bài 6. Phương tiện giao thông qua các thời đại",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 6 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 6).",
            "Nhận biết và gọi tên đúng các loại linh kiện: khung, chốt nối, động cơ Dynamixel, bộ điều khiển CM-150.",
            "Kết nối đúng dây cáp và vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "45 phút",
    },
    6: {
        "kit": "OLLO Excel 1",
        "lop_label": "6A1",
        "to": "Tổ chuyên môn THCS",
        "bai_da_hoc": [
            "Bài 1. Động cơ là gì",
            "Bài 2. Robot cơ – Đỉnh vật",
            "Bài 3. Robot nhận biết âm thanh",
            "Bài 4. Bộ điều khiển từ xa",
            "Bài 5. Động cơ Dynamixel hoạt động thế nào",
            "Bài 6. Phương tiện giao thông qua các thời đại",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 6 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 6).",
            "Nhận biết và gọi tên đúng các loại linh kiện: khung, chốt nối, động cơ Dynamixel, bộ điều khiển CM-150.",
            "Kết nối đúng dây cáp và vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "45 phút",
    },
    7: {
        "kit": "OLLO Excel 1",
        "lop_label": "7A1",
        "to": "Tổ chuyên môn THCS",
        "bai_da_hoc": [
            "Bài 1. Động cơ là gì",
            "Bài 2. Robot cơ – Đỉnh vật",
            "Bài 3. Robot nhận biết âm thanh",
            "Bài 4. Bộ điều khiển từ xa",
            "Bài 5. Động cơ Dynamixel hoạt động thế nào",
            "Bài 6. Phương tiện giao thông qua các thời đại",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 6 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 6).",
            "Nhận biết và gọi tên đúng các loại linh kiện: khung, chốt nối, động cơ Dynamixel, bộ điều khiển CM-150.",
            "Kết nối đúng dây cáp và vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "45 phút",
    },
    8: {
        "kit": "OLLO Excel 1",
        "lop_label": "8A1",
        "to": "Tổ chuyên môn THCS",
        "bai_da_hoc": [
            "Bài 1. Động cơ là gì",
            "Bài 2. Robot cơ – Đỉnh vật",
            "Bài 3. Robot nhận biết âm thanh",
            "Bài 4. Bộ điều khiển từ xa",
            "Bài 5. Động cơ Dynamixel hoạt động thế nào",
            "Bài 6. Phương tiện giao thông qua các thời đại",
        ],
        "mo_hinh_kiem_tra": "Mô hình tự chọn (1 trong 6 bài đã học)",
        "yeu_cau_thuc_hanh": [
            "Lắp ghép hoàn chỉnh 1 mô hình robot tự chọn từ các bài đã học (Bài 1 – Bài 6).",
            "Nhận biết và gọi tên đúng các loại linh kiện: khung, chốt nối, động cơ Dynamixel, bộ điều khiển CM-150.",
            "Kết nối đúng dây cáp và vận hành mô hình robot hoạt động đúng chức năng.",
            "Thu dọn linh kiện gọn gàng, đúng vị trí sau khi hoàn thành.",
        ],
        "thoi_gian": "45 phút",
    },
}

# Tiêu chí chấm điểm chung
TIEU_CHI_CHAM_DIEM_TH = [
    ("Lắp ghép đúng mô hình, chắc chắn, không rơi rời", "4 điểm"),
    ("Gọi tên đúng linh kiện đã sử dụng", "2 điểm"),
    ("Mô hình hoạt động đúng chức năng", "2 điểm"),
    ("Thu dọn gọn gàng, đúng vị trí", "1 điểm"),
    ("Thái độ nghiêm túc, hợp tác tốt trong nhóm", "1 điểm"),
]

TIEU_CHI_CHAM_DIEM_THCS = [
    ("Lắp ghép đúng mô hình, chắc chắn, không rơi rời", "3 điểm"),
    ("Nhận biết và gọi tên đúng các loại linh kiện", "2 điểm"),
    ("Kết nối đúng dây cáp, mô hình hoạt động đúng chức năng", "2 điểm"),
    ("Giải thích được nguyên lý hoạt động cơ bản", "1 điểm"),
    ("Thu dọn gọn gàng, đúng vị trí", "1 điểm"),
    ("Thái độ nghiêm túc, hợp tác tốt trong nhóm", "1 điểm"),
]


# ============================================================
# HÀM TIỆN ÍCH
# ============================================================

def set_font(run, name="Times New Roman", size=13, bold=False, italic=False, color=None):
    """Thiết lập font cho run."""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = color


def add_paragraph_styled(doc, text, size=13, bold=False, italic=False, alignment=None, space_before=0, space_after=0):
    """Thêm paragraph với style."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    if alignment is not None:
        p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    return p


def set_table_borders(table, border_type="single", size="4", color="000000"):
    """Thiết lập viền cho bảng."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders_xml = f'''<w:tblBorders {nsdecls("w")}>
        <w:top w:val="{border_type}" w:sz="{size}" w:space="0" w:color="{color}"/>
        <w:left w:val="{border_type}" w:sz="{size}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{border_type}" w:sz="{size}" w:space="0" w:color="{color}"/>
        <w:right w:val="{border_type}" w:sz="{size}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{border_type}" w:sz="{size}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{border_type}" w:sz="{size}" w:space="0" w:color="{color}"/>
    </w:tblBorders>'''
    tblPr.append(parse_xml(borders_xml))
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def set_no_borders(table):
    """Xóa viền bảng."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders_xml = f'''<w:tblBorders {nsdecls("w")}>
        <w:top w:val="nil"/>
        <w:left w:val="nil"/>
        <w:bottom w:val="nil"/>
        <w:right w:val="nil"/>
        <w:insideH w:val="nil"/>
        <w:insideV w:val="nil"/>
    </w:tblBorders>'''
    tblPr.append(parse_xml(borders_xml))
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def set_cell_text(cell, text, size=13, bold=False, italic=False, alignment=None):
    """Thiết lập text cho cell."""
    # Clear existing paragraphs
    for p in cell.paragraphs:
        p.clear()
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    if alignment is not None:
        p.alignment = alignment


def clear_body_keep_sectpr(doc):
    """Xóa toàn bộ body nhưng giữ lại sectPr (chứa header/footer refs)."""
    body = doc.element.body
    children_to_remove = []
    for child in body:
        if not child.tag.endswith('sectPr'):
            children_to_remove.append(child)
    for child in children_to_remove:
        body.remove(child)


# ============================================================
# HÀM TẠO ĐỀ KIỂM TRA
# ============================================================

def create_exam(grade, data, template_path, output_path):
    """Tạo đề kiểm tra Robotics ĐK1 cho 1 lớp."""
    
    # Mở template để giữ header/footer
    doc = Document(template_path)
    
    # Xóa body cũ, giữ sectPr
    clear_body_keep_sectpr(doc)
    
    is_thcs = grade >= 6
    tieu_chi = TIEU_CHI_CHAM_DIEM_THCS if is_thcs else TIEU_CHI_CHAM_DIEM_TH
    
    # ============================================
    # BẢNG THÔNG TIN ĐẦU BÀI (Table 0 - NO BORDER)
    # ============================================
    info_table = doc.add_table(rows=3, cols=2)
    set_no_borders(info_table)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Row 0
    set_cell_text(info_table.cell(0, 0), "Trường: Tiểu học và THCS UNIGO", bold=True)
    set_cell_text(info_table.cell(0, 1), "Ngày kiểm tra: ....../....../2026", bold=True)
    # Row 1
    set_cell_text(info_table.cell(1, 0), "GV: Đậu Đình Nguyên")
    set_cell_text(info_table.cell(1, 1), f"Lớp: {data['lop_label']}")
    # Row 2
    set_cell_text(info_table.cell(2, 0), f"Tổ: {data['to']}")
    set_cell_text(info_table.cell(2, 1), f"Thời gian: {data['thoi_gian']}")
    
    # ============================================
    # TIÊU ĐỀ ĐỀ KIỂM TRA
    # ============================================
    doc.add_paragraph()  # Dòng trống
    
    if is_thcs:
        # THCS format
        add_paragraph_styled(doc, "ĐỀ KIỂM TRA ĐÁNH GIÁ ĐỊNH KỲ 1", 
                           size=14, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_paragraph_styled(doc, "Môn học: Robotics", 
                           size=13, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_paragraph_styled(doc, f"Thời lượng: {data['thoi_gian']}", 
                           size=13, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_paragraph_styled(doc, f"Bộ thiết bị: {data['kit']}", 
                           size=13, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        # Tiểu học format
        add_paragraph_styled(doc, "ĐỀ KIỂM TRA ĐÁNH GIÁ ĐỊNH KỲ 1", 
                           size=14, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_paragraph_styled(doc, "MÔN: ROBOTICS", 
                           size=13, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_paragraph_styled(doc, f"BỘ THIẾT BỊ: {data['kit'].upper()}", 
                           size=13, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_paragraph_styled(doc, f"Thời gian: {data['thoi_gian']}", 
                           size=13, bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    
    doc.add_paragraph()  # Dòng trống
    
    # ============================================
    # PHẦN HƯỚNG DẪN
    # ============================================
    add_paragraph_styled(doc, "HƯỚNG DẪN:", size=13, bold=True, space_before=6)
    add_paragraph_styled(doc, "- Hình thức kiểm tra: Thực hành trên máy (lắp ghép robot bằng bộ Kit Robotics).", 
                        size=13, space_before=2)
    add_paragraph_styled(doc, "- Học sinh làm việc theo nhóm 2 bạn, mỗi nhóm tự chọn 1 mô hình từ các bài đã học.", 
                        size=13, space_before=2)
    add_paragraph_styled(doc, f"- Thời gian thực hành: {data['thoi_gian']} (bao gồm cả thu dọn).", 
                        size=13, space_before=2)
    
    doc.add_paragraph()  # Dòng trống
    
    # ============================================
    # NỘI DUNG KIỂM TRA
    # ============================================
    add_paragraph_styled(doc, "NỘI DUNG KIỂM TRA: THỰC HÀNH LẮP GHÉP BÀI KIỂM TRA SỐ 1", 
                        size=13, bold=True, space_before=6)
    
    doc.add_paragraph()
    
    # Phạm vi kiến thức
    add_paragraph_styled(doc, "I. PHẠM VI KIẾN THỨC (Các bài đã học):", size=13, bold=True, space_before=6)
    for bai in data["bai_da_hoc"]:
        add_paragraph_styled(doc, f"    - {bai}", size=13, space_before=2)
    
    doc.add_paragraph()
    
    # Yêu cầu thực hành
    add_paragraph_styled(doc, "II. YÊU CẦU THỰC HÀNH:", size=13, bold=True, space_before=6)
    add_paragraph_styled(doc, f"Mô hình kiểm tra: {data['mo_hinh_kiem_tra']}", 
                        size=13, italic=True, space_before=4)
    doc.add_paragraph()
    for i, yc in enumerate(data["yeu_cau_thuc_hanh"], 1):
        add_paragraph_styled(doc, f"    {i}. {yc}", size=13, space_before=2)
    
    doc.add_paragraph()
    
    # ============================================
    # HÌNH ẢNH MÔ HÌNH (Chỗ trống để chèn)
    # ============================================
    add_paragraph_styled(doc, "III. HÌNH ẢNH MÔ HÌNH THAM KHẢO:", size=13, bold=True, space_before=6)
    add_paragraph_styled(doc, "(GV chèn hình ảnh các mô hình robot đã học vào đây)", 
                        size=13, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=4)
    
    # Tạo bảng placeholder cho hình ảnh
    if grade <= 4:
        # Lớp 1-4: 7 bài → bảng 2 hàng x 4 cột (ô cuối để trống)
        img_table = doc.add_table(rows=4, cols=4)
        set_table_borders(img_table)
        img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        bai_names_short = [f"Bài {i+1}" for i in range(7)]
        bai_names_short.append("")  # ô trống
        
        for idx in range(8):
            r = (idx // 4) * 2
            c = idx % 4
            # Hàng tiêu đề bài
            set_cell_text(img_table.cell(r, c), bai_names_short[idx] if idx < len(bai_names_short) else "", 
                         bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            # Hàng hình ảnh
            if idx < 7:
                p = img_table.cell(r+1, c).paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run("\n\n(Chèn hình)\n\n")
                set_font(run, size=10, italic=True, color=RGBColor(128, 128, 128))
    else:
        # Lớp 5-8: 6 bài → bảng 2 hàng x 3 cột
        img_table = doc.add_table(rows=4, cols=3)
        set_table_borders(img_table)
        img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        bai_names_short = [f"Bài {i+1}" for i in range(6)]
        
        for idx in range(6):
            r = (idx // 3) * 2
            c = idx % 3
            # Hàng tiêu đề bài
            set_cell_text(img_table.cell(r, c), bai_names_short[idx], 
                         bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            # Hàng hình ảnh
            p = img_table.cell(r+1, c).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("\n\n(Chèn hình)\n\n")
            set_font(run, size=10, italic=True, color=RGBColor(128, 128, 128))
    
    doc.add_paragraph()
    
    # ============================================
    # TIÊU CHÍ ĐÁNH GIÁ (Bảng có viền)
    # ============================================
    add_paragraph_styled(doc, "IV. TIÊU CHÍ ĐÁNH GIÁ:", size=13, bold=True, space_before=6)
    doc.add_paragraph()
    
    tc_table = doc.add_table(rows=len(tieu_chi) + 2, cols=3)
    set_table_borders(tc_table)
    tc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    headers = ["STT", "Tiêu chí", "Điểm"]
    for i, h in enumerate(headers):
        set_cell_text(tc_table.cell(0, i), h, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Data rows
    for i, (tc, diem) in enumerate(tieu_chi, 1):
        set_cell_text(tc_table.cell(i, 0), str(i), alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(tc_table.cell(i, 1), tc)
        set_cell_text(tc_table.cell(i, 2), diem, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Tổng điểm
    set_cell_text(tc_table.cell(len(tieu_chi) + 1, 0), "", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(tc_table.cell(len(tieu_chi) + 1, 1), "TỔNG CỘNG", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(tc_table.cell(len(tieu_chi) + 1, 2), "10 điểm", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Set column widths
    for row in tc_table.rows:
        row.cells[0].width = Cm(1.5)
        row.cells[1].width = Cm(12)
        row.cells[2].width = Cm(3)
    
    doc.add_paragraph()
    
    # ============================================
    # PHẦN GHI CHÚ CHO GIÁO VIÊN
    # ============================================
    add_paragraph_styled(doc, "V. GHI CHÚ CHO GIÁO VIÊN:", size=13, bold=True, space_before=6)
    add_paragraph_styled(doc, "- Quan sát quá trình lắp ghép của từng nhóm, ghi nhận xét vào phiếu chấm điểm.", 
                        size=13, space_before=2)
    add_paragraph_styled(doc, "- Đánh giá cả quá trình thực hành (không chỉ sản phẩm cuối cùng).", 
                        size=13, space_before=2)
    add_paragraph_styled(doc, "- Nhắc nhở HS tuân thủ quy tắc an toàn khi sử dụng linh kiện điện tử.", 
                        size=13, space_before=2)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # ============================================
    # DÒNG KẾT THÚC
    # ============================================
    add_paragraph_styled(doc, "-------------------------------- Hết --------------------------------", 
                        size=13, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=6)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # ============================================
    # BẢNG CHỮ KÝ (Table cuối - NO BORDER)
    # ============================================
    sign_table = doc.add_table(rows=3, cols=3)
    set_no_borders(sign_table)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Row 0
    set_cell_text(sign_table.cell(0, 0), "DUYỆT CỦA BGH", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(sign_table.cell(0, 1), "DUYỆT CỦA TỔ CM", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(sign_table.cell(0, 2), "NGƯỜI RA ĐỀ", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Row 1
    set_cell_text(sign_table.cell(1, 0), "(Ký, ghi rõ họ tên)", italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(sign_table.cell(1, 1), "(Ký, ghi rõ họ tên)", italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(sign_table.cell(1, 2), "(Ký, ghi rõ họ tên)", italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Row 2
    set_cell_text(sign_table.cell(2, 0), "\n\n\n")
    set_cell_text(sign_table.cell(2, 1), "\n\n\n")
    # Người soạn: In đậm tên
    p = sign_table.cell(2, 2).paragraphs[0]
    p.clear()
    run = p.add_run("\n\n\nĐậu Đình Nguyên")
    set_font(run, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # ============================================
    # LƯU FILE
    # ============================================
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    try:
        doc.save(output_path)
        print(f"  ✅ Đã tạo: {output_path}")
    except PermissionError:
        alt_path = output_path.replace(".docx", "_new.docx")
        doc.save(alt_path)
        print(f"  ⚠️  File đang mở, lưu tại: {alt_path}")


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 70)
    print("TẠO ĐỀ KIỂM TRA ROBOTICS - ĐÁNH GIÁ ĐỊNH KỲ 1")
    print("=" * 70)
    
    results = []
    
    for grade in range(1, 9):
        data = EXAM_DATA[grade]
        
        # Tìm template file (file KHBD Đánh giá ĐK1 hiện có)
        khbd_dir = rf"d:\UNIGO\KHBD_Robotics\Lớp_{grade}"
        
        # Tìm file đánh giá ĐK1 làm template
        template_path = None
        for root, dirs, files in os.walk(khbd_dir):
            for f in files:
                if "anh_gia_inh_ky_1" in f and f.endswith(".docx"):
                    template_path = os.path.join(root, f)
                    break
            if template_path:
                break
        
        if not template_path:
            # Fallback: dùng bất kỳ file KHBD nào
            for root, dirs, files in os.walk(khbd_dir):
                for f in files:
                    if f.endswith(".docx"):
                        template_path = os.path.join(root, f)
                        break
                if template_path:
                    break
        
        if not template_path:
            print(f"  ❌ Lớp {grade}: Không tìm thấy file template!")
            continue
        
        print(f"\n📝 Lớp {grade} (Kit: {data['kit']}):")
        print(f"   Template: {os.path.basename(template_path)}")
        
        # Xác định thư mục output (cùng thư mục với file KHBD ĐK1)
        # Tìm thư mục Tuần chứa file ĐK1
        tuan_dir = os.path.dirname(template_path)
        output_path = os.path.join(tuan_dir, f"De_kiem_tra_Robotics_DK1_Lop_{grade}.docx")
        
        create_exam(grade, data, template_path, output_path)
        results.append((grade, output_path))
    
    print("\n" + "=" * 70)
    print(f"HOÀN TẤT: Đã tạo {len(results)} đề kiểm tra")
    print("=" * 70)
    for grade, path in results:
        print(f"  Lớp {grade}: {path}")


if __name__ == "__main__":
    main()
