#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
check_cuoc_thi.py - Thống kê cuộc thi học sinh UNIGO
=====================================================
Tự động quét thư mục Check_các_cuộc_thi, đọc kết quả thi từ các lần,
so khớp với danh sách HS toàn trường, xuất file Excel thống kê tổng thể
và xuất file thống kê chi tiết cho từng lần thi trực tiếp trong thư mục tương ứng.

Usage:
    python scripts/check_cuoc_thi.py
"""

import os
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
    numbers
)
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIG
# ============================================================
BASE_DIR = Path(r"D:\UNIGO\Check_các_cuộc_thi")
DS_HS_FILE = BASE_DIR / "DANH SÁCH HỌC SINH NĂM HỌC 2026 - 2027.xlsx"
OUTPUT_FILE = BASE_DIR / "Thống kê cuộc thi.xlsx"

# Folders / files to ignore (not contests)
IGNORE_FOLDERS = {"Thống kê cuộc thi.xlsx"}
IGNORE_PREFIXES = ("~$",)

# ============================================================
# STYLES
# ============================================================
THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)
HEADER_FONT = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SUBHEADER_FONT = Font(name="Times New Roman", size=11, bold=True, color="1F4E79")
DATA_FONT = Font(name="Times New Roman", size=11)
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GREEN_FONT = Font(name="Times New Roman", size=11, color="006100", bold=True)
RED_FONT = Font(name="Times New Roman", size=11, color="9C0006", bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
TITLE_FONT = Font(name="Times New Roman", size=14, bold=True, color="2F5496")
CLASS_HEADER_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
CLASS_HEADER_FONT = Font(name="Times New Roman", size=12, bold=True, color="BF8F00")

# Fills for leaderboards
GOLD_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
SILVER_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
BRONZE_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")


# ============================================================
# HELPERS
# ============================================================

def normalize_name(name: str) -> str:
    """Normalize Vietnamese name for matching: lowercase, strip, collapse spaces."""
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r"\s+", " ", name)
    name = name.lower()
    return name


def strip_accents(s: str) -> str:
    """Remove Vietnamese diacritics for fallback matching."""
    if not s:
        return ""
    s = unicodedata.normalize("NFD", str(s))
    s = re.sub(r"[\u0300-\u036f]", "", s)
    s = s.replace("đ", "d").replace("Đ", "d")
    return re.sub(r"\s+", " ", s).lower().strip()


def normalize_class(cls: str) -> str:
    """Normalize class name: uppercase, strip."""
    if not cls:
        return ""
    return str(cls).strip().upper()


def extract_lan_number(folder_name: str) -> int:
    """Extract round number from folder name like 'Lần 1', 'Lần 2'."""
    m = re.search(r"[Ll]ần\s*(\d+)", folder_name)
    if m:
        return int(m.group(1))
    m = re.search(r"(\d+)", folder_name)
    if m:
        return int(m.group(1))
    return 0


def find_header_columns(ws):
    """
    Auto-detect column indices for name and class in a contest result sheet.
    Returns (name_col_idx, class_col_idx, extra_cols) — 0-based indices.
    """
    name_col = None
    class_col = None
    extra_cols = {}

    for row in ws.iter_rows(min_row=1, max_row=5, values_only=False):
        for cell in row:
            val = str(cell.value or "").strip().lower()
            if any(kw in val for kw in ["họ và tên", "họ tên", "ho va ten", "ho ten"]):
                name_col = cell.column - 1  # 0-based
            elif val in ["lớp", "lop"]:
                class_col = cell.column - 1
            elif val in ["khối", "khoi"]:
                if class_col is None:
                    extra_cols["khối"] = cell.column - 1
            # Extra columns
            if any(kw in val for kw in ["tổng điểm", "điểm thi", "diem thi"]):
                extra_cols["điểm"] = cell.column - 1
            if any(kw in val for kw in ["vòng", "vong"]):
                extra_cols["vòng"] = cell.column - 1
            if any(kw in val for kw in ["thời gian", "thoi gian"]):
                extra_cols["thời_gian"] = cell.column - 1
            if any(kw in val for kw in ["id", "mã tài khoản"]):
                extra_cols["id"] = cell.column - 1

        if name_col is not None:
            break

    return name_col, class_col, extra_cols


# ============================================================
# STEP 1: Read master student list
# ============================================================

def read_master_students(filepath: Path) -> dict:
    """
    Read master student list.
    Returns: {
        'all_students': [(họ_tên, lớp), ...],
        'by_class': {'1A1': [họ_tên, ...], ...},
        'classes_order': ['1A1', '1C1', ...]
    }
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_students = []
    by_class = {}
    classes_order = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        class_name = normalize_class(sheet_name)
        classes_order.append(class_name)
        by_class[class_name] = []

        # Dynamically detect name column and header row
        name_col_idx = 2
        header_row_idx = 2
        for r_idx, row_vals in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), start=1):
            for c_i, v in enumerate(row_vals):
                val_str = str(v or "").strip().lower()
                if any(kw in val_str for kw in ["họ tên", "họ và tên", "ho ten"]):
                    name_col_idx = c_i
                    header_row_idx = r_idx
                    break

        for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row, values_only=True):
            stt = row[0]
            if stt is None:
                continue
            stt_str = str(stt).replace(".0", "").strip()
            if not stt_str.isdigit():
                continue

            if name_col_idx >= len(row):
                continue
            ho_ten = str(row[name_col_idx] or "").strip()
            if not ho_ten:
                continue

            all_students.append((ho_ten, class_name))
            by_class[class_name].append(ho_ten)

    wb.close()
    return {
        "all_students": all_students,
        "by_class": by_class,
        "classes_order": classes_order,
    }


# ============================================================
# STEP 2: Scan contest directories
# ============================================================

def scan_contests(base_dir: Path) -> dict:
    """
    Scan contest directories.
    Returns: {
        'IOE': {
            1: [Path_to_excel, ...],  # Lần 1
            2: [Path_to_excel, ...],  # Lần 2
        },
        ...
    }
    """
    contests = {}

    for item in sorted(base_dir.iterdir()):
        if not item.is_dir():
            continue
        if item.name.startswith(tuple(IGNORE_PREFIXES)):
            continue

        contest_name = item.name
        contests[contest_name] = {}

        for sub in sorted(item.iterdir()):
            if not sub.is_dir():
                continue
            lan_num = extract_lan_number(sub.name)
            if lan_num == 0:
                continue

            excel_files = []
            for f in sub.iterdir():
                # Avoid reading the generated summary files as contest source
                if f.name.startswith("Thống kê") or f.name.startswith("Danh sách học sinh tham gia"):
                    continue
                if f.suffix.lower() in (".xlsx", ".xls") and not f.name.startswith("~$"):
                    excel_files.append(f)

            if excel_files:
                contests[contest_name][lan_num] = excel_files

    return contests


# ============================================================
# STEP 3: Read contest results
# ============================================================

def read_contest_results(excel_files: list) -> list:
    """
    Read contest result files.
    Returns: [(normalized_name, class_name, {extra_info}), ...]
    """
    results = []

    for fpath in excel_files:
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
        except Exception as e:
            print(f"  ⚠ Không đọc được file: {fpath.name} — {e}")
            continue

        for ws in wb.worksheets:
            name_col, class_col, extra_cols = find_header_columns(ws)

            if name_col is None:
                print(f"  ⚠ Không tìm thấy cột tên HS trong sheet '{ws.title}' của file {fpath.name}")
                continue

            # Find header row
            header_row = 1
            for r in range(1, 6):
                cell_val = str(ws.cell(row=r, column=name_col + 1).value or "").strip().lower()
                if any(kw in cell_val for kw in ["họ và tên", "họ tên"]):
                    header_row = r
                    break

            for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, values_only=True):
                if name_col >= len(row):
                    continue
                raw_name = row[name_col]
                if not raw_name or str(raw_name).strip() == "":
                    continue

                ho_ten = str(raw_name).strip()
                norm_name = normalize_name(ho_ten)

                # Get class
                cls = ""
                if class_col is not None and class_col < len(row):
                    cls = normalize_class(str(row[class_col] or ""))

                # Get extra info
                info = {"tên_gốc": ho_ten}
                for key, idx in extra_cols.items():
                    if idx < len(row) and row[idx] is not None:
                        info[key] = row[idx]

                results.append((norm_name, cls, info))

        wb.close()

    return results


# ============================================================
# STEP 4: Match students (Name uniqueness & class-first matching)
# ============================================================

def match_students(master_students: list, contest_results: list) -> dict:
    """
    Match contest results to master student list accurately.
    Returns: {(họ_tên, lớp): {extra_info} or None}
    """
    # Count frequency of normalized name and unaccented name across all master students
    name_counts = Counter(normalize_name(h) for h, _ in master_students)
    no_acc_counts = Counter(strip_accents(h) for h, _ in master_students)

    # Build lookup from contest results
    contest_lookup = {}
    contest_lookup_no_acc = {}
    for norm_name, cls, info in contest_results:
        if norm_name not in contest_lookup:
            contest_lookup[norm_name] = []
        contest_lookup[norm_name].append((cls, info))

        no_acc = strip_accents(norm_name)
        if no_acc not in contest_lookup_no_acc:
            contest_lookup_no_acc[no_acc] = []
        contest_lookup_no_acc[no_acc].append((cls, info))

    def get_score_key(item):
        info = item[1]
        try:
            v = int(info.get("vòng", 0) or 0)
        except (ValueError, TypeError):
            v = 0
        try:
            d = int(info.get("điểm", 0) or 0)
        except (ValueError, TypeError):
            d = 0
        return (v, d)

    # For students with multiple attempts (e.g. 2 accounts), sort candidates by (vong desc, diem desc)
    for n in contest_lookup:
        contest_lookup[n].sort(key=get_score_key, reverse=True)
    for n in contest_lookup_no_acc:
        contest_lookup_no_acc[n].sort(key=get_score_key, reverse=True)

    matched = {}
    used_records = set()

    for ho_ten, lop in master_students:
        norm = normalize_name(ho_ten)
        no_acc = strip_accents(ho_ten)
        key = (ho_ten, lop)
        found = None

        if norm in contest_lookup:
            candidates = contest_lookup[norm]
            # 1. Match by exact class first
            for cls, info in candidates:
                if cls == lop and id(info) not in used_records:
                    found = info
                    used_records.add(id(info))
                    break

            # 2. Fallback only if name is completely unique across the whole school
            if found is None and name_counts[norm] == 1:
                for cls, info in candidates:
                    if id(info) not in used_records:
                        found = info
                        used_records.add(id(info))
                        break

        # 3. Unaccented matching fallback (for names typed without Vietnamese diacritics in contest system)
        if found is None and no_acc in contest_lookup_no_acc:
            candidates_no_acc = contest_lookup_no_acc[no_acc]
            # 3a. Same class
            for cls, info in candidates_no_acc:
                if cls == lop and id(info) not in used_records:
                    found = info
                    used_records.add(id(info))
                    break

            # 3b. Fallback only if unaccented name is completely unique across the whole school
            if found is None and no_acc_counts[no_acc] == 1:
                for cls, info in candidates_no_acc:
                    if id(info) not in used_records:
                        found = info
                        used_records.add(id(info))
                        break

        matched[key] = found

    return matched


# ============================================================
# STEP 5: Generate per-round Excel report (Saved in round folder)
# ============================================================

def apply_cell_style(cell, font=None, fill=None, alignment=None, border=None):
    """Apply styles to a cell."""
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border


def auto_fit_columns(ws, min_width=8, max_width=40):
    """Auto-fit column widths."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                cell_len = len(str(cell.value))
                max_len = max(max_len, cell_len)
        adjusted = min(max(max_len + 3, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted


def generate_round_report(master_data: dict, contest_name: str, lan_num: int, matches: dict, round_dir: Path) -> Path:
    """
    Generate an individual detailed Excel report for a specific round of a contest,
    saved directly inside that round's folder.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    all_students = master_data["all_students"]
    classes_order = master_data["classes_order"]
    by_class = master_data["by_class"]
    total_school = len(all_students)

    da_thi_count = sum(1 for v in matches.values() if v is not None)
    chua_thi_count = total_school - da_thi_count
    pct_da_thi = (da_thi_count * 100 // total_school) if total_school > 0 else 0

    # ----------------------------------------------------
    # SHEET 1: TỔNG HỢP TOÀN TRƯỜNG
    # ----------------------------------------------------
    ws_all = wb.create_sheet("Tổng hợp toàn trường")

    # Title
    ws_all.merge_cells("A1:H1")
    t_cell = ws_all.cell(row=1, column=1, value=f"THỐNG KÊ KẾT QUẢ {contest_name.upper()} - LẦN {lan_num}")
    apply_cell_style(t_cell, font=TITLE_FONT, alignment=CENTER)
    ws_all.row_dimensions[1].height = 28

    # Subtitle
    ws_all.merge_cells("A2:H2")
    sub_cell = ws_all.cell(
        row=2, column=1,
        value=f"Năm học 2026 - 2027 | Sĩ số: {total_school} HS | Đã thi: {da_thi_count} HS ({pct_da_thi}%) | Chưa thi: {chua_thi_count} HS ({100 - pct_da_thi}%)"
    )
    apply_cell_style(sub_cell, font=Font(name="Times New Roman", size=11, italic=True, color="595959"), alignment=CENTER)
    ws_all.row_dimensions[2].height = 20

    # Header row
    headers = ["STT", "Họ và tên", "Lớp", "Trạng thái", "Vòng tự luyện", "Tổng điểm", "Thời gian (giây)", "Ghi chú"]
    for col_idx, h in enumerate(headers, 1):
        c = ws_all.cell(row=4, column=col_idx, value=h)
        apply_cell_style(c, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=THIN_BORDER)
    ws_all.row_dimensions[4].height = 28

    cur_row = 5
    stt_all = 0
    for cls in classes_order:
        # Class header row
        ws_all.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=len(headers))
        cls_cell = ws_all.cell(row=cur_row, column=1, value=f"Lớp {cls}")
        apply_cell_style(cls_cell, font=CLASS_HEADER_FONT, fill=CLASS_HEADER_FILL, alignment=LEFT, border=THIN_BORDER)
        for c in range(2, len(headers) + 1):
            apply_cell_style(ws_all.cell(row=cur_row, column=c), border=THIN_BORDER)
        ws_all.row_dimensions[cur_row].height = 22
        cur_row += 1

        for ho_ten in by_class[cls]:
            stt_all += 1
            key = (ho_ten, cls)
            info = matches.get(key)

            ws_all.cell(row=cur_row, column=1, value=stt_all)
            apply_cell_style(ws_all.cell(row=cur_row, column=1), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

            ws_all.cell(row=cur_row, column=2, value=ho_ten)
            apply_cell_style(ws_all.cell(row=cur_row, column=2), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)

            ws_all.cell(row=cur_row, column=3, value=cls)
            apply_cell_style(ws_all.cell(row=cur_row, column=3), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

            status_cell = ws_all.cell(row=cur_row, column=4)
            if info is not None:
                status_cell.value = "✅ Đã thi"
                apply_cell_style(status_cell, font=GREEN_FONT, fill=GREEN_FILL, alignment=CENTER, border=THIN_BORDER)
                ws_all.cell(row=cur_row, column=5, value=info.get("vòng", ""))
                ws_all.cell(row=cur_row, column=6, value=info.get("điểm", ""))
                ws_all.cell(row=cur_row, column=7, value=info.get("thời_gian", ""))
            else:
                status_cell.value = "❌ Chưa thi"
                apply_cell_style(status_cell, font=RED_FONT, fill=RED_FILL, alignment=CENTER, border=THIN_BORDER)
                ws_all.cell(row=cur_row, column=5, value="")
                ws_all.cell(row=cur_row, column=6, value="")
                ws_all.cell(row=cur_row, column=7, value="")

            for c in range(5, 9):
                apply_cell_style(ws_all.cell(row=cur_row, column=c), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

            ws_all.row_dimensions[cur_row].height = 20
            cur_row += 1

    # Total row
    cur_row += 1
    ws_all.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=3)
    tot_c = ws_all.cell(row=cur_row, column=1, value="TỔNG TOÀN TRƯỜNG:")
    apply_cell_style(tot_c, font=Font(name="Times New Roman", size=11, bold=True), alignment=LEFT)
    stat_c = ws_all.cell(row=cur_row, column=4, value=f"{da_thi_count}/{total_school} ({pct_da_thi}%)")
    apply_cell_style(stat_c, font=Font(name="Times New Roman", size=11, bold=True, color="006100"), alignment=CENTER)

    auto_fit_columns(ws_all)
    ws_all.column_dimensions["A"].width = 6
    ws_all.column_dimensions["B"].width = 28
    ws_all.freeze_panes = "D5"

    # ----------------------------------------------------
    # SHEET 2: BẢNG XẾP HẠNG (ĐÃ THI)
    # ----------------------------------------------------
    ws_rank = wb.create_sheet("Bảng xếp hạng (Đã thi)")

    # Collect and sort completed students
    completed_list = []
    for (ho_ten, cls), info in matches.items():
        if info is not None:
            try: v = int(info.get("vòng", 0) or 0)
            except: v = 0
            try: d = int(info.get("điểm", 0) or 0)
            except: d = 0
            try: t = int(info.get("thời_gian", 0) or 0)
            except: t = 0
            completed_list.append({
                "ho_ten": ho_ten,
                "lop": cls,
                "vong": v,
                "diem": d,
                "thoi_gian": t,
                "id": info.get("id", "")
            })

    # Sort: Vòng desc, Điểm desc, Thời gian asc
    completed_list.sort(key=lambda x: (-x["vong"], -x["diem"], x["thoi_gian"]))

    ws_rank.merge_cells("A1:G1")
    tr_cell = ws_rank.cell(row=1, column=1, value=f"BẢNG XẾP HẠNG {contest_name.upper()} - LẦN {lan_num}")
    apply_cell_style(tr_cell, font=TITLE_FONT, alignment=CENTER)
    ws_rank.row_dimensions[1].height = 28

    ws_rank.merge_cells("A2:G2")
    subr_cell = ws_rank.cell(row=2, column=1, value=f"Tổng số học sinh đã hoàn thành: {len(completed_list)} học sinh")
    apply_cell_style(subr_cell, font=Font(name="Times New Roman", size=11, italic=True, color="595959"), alignment=CENTER)
    ws_rank.row_dimensions[2].height = 20

    rank_headers = ["Xếp hạng", "Họ và tên", "Lớp", "Vòng tự luyện", "Tổng điểm", "Thời gian (giây)", "Danh hiệu / Khen thưởng"]
    for col_idx, h in enumerate(rank_headers, 1):
        c = ws_rank.cell(row=4, column=col_idx, value=h)
        apply_cell_style(c, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=THIN_BORDER)
    ws_rank.row_dimensions[4].height = 28

    for idx, item in enumerate(completed_list, 1):
        r_idx = 4 + idx
        fill_to_use = None
        award = ""
        if idx == 1:
            fill_to_use = GOLD_FILL
            award = "🥇 Thủ khoa toàn trường"
        elif idx in (2, 3):
            fill_to_use = SILVER_FILL
            award = f"🥈 Top {idx} toàn trường"
        elif idx in (4, 5):
            fill_to_use = BRONZE_FILL
            award = f"🥉 Top {idx} toàn trường"

        c1 = ws_rank.cell(row=r_idx, column=1, value=idx)
        apply_cell_style(c1, font=Font(name="Times New Roman", size=11, bold=(idx <= 3)), fill=fill_to_use, alignment=CENTER, border=THIN_BORDER)

        c2 = ws_rank.cell(row=r_idx, column=2, value=item["ho_ten"])
        apply_cell_style(c2, font=Font(name="Times New Roman", size=11, bold=(idx <= 3)), fill=fill_to_use, alignment=LEFT, border=THIN_BORDER)

        c3 = ws_rank.cell(row=r_idx, column=3, value=item["lop"])
        apply_cell_style(c3, font=DATA_FONT, fill=fill_to_use, alignment=CENTER, border=THIN_BORDER)

        c4 = ws_rank.cell(row=r_idx, column=4, value=item["vong"])
        apply_cell_style(c4, font=Font(name="Times New Roman", size=11, bold=True), fill=fill_to_use, alignment=CENTER, border=THIN_BORDER)

        c5 = ws_rank.cell(row=r_idx, column=5, value=item["diem"])
        apply_cell_style(c5, font=Font(name="Times New Roman", size=11, bold=True, color="006100"), fill=fill_to_use, alignment=CENTER, border=THIN_BORDER)

        c6 = ws_rank.cell(row=r_idx, column=6, value=item["thoi_gian"])
        apply_cell_style(c6, font=DATA_FONT, fill=fill_to_use, alignment=CENTER, border=THIN_BORDER)

        c7 = ws_rank.cell(row=r_idx, column=7, value=award)
        apply_cell_style(c7, font=Font(name="Times New Roman", size=11, italic=True), fill=fill_to_use, alignment=LEFT, border=THIN_BORDER)

        ws_rank.row_dimensions[r_idx].height = 21

    auto_fit_columns(ws_rank)
    ws_rank.column_dimensions["A"].width = 10
    ws_rank.column_dimensions["B"].width = 28
    ws_rank.column_dimensions["G"].width = 26
    ws_rank.freeze_panes = "A5"

    # ----------------------------------------------------
    # SHEET 3: NHẮC NHỞ GVCN (CHƯA THI)
    # ----------------------------------------------------
    ws_gv = wb.create_sheet("Nhắc nhở GVCN (Chưa thi)")

    ws_gv.merge_cells("A1:E1")
    tgv_cell = ws_gv.cell(row=1, column=1, value=f"DANH SÁCH HỌC SINH CHƯA THI {contest_name.upper()} - LẦN {lan_num}")
    apply_cell_style(tgv_cell, font=TITLE_FONT, alignment=CENTER)
    ws_gv.row_dimensions[1].height = 28

    ws_gv.merge_cells("A2:E2")
    subgv_cell = ws_gv.cell(row=2, column=1, value="Dành cho GVCN đôn đốc, nhắc nhở học sinh hoàn thành tự luyện đúng hạn")
    apply_cell_style(subgv_cell, font=Font(name="Times New Roman", size=11, italic=True, color="595959"), alignment=CENTER)
    ws_gv.row_dimensions[2].height = 20

    row_gv = 4
    for cls in classes_order:
        chua_thi_lop = [h for h in by_class[cls] if matches.get((h, cls)) is None]
        total_lop = len(by_class[cls])
        da_thi_lop = total_lop - len(chua_thi_lop)

        if not chua_thi_lop:
            continue

        ws_gv.merge_cells(start_row=row_gv, start_column=1, end_row=row_gv, end_column=5)
        h_text = f"Lớp {cls} — {da_thi_lop}/{total_lop} đã thi ({da_thi_lop*100//total_lop}%) | {len(chua_thi_lop)}/{total_lop} chưa thi"
        h_cell = ws_gv.cell(row=row_gv, column=1, value=h_text)
        apply_cell_style(h_cell, font=CLASS_HEADER_FONT, fill=CLASS_HEADER_FILL, alignment=LEFT, border=THIN_BORDER)
        for c in range(2, 6):
            apply_cell_style(ws_gv.cell(row=row_gv, column=c), border=THIN_BORDER)
        ws_gv.row_dimensions[row_gv].height = 22
        row_gv += 1

        sub_h = ["STT", "Họ và tên", "Lớp", "Trạng thái", "Ghi chú"]
        for col_idx, sh in enumerate(sub_h, 1):
            sc = ws_gv.cell(row=row_gv, column=col_idx, value=sh)
            apply_cell_style(sc, font=SUBHEADER_FONT, fill=SUBHEADER_FILL, alignment=CENTER, border=THIN_BORDER)
        ws_gv.row_dimensions[row_gv].height = 22
        row_gv += 1

        for idx, ho_ten in enumerate(chua_thi_lop, 1):
            ws_gv.cell(row=row_gv, column=1, value=idx)
            apply_cell_style(ws_gv.cell(row=row_gv, column=1), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

            ws_gv.cell(row=row_gv, column=2, value=ho_ten)
            apply_cell_style(ws_gv.cell(row=row_gv, column=2), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)

            ws_gv.cell(row=row_gv, column=3, value=cls)
            apply_cell_style(ws_gv.cell(row=row_gv, column=3), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

            status = ws_gv.cell(row=row_gv, column=4, value="❌ Chưa thi")
            apply_cell_style(status, font=RED_FONT, fill=RED_FILL, alignment=CENTER, border=THIN_BORDER)

            apply_cell_style(ws_gv.cell(row=row_gv, column=5), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)
            ws_gv.row_dimensions[row_gv].height = 20
            row_gv += 1

        row_gv += 1

    auto_fit_columns(ws_gv)
    ws_gv.column_dimensions["A"].width = 6
    ws_gv.column_dimensions["B"].width = 28
    ws_gv.freeze_panes = "A4"

    # ----------------------------------------------------
    # SHEET 4: THỐNG KÊ THEO LỚP
    # ----------------------------------------------------
    ws_cls = wb.create_sheet("Thống kê theo lớp")

    ws_cls.merge_cells("A1:H1")
    tcls_cell = ws_cls.cell(row=1, column=1, value=f"BẢNG THỐNG KÊ TỶ LỆ THAM GIA THEO TỪNG LỚP - {contest_name.upper()} LẦN {lan_num}")
    apply_cell_style(tcls_cell, font=TITLE_FONT, alignment=CENTER)
    ws_cls.row_dimensions[1].height = 28

    cls_headers = ["STT", "Lớp", "Sĩ số", "Đã thi", "Chưa thi", "Tỷ lệ tham gia", "Điểm cao nhất", "Thủ khoa của lớp"]
    for col_idx, ch in enumerate(cls_headers, 1):
        cc = ws_cls.cell(row=3, column=col_idx, value=ch)
        apply_cell_style(cc, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=THIN_BORDER)
    ws_cls.row_dimensions[3].height = 28

    row_c = 4
    for idx, cls in enumerate(classes_order, 1):
        total_lop = len(by_class[cls])
        da_thi_lop = sum(1 for h in by_class[cls] if matches.get((h, cls)) is not None)
        chua_thi_lop = total_lop - da_thi_lop
        pct = (da_thi_lop * 100 // total_lop) if total_lop > 0 else 0

        # Find top scorer in class
        top_scorer = ""
        max_score = ""
        lop_completed = [
            (h, matches.get((h, cls)))
            for h in by_class[cls]
            if matches.get((h, cls)) is not None
        ]
        if lop_completed:
            def sort_lop_key(x):
                info = x[1]
                try: v = int(info.get("vòng", 0) or 0)
                except: v = 0
                try: d = int(info.get("điểm", 0) or 0)
                except: d = 0
                return (v, d)
            lop_completed.sort(key=sort_lop_key, reverse=True)
            top_scorer = lop_completed[0][0]
            max_score = f"{lop_completed[0][1].get('điểm', '')} đ (Vòng {lop_completed[0][1].get('vòng', '')})"

        ws_cls.cell(row=row_c, column=1, value=idx)
        apply_cell_style(ws_cls.cell(row=row_c, column=1), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

        ws_cls.cell(row=row_c, column=2, value=cls)
        apply_cell_style(ws_cls.cell(row=row_c, column=2), font=Font(name="Times New Roman", size=11, bold=True), alignment=CENTER, border=THIN_BORDER)

        ws_cls.cell(row=row_c, column=3, value=total_lop)
        apply_cell_style(ws_cls.cell(row=row_c, column=3), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

        ws_cls.cell(row=row_c, column=4, value=da_thi_lop)
        apply_cell_style(ws_cls.cell(row=row_c, column=4), font=Font(name="Times New Roman", size=11, bold=True, color="006100"), alignment=CENTER, border=THIN_BORDER)

        ws_cls.cell(row=row_c, column=5, value=chua_thi_lop)
        apply_cell_style(ws_cls.cell(row=row_c, column=5), font=Font(name="Times New Roman", size=11, bold=True, color="9C0006"), alignment=CENTER, border=THIN_BORDER)

        ws_cls.cell(row=row_c, column=6, value=f"{pct}%")
        apply_cell_style(ws_cls.cell(row=row_c, column=6), font=Font(name="Times New Roman", size=11, bold=True), alignment=CENTER, border=THIN_BORDER)

        ws_cls.cell(row=row_c, column=7, value=max_score)
        apply_cell_style(ws_cls.cell(row=row_c, column=7), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

        ws_cls.cell(row=row_c, column=8, value=top_scorer)
        apply_cell_style(ws_cls.cell(row=row_c, column=8), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)

        ws_cls.row_dimensions[row_c].height = 21
        row_c += 1

    # Total row
    ws_cls.merge_cells(start_row=row_c, start_column=1, end_row=row_c, end_column=2)
    t_sum = ws_cls.cell(row=row_c, column=1, value="TOÀN TRƯỜNG")
    apply_cell_style(t_sum, font=Font(name="Times New Roman", size=11, bold=True), alignment=CENTER, border=THIN_BORDER)
    apply_cell_style(ws_cls.cell(row=row_c, column=2), border=THIN_BORDER)

    ws_cls.cell(row=row_c, column=3, value=total_school)
    apply_cell_style(ws_cls.cell(row=row_c, column=3), font=Font(name="Times New Roman", size=11, bold=True), alignment=CENTER, border=THIN_BORDER)

    ws_cls.cell(row=row_c, column=4, value=da_thi_count)
    apply_cell_style(ws_cls.cell(row=row_c, column=4), font=Font(name="Times New Roman", size=11, bold=True, color="006100"), alignment=CENTER, border=THIN_BORDER)

    ws_cls.cell(row=row_c, column=5, value=chua_thi_count)
    apply_cell_style(ws_cls.cell(row=row_c, column=5), font=Font(name="Times New Roman", size=11, bold=True, color="9C0006"), alignment=CENTER, border=THIN_BORDER)

    ws_cls.cell(row=row_c, column=6, value=f"{pct_da_thi}%")
    apply_cell_style(ws_cls.cell(row=row_c, column=6), font=Font(name="Times New Roman", size=11, bold=True), alignment=CENTER, border=THIN_BORDER)

    apply_cell_style(ws_cls.cell(row=row_c, column=7), border=THIN_BORDER)
    apply_cell_style(ws_cls.cell(row=row_c, column=8), border=THIN_BORDER)
    ws_cls.row_dimensions[row_c].height = 22

    auto_fit_columns(ws_cls)
    ws_cls.column_dimensions["A"].width = 6
    ws_cls.column_dimensions["B"].width = 10
    ws_cls.column_dimensions["H"].width = 25
    ws_cls.freeze_panes = "A4"

    # Save to round folder
    out_file = round_dir / f"Thống kê {contest_name} - Lần {lan_num}.xlsx"
    try:
        wb.save(out_file)
        print(f"   ✅ Đã tạo file thống kê đợt thi: {out_file}")
    except PermissionError:
        alt_f = round_dir / f"Thống kê {contest_name} - Lần {lan_num} (new).xlsx"
        wb.save(alt_f)
        print(f"   ⚠ File đang mở, lưu bản mới: {alt_f}")

    wb.close()
    return out_file


# ============================================================
# STEP 6: Generate Master output Excel
# ============================================================

def generate_output(master_data, contests_data, all_matches, output_path):
    """Generate the final Master Excel report."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    all_students = master_data["all_students"]
    classes_order = master_data["classes_order"]
    by_class = master_data["by_class"]

    # ========================================
    # SHEET 1: TỔNG HỢP
    # ========================================
    ws_tong = wb.create_sheet("Tổng hợp")

    # Build column structure
    contest_columns = []
    for contest_name in sorted(contests_data.keys()):
        for lan_num in sorted(contests_data[contest_name].keys()):
            contest_columns.append((contest_name, lan_num))

    # Title row
    ws_tong.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3 + len(contest_columns))
    title_cell = ws_tong.cell(row=1, column=1, value="THỐNG KÊ CUỘC THI HỌC SINH UNIGO - NĂM HỌC 2026-2027")
    apply_cell_style(title_cell, font=TITLE_FONT, alignment=CENTER)
    ws_tong.row_dimensions[1].height = 30

    # Header row
    row = 3
    headers = ["STT", "Họ và tên", "Lớp"]
    for contest_name, lan_num in contest_columns:
        headers.append(f"{contest_name}\nLần {lan_num}")

    if contest_columns:
        headers.append("Tổng cuộc thi\nđã tham gia")

    for col_idx, header in enumerate(headers, 1):
        cell = ws_tong.cell(row=row, column=col_idx, value=header)
        apply_cell_style(cell, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=THIN_BORDER)

    ws_tong.row_dimensions[row].height = 35

    # Data rows — grouped by class
    data_row = row + 1
    stt = 0
    for cls in classes_order:
        ws_tong.merge_cells(
            start_row=data_row, start_column=1,
            end_row=data_row, end_column=len(headers)
        )
        cls_cell = ws_tong.cell(row=data_row, column=1, value=f"Lớp {cls}")
        apply_cell_style(cls_cell, font=CLASS_HEADER_FONT, fill=CLASS_HEADER_FILL, alignment=LEFT, border=THIN_BORDER)
        for c in range(2, len(headers) + 1):
            apply_cell_style(ws_tong.cell(row=data_row, column=c), border=THIN_BORDER)
        ws_tong.row_dimensions[data_row].height = 22
        data_row += 1

        for ho_ten in by_class[cls]:
            stt += 1
            key = (ho_ten, cls)

            ws_tong.cell(row=data_row, column=1, value=stt)
            apply_cell_style(ws_tong.cell(row=data_row, column=1), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

            ws_tong.cell(row=data_row, column=2, value=ho_ten)
            apply_cell_style(ws_tong.cell(row=data_row, column=2), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)

            ws_tong.cell(row=data_row, column=3, value=cls)
            apply_cell_style(ws_tong.cell(row=data_row, column=3), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

            total_participated = 0
            for col_offset, (contest_name, lan_num) in enumerate(contest_columns):
                match_key = (contest_name, lan_num)
                match_result = all_matches.get(match_key, {}).get(key)

                col = 4 + col_offset
                cell = ws_tong.cell(row=data_row, column=col)

                if match_result is not None:
                    cell.value = "✅ Đã thi"
                    apply_cell_style(cell, font=GREEN_FONT, fill=GREEN_FILL, alignment=CENTER, border=THIN_BORDER)
                    total_participated += 1
                else:
                    cell.value = "❌ Chưa thi"
                    apply_cell_style(cell, font=RED_FONT, fill=RED_FILL, alignment=CENTER, border=THIN_BORDER)

            # Summary column
            if contest_columns:
                sum_col = 4 + len(contest_columns)
                sum_cell = ws_tong.cell(row=data_row, column=sum_col, value=total_participated)
                apply_cell_style(sum_cell, font=Font(name="Times New Roman", size=11, bold=True),
                                alignment=CENTER, border=THIN_BORDER)

            data_row += 1

    # Statistics row
    data_row += 1
    ws_tong.merge_cells(start_row=data_row, start_column=1, end_row=data_row, end_column=3)
    stat_cell = ws_tong.cell(row=data_row, column=1, value="THỐNG KÊ TỔNG:")
    apply_cell_style(stat_cell, font=Font(name="Times New Roman", size=12, bold=True), alignment=LEFT)

    for col_offset, (contest_name, lan_num) in enumerate(contest_columns):
        match_key = (contest_name, lan_num)
        matches = all_matches.get(match_key, {})
        participated = sum(1 for v in matches.values() if v is not None)
        total = len(matches)
        col = 4 + col_offset
        cell = ws_tong.cell(row=data_row, column=col,
                            value=f"{participated}/{total} ({participated*100//total}%)")
        apply_cell_style(cell, font=Font(name="Times New Roman", size=11, bold=True), alignment=CENTER)

    auto_fit_columns(ws_tong)
    ws_tong.column_dimensions["A"].width = 6
    ws_tong.column_dimensions["B"].width = 28
    ws_tong.column_dimensions["C"].width = 8
    ws_tong.freeze_panes = "D4"

    # ========================================
    # SHEET 2+: Chi tiết từng cuộc thi
    # ========================================
    for contest_name in sorted(contests_data.keys()):
        ws_ct = wb.create_sheet(f"Chi tiết {contest_name}")

        ws_ct.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        title = ws_ct.cell(row=1, column=1,
                           value=f"CHI TIẾT KẾT QUẢ {contest_name.upper()} - NĂM HỌC 2026-2027")
        apply_cell_style(title, font=TITLE_FONT, alignment=CENTER)
        ws_ct.row_dimensions[1].height = 30

        row = 3
        for lan_num in sorted(contests_data[contest_name].keys()):
            match_key = (contest_name, lan_num)
            matches = all_matches.get(match_key, {})

            # Lần header
            ws_ct.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            lan_cell = ws_ct.cell(row=row, column=1, value=f"Lần {lan_num}")
            apply_cell_style(lan_cell, font=SUBHEADER_FONT, fill=SUBHEADER_FILL, alignment=LEFT, border=THIN_BORDER)
            for c in range(2, 9):
                apply_cell_style(ws_ct.cell(row=row, column=c), border=THIN_BORDER)
            row += 1

            # Detail headers
            detail_headers = ["STT", "Họ và tên", "Lớp", "Trạng thái", "Vòng", "Điểm", "Thời gian", "Ghi chú"]
            for col_idx, h in enumerate(detail_headers, 1):
                cell = ws_ct.cell(row=row, column=col_idx, value=h)
                apply_cell_style(cell, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=THIN_BORDER)
            row += 1

            # Data - group by class
            stt_detail = 0
            for cls in classes_order:
                for ho_ten in by_class[cls]:
                    stt_detail += 1
                    key = (ho_ten, cls)
                    info = matches.get(key)

                    ws_ct.cell(row=row, column=1, value=stt_detail)
                    apply_cell_style(ws_ct.cell(row=row, column=1), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

                    ws_ct.cell(row=row, column=2, value=ho_ten)
                    apply_cell_style(ws_ct.cell(row=row, column=2), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)

                    ws_ct.cell(row=row, column=3, value=cls)
                    apply_cell_style(ws_ct.cell(row=row, column=3), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

                    status_cell = ws_ct.cell(row=row, column=4)
                    if info is not None:
                        status_cell.value = "✅ Đã thi"
                        apply_cell_style(status_cell, font=GREEN_FONT, fill=GREEN_FILL, alignment=CENTER, border=THIN_BORDER)

                        vong = info.get("vòng", "")
                        diem = info.get("điểm", "")
                        thoi_gian = info.get("thời_gian", "")

                        ws_ct.cell(row=row, column=5, value=vong)
                        ws_ct.cell(row=row, column=6, value=diem)
                        ws_ct.cell(row=row, column=7, value=thoi_gian)
                    else:
                        status_cell.value = "❌ Chưa thi"
                        apply_cell_style(status_cell, font=RED_FONT, fill=RED_FILL, alignment=CENTER, border=THIN_BORDER)

                    for c in range(5, 9):
                        apply_cell_style(ws_ct.cell(row=row, column=c), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

                    row += 1

            row += 1

        auto_fit_columns(ws_ct)
        ws_ct.column_dimensions["A"].width = 6
        ws_ct.column_dimensions["B"].width = 28
        ws_ct.freeze_panes = "D4"

    # ========================================
    # SHEET LAST: NHẮC NHỞ GV
    # ========================================
    ws_gv = wb.create_sheet("Nhắc nhở GV")

    ws_gv.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    title = ws_gv.cell(row=1, column=1,
                       value="NHẮC NHỞ GIÁO VIÊN CHỦ NHIỆM - HỌC SINH CHƯA THI")
    apply_cell_style(title, font=TITLE_FONT, alignment=CENTER)
    ws_gv.row_dimensions[1].height = 30

    row = 3
    for cls in classes_order:
        for contest_name in sorted(contests_data.keys()):
            for lan_num in sorted(contests_data[contest_name].keys()):
                match_key = (contest_name, lan_num)
                matches = all_matches.get(match_key, {})

                chua_thi = []
                for ho_ten in by_class[cls]:
                    key = (ho_ten, cls)
                    if matches.get(key) is None:
                        chua_thi.append(ho_ten)

                if not chua_thi:
                    continue

                total_class = len(by_class[cls])
                da_thi = total_class - len(chua_thi)

                ws_gv.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
                header_text = (
                    f"Lớp {cls} — {contest_name} Lần {lan_num} "
                    f"({da_thi}/{total_class} đã thi, "
                    f"{len(chua_thi)}/{total_class} chưa thi)"
                )
                h_cell = ws_gv.cell(row=row, column=1, value=header_text)
                apply_cell_style(h_cell, font=CLASS_HEADER_FONT, fill=CLASS_HEADER_FILL, alignment=LEFT, border=THIN_BORDER)
                for c in range(2, 6):
                    apply_cell_style(ws_gv.cell(row=row, column=c), border=THIN_BORDER)
                row += 1

                sub_headers = ["STT", "Họ và tên", "Lớp", "Trạng thái", "Ghi chú"]
                for col_idx, h in enumerate(sub_headers, 1):
                    cell = ws_gv.cell(row=row, column=col_idx, value=h)
                    apply_cell_style(cell, font=SUBHEADER_FONT, fill=SUBHEADER_FILL, alignment=CENTER, border=THIN_BORDER)
                row += 1

                for idx, ho_ten in enumerate(chua_thi, 1):
                    ws_gv.cell(row=row, column=1, value=idx)
                    apply_cell_style(ws_gv.cell(row=row, column=1), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

                    ws_gv.cell(row=row, column=2, value=ho_ten)
                    apply_cell_style(ws_gv.cell(row=row, column=2), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)

                    ws_gv.cell(row=row, column=3, value=cls)
                    apply_cell_style(ws_gv.cell(row=row, column=3), font=DATA_FONT, alignment=CENTER, border=THIN_BORDER)

                    status = ws_gv.cell(row=row, column=4, value="❌ Chưa thi")
                    apply_cell_style(status, font=RED_FONT, fill=RED_FILL, alignment=CENTER, border=THIN_BORDER)

                    apply_cell_style(ws_gv.cell(row=row_gv if 'row_gv' in locals() else row, column=5), font=DATA_FONT, alignment=LEFT, border=THIN_BORDER)

                    row += 1

                row += 1

    auto_fit_columns(ws_gv)
    ws_gv.column_dimensions["A"].width = 6
    ws_gv.column_dimensions["B"].width = 28
    ws_gv.freeze_panes = "A3"

    # Save Master
    try:
        wb.save(output_path)
        print(f"\n✅ Đã lưu file thống kê tổng thể: {output_path}")
    except PermissionError:
        alt_path = output_path.with_stem(output_path.stem + " (new)")
        wb.save(alt_path)
        print(f"\n⚠ File gốc đang mở, đã lưu bản mới: {alt_path}")

    wb.close()


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print("  THỐNG KÊ CUỘC THI HỌC SINH UNIGO")
    print("=" * 60)

    # Step 1: Read master student list
    print(f"\n📋 Đọc danh sách HS: {DS_HS_FILE.name}")
    master = read_master_students(DS_HS_FILE)
    print(f"   → {len(master['all_students'])} học sinh, {len(master['classes_order'])} lớp")
    for cls in master["classes_order"]:
        print(f"     • {cls}: {len(master['by_class'][cls])} HS")

    # Step 2: Scan contests
    print(f"\n🔍 Quét thư mục cuộc thi: {BASE_DIR}")
    contests = scan_contests(BASE_DIR)
    if not contests:
        print("   ⚠ Không tìm thấy cuộc thi nào!")
        return

    for contest_name, rounds in contests.items():
        print(f"   📁 {contest_name}:")
        for lan_num, files in sorted(rounds.items()):
            print(f"      • Lần {lan_num}: {len(files)} file(s)")

    # Step 3 & 4: Read results and match
    print(f"\n📊 Đọc kết quả và so khớp...")
    all_matches = {}  # {(contest_name, lan_num): {(ho_ten, lop): info_or_None}}

    for contest_name, rounds in contests.items():
        for lan_num, files in sorted(rounds.items()):
            print(f"\n   📄 {contest_name} - Lần {lan_num}:")
            results = read_contest_results(files)
            print(f"      → Đọc được {len(results)} bản ghi từ file raw")

            matched = match_students(master["all_students"], results)
            participated = sum(1 for v in matched.values() if v is not None)
            not_participated = sum(1 for v in matched.values() if v is None)
            print(f"      ✅ Đã thi: {participated} | ❌ Chưa thi: {not_participated}")

            all_matches[(contest_name, lan_num)] = matched

            # Step 5: Generate round-specific file directly in that round's folder!
            round_dir = files[0].parent
            generate_round_report(master, contest_name, lan_num, matched, round_dir)

    # Step 6: Generate master output
    print(f"\n📝 Tạo file thống kê tổng thể...")
    generate_output(master, contests, all_matches, OUTPUT_FILE)

    # Summary
    print(f"\n{'=' * 60}")
    print(f"  HOÀN TẤT!")
    print(f"{'=' * 60}")
    print(f"  📁 Master Output: {OUTPUT_FILE}")
    for contest_name, rounds in contests.items():
        for lan_num, files in sorted(rounds.items()):
            round_file = files[0].parent / f"Thống kê {contest_name} - Lần {lan_num}.xlsx"
            print(f"  📁 Round Output:  {round_file}")


if __name__ == "__main__":
    main()
