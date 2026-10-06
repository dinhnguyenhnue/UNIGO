# -*- coding: utf-8 -*-
"""
Script chuẩn hóa toàn bộ 6 đề kiểm tra Tin học Đánh giá định kỳ 1 (Lớp 3 - Lớp 8).

Nhiệm vụ:
1. Đảo và chia đều đáp án trắc nghiệm:
   - Mỗi đề 20 câu.
   - Chia đều tuyệt đối: đúng 5 câu A, 5 câu B, 5 câu C, 5 câu D (mỗi chữ cái 25%).
   - Bảo toàn 100% nội dung câu hỏi và tính đúng đắn của phương án đúng.
   - Cập nhật tự động Bảng Table 2 & Table 3 (bảng đáp án).
2. Chuẩn hóa Phần II (Thực hành & Bài tập):
   - Ghi rõ ràng: [Thực hành trên máy tính] và [Trả lời viết ra trên giấy].
   - Bổ sung các hàng dòng chấm '......' đầy đủ để học sinh làm bài trực tiếp vào đề thi.
   - Cập nhật Hướng dẫn chấm chi tiết tương ứng.
3. Bảo tồn tuyệt đối Header (logo drawing UNIGO) và Footer của file mẫu.
4. Đảm bảo toàn bộ văn bản chuẩn Times New Roman 13pt.
"""

import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
import re
import random
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ============================================================
# CẤU HÌNH PHẦN II CHO TỪNG LỚP (GHI RÕ MÁY / GIẤY & DÒNG CHẤM)
# ============================================================

PART2_CONTENT = {
    3: {
        "title": "PHẦN II. THỰC HÀNH VÀ BÀI TẬP (2 câu – 5,0 điểm)",
        "tasks": [
            {
                "header": "Câu 1 (2,0 điểm) – Trả lời viết ra trên giấy: Quan sát máy tính để bàn trong phòng Tin học và ghi tên 4 bộ phận chính:",
                "items": [
                    "+ Bộ phận 1: ....................................................................................................................................",
                    "+ Bộ phận 2: ....................................................................................................................................",
                    "+ Bộ phận 3: ....................................................................................................................................",
                    "+ Bộ phận 4: ...................................................................................................................................."
                ]
            },
            {
                "header": "Câu 2 (3,0 điểm) – Thực hành trên máy tính: Em hãy thực hiện các thao tác sau dưới sự quan sát của giáo viên:",
                "items": [
                    "a) Khởi động (bật) máy tính đúng quy trình, chờ màn hình làm việc xuất hiện. (1,0 điểm)",
                    "b) Mở phần mềm luyện gõ bàn phím (hoặc phần mềm Paint), thực hiện thao tác nhấp chuột di chuyển con trỏ. (1,0 điểm)",
                    "c) Tắt máy tính đúng cách qua menu Start → Power → Shut Down. (1,0 điểm)"
                ]
            }
        ],
        "grading": [
            "Câu 1 (2,0 điểm) – Trả lời viết ra trên giấy:",
            "- Ghi đúng tên 4 bộ phận chính: Màn hình, Thân máy, Bàn phím, Chuột (mỗi bộ phận đúng được 0,5 điểm).",
            "Câu 2 (3,0 điểm) – Thực hành trên máy tính:",
            "a) Nhấn nút nguồn (Power), bật máy tính đúng quy trình (1,0 điểm).",
            "b) Mở đúng phần mềm và thao tác chuột thành thạo (1,0 điểm).",
            "c) Thực hiện thao tác tắt máy tính đúng quy trình Start → Shut Down (1,0 điểm)."
        ]
    },
    4: {
        "title": "PHẦN II. THỰC HÀNH VÀ BÀI TẬP (2 câu – 5,0 điểm)",
        "tasks": [
            {
                "header": "Câu 1 (2,0 điểm) – Thực hành trên máy tính: Soạn thảo văn bản và lưu tệp:",
                "items": [
                    "a) Mở phần mềm soạn thảo văn bản Microsoft Word (hoặc Notepad). (0,5 điểm)",
                    "b) Gõ đúng đoạn thơ sau (sử dụng kiểu gõ Telex hoặc Vni, ngồi đúng tư thế): (1,0 điểm)",
                    "   \"Hạt gạo làng ta",
                    "   Có vị phù sa",
                    "   Của sông Kinh Thầy...\"",
                    "c) Lưu tệp văn bản vào màn hình Desktop với tên: 'BaiKiemTra_HoVaTen.docx'. (0,5 điểm)"
                ]
            },
            {
                "header": "Câu 2 (3,0 điểm) – Thực hành tìm kiếm trên máy tính và trả lời viết ra trên giấy:",
                "items": [
                    "a) [Thực hành trên máy tính] Mở trình duyệt web Google Chrome, tìm kiếm với từ khóa: \"Danh lam thang canh Viet Nam\". (1,0 điểm)",
                    "b) [Trả lời viết ra trên giấy] Chọn 1 bài viết uy tín từ kết quả tìm kiếm, ghi lại tên 3 thắng cảnh nổi tiếng của Việt Nam: (1,0 điểm)",
                    "- Thắng cảnh 1: ................................................................................................................................",
                    "- Thắng cảnh 2: ................................................................................................................................",
                    "- Thắng cảnh 3: ................................................................................................................................",
                    "c) [Trả lời viết ra trên giấy] Ghi lại tên trang web hoặc địa chỉ liên kết (URL) mà em đã tham khảo thông tin: (1,0 điểm)",
                    "- Tên trang web: ................................................................................................................................",
                    "- Địa chỉ trang web tham khảo: ............................................................................................................"
                ]
            }
        ],
        "grading": [
            "Câu 1 (2,0 điểm) – Thực hành trên máy tính:",
            "a) Mở đúng phần mềm Microsoft Word hoặc Notepad (0,5 điểm).",
            "b) Gõ đúng chính tả đoạn thơ, gõ có dấu Tiếng Việt, ngồi đúng tư thế (1,0 điểm).",
            "c) Lưu tệp đúng tên quy định vào màn hình Desktop (0,5 điểm).",
            "Câu 2 (3,0 điểm) – Thực hành tìm kiếm và trả lời viết ra trên giấy:",
            "a) Mở trình duyệt Google Chrome và nhập đúng từ khóa tìm kiếm (1,0 điểm).",
            "b) Ghi đúng tên 3 thắng cảnh nổi tiếng của Việt Nam (1,0 điểm).",
            "c) Ghi đúng tên trang web và địa chỉ đường liên kết uy tín (1,0 điểm)."
        ]
    },
    5: {
        "title": "PHẦN II. THỰC HÀNH VÀ BÀI TẬP (2 câu – 5,0 điểm)",
        "tasks": [
            {
                "header": "Câu 1 (2,0 điểm) – Thực hành trên máy tính: Tổ chức cây thư mục trên máy tính:",
                "items": [
                    "a) Mở File Explorer, tạo thư mục có tên 'HOC_TAP' trên Desktop. (0,5 điểm)",
                    "b) Trong thư mục 'HOC_TAP', tạo 3 thư mục con: 'TOAN', 'TIENG_VIET', 'TIN_HOC'. (0,75 điểm)",
                    "c) Tạo 1 tệp văn bản Word tên 'BaiTap1.docx' đặt bên trong thư mục 'TIN_HOC'. (0,75 điểm)"
                ]
            },
            {
                "header": "Câu 2 (3,0 điểm) – Thực hành tìm kiếm trên máy tính và trả lời viết ra trên giấy:",
                "items": [
                    "Tình huống: Em cần tìm hiểu về \"5 loài động vật quý hiếm ở Việt Nam\" để chuẩn bị bài thuyết trình.",
                    "a) [Thực hành trên máy tính] Mở trình duyệt Google Chrome, nhập từ khóa tìm kiếm phù hợp để tra cứu thông tin. (1,0 điểm)",
                    "b) [Trả lời viết ra trên giấy] Chọn kết quả từ một nguồn thông tin đáng tin cậy, ghi lại tên 5 loài động vật quý hiếm: (1,0 điểm)",
                    "1. ......................................................................  2. ......................................................................",
                    "3. ......................................................................  4. ......................................................................",
                    "5. ................................................................................................................................................",
                    "c) [Trả lời viết ra trên giấy] Ghi rõ nguồn tham khảo gồm tên trang web và đường liên kết dẫn tới bài viết: (1,0 điểm)",
                    "- Tên trang web: ................................................................................................................................",
                    "- Đường liên kết (URL): ........................................................................................................................"
                ]
            }
        ],
        "grading": [
            "Câu 1 (2,0 điểm) – Thực hành trên máy tính:",
            "a) Tạo đúng thư mục HOC_TAP trên Desktop (0,5 điểm).",
            "b) Tạo đúng 3 thư mục con TOAN, TIENG_VIET, TIN_HOC (0,75 điểm).",
            "c) Tạo đúng tệp BaiTap1.docx trong thư mục TIN_HOC (0,75 điểm).",
            "Câu 2 (3,0 điểm) – Thực hành tìm kiếm và trả lời viết ra trên giấy:",
            "a) Mở trình duyệt và tìm kiếm từ khóa phù hợp (1,0 điểm).",
            "b) Chọn nguồn đáng tin cậy, ghi đúng 5 loài động vật quý hiếm (1,0 điểm).",
            "c) Ghi đầy đủ tên trang web và đường liên kết tham khảo (1,0 điểm)."
        ]
    },
    6: {
        "title": "PHẦN II. THỰC HÀNH VÀ BÀI TẬP (2 câu – 5,0 điểm)",
        "tasks": [
            {
                "header": "Câu 1 (2,0 điểm) – Trả lời viết ra trên giấy: Tính toán quy đổi đơn vị dung lượng và giải bài toán lưu trữ:",
                "items": [
                    "a) Thực hiện quy đổi các đơn vị đo dung lượng thông tin: (1,0 điểm)",
                    "   1 GB = ................................... MB = ................................... KB = ................................... Byte",
                    "b) Một ổ cứng di động có dung lượng 256 GB. Hỏi ổ cứng đó có thể lưu trữ tối đa được bao nhiêu bộ phim chất lượng cao, biết mỗi bộ phim có dung lượng trung bình là 4 GB? (1,0 điểm)",
                    "- Phép tính / Lời giải: ......................................................................................................................",
                    "....................................................................................................................................................",
                    "- Đáp số: ........................................................................................................................................"
                ]
            },
            {
                "header": "Câu 2 (3,0 điểm) – Nhận diện thiết bị mạng trên máy và trả lời viết ra trên giấy:",
                "items": [
                    "Em hãy quan sát hệ thống mạng trong phòng Tin học và thực hiện:",
                    "a) [Thực hành] Chỉ ra vị trí của thiết bị chuyển mạch (Switch), dây cáp mạng và bộ định tuyến (Router). (1,0 điểm)",
                    "b) [Trả lời viết ra trên giấy] Em hãy giải thích ngắn gọn cách các máy tính trong phòng Tin học được kết nối với nhau để tạo thành mạng cục bộ (LAN): (1,0 điểm)",
                    "....................................................................................................................................................",
                    "....................................................................................................................................................",
                    "c) [Thực hành trên máy tính] Mở trình duyệt web, truy cập trang google.com để kiểm tra máy tính có kết nối Internet hay không và báo cáo kết quả cho giáo viên. (1,0 điểm)"
                ]
            }
        ],
        "grading": [
            "Câu 1 (2,0 điểm) – Trả lời viết ra trên giấy:",
            "a) 1 GB = 1024 MB = 1.048.576 KB = 1.073.741.824 Byte (1,0 điểm).",
            "b) Thực hiện đúng phép chia: 256 ÷ 4 = 64 (bộ phim) (1,0 điểm).",
            "Câu 2 (3,0 điểm) – Nhận diện thiết bị mạng và trả lời viết ra trên giấy:",
            "a) Nhận diện đúng các thiết bị mạng trong phòng học: Switch, cáp mạng, Router (1,0 điểm).",
            "b) Giải thích đúng: các máy tính nối dây mạng vào Switch/Router để chia sẻ dữ liệu và kết nối Internet (1,0 điểm).",
            "c) Mở trình duyệt web và truy cập thành công trang google.com (1,0 điểm)."
        ]
    },
    7: {
        "title": "PHẦN II. THỰC HÀNH VÀ BÀI TẬP (2 câu – 5,0 điểm)",
        "tasks": [
            {
                "header": "Câu 1 (2,0 điểm) – Thực hành trên máy tính: Quản lí tệp và thư mục:",
                "items": [
                    "Em hãy thực hiện các thao tác sau trên máy tính:",
                    "a) Mở File Explorer, tạo thư mục 'BAI_KIEM_TRA' trên Desktop. (0,5 điểm)",
                    "b) Trong thư mục 'BAI_KIEM_TRA', tạo 2 thư mục con: 'LY_THUYET' và 'THUC_HANH'. (0,5 điểm)",
                    "c) Tạo 1 tệp văn bản Word đặt tên là 'TraLoi.docx' bên trong thư mục 'LY_THUYET'. (0,5 điểm)",
                    "d) Sử dụng phím tắt (Ctrl+C, Ctrl+V) để sao chép tệp 'TraLoi.docx' sang thư mục 'THUC_HANH'. (0,5 điểm)"
                ]
            },
            {
                "header": "Câu 2 (3,0 điểm) – Trả lời viết ra trên giấy và thực hành soạn thảo trên máy tính:",
                "items": [
                    "a) [Trả lời viết ra trên giấy] Em hãy kể tên 3 phần mềm hệ thống và 3 phần mềm ứng dụng đang được cài đặt trên máy tính: (1,0 điểm)",
                    "- Ba phần mềm hệ thống: ....................................................................................................................",
                    "- Ba phần mềm ứng dụng: ....................................................................................................................",
                    "b) [Thực hành trên máy tính] Mở Microsoft Word, soạn thảo một đoạn văn ngắn (từ 3 đến 5 câu) giới thiệu về ngôi trường Tiểu học & THCS UNIGO nơi em đang học. (1,0 điểm)",
                    "c) [Thực hành trên máy tính] Lưu tệp văn bản vừa soạn với tên 'GioiThieu_UNIGO.docx' vào thư mục 'BAI_KIEM_TRA' trên Desktop. (1,0 điểm)"
                ]
            }
        ],
        "grading": [
            "Câu 1 (2,0 điểm) – Thực hành trên máy tính:",
            "a) Mở File Explorer, tạo đúng thư mục BAI_KIEM_TRA (0,5 điểm).",
            "b) Tạo đúng 2 thư mục con LY_THUYET và THUC_HANH (0,5 điểm).",
            "c) Tạo đúng tệp Word TraLoi.docx trong LY_THUYET (0,5 điểm).",
            "d) Sao chép thành công tệp sang thư mục THUC_HANH (0,5 điểm).",
            "Câu 2 (3,0 điểm) – Trả lời viết ra trên giấy và thực hành soạn thảo:",
            "a) Nêu đúng 3 phần mềm hệ thống (Windows, iOS, Driver, Linux...) và 3 phần mềm ứng dụng (Word, Excel, Chrome, Paint...) (1,0 điểm).",
            "b) Mở Word, soạn thảo đoạn văn đúng chính tả, đủ 3-5 câu (1,0 điểm).",
            "c) Lưu đúng tên GioiThieu_UNIGO.docx vào thư mục BAI_KIEM_TRA (1,0 điểm)."
        ]
    },
    8: {
        "title": "PHẦN II. THỰC HÀNH VÀ BÀI TẬP (2 câu – 5,0 điểm)",
        "tasks": [
            {
                "header": "Câu 1 (2,0 điểm) – Thực hành trên máy tính: Tạo lập và định dạng bảng tính Excel:",
                "items": [
                    "Em hãy thực hiện trên máy tính các yêu cầu sau:",
                    "a) Khởi động phần mềm bảng tính Microsoft Excel. (0,5 điểm)",
                    "b) Tạo một bảng tính với tiêu đề \"BẢNG ĐIỂM HỌC TẬP LỚP 8A1\" gồm 6 cột: STT, Họ và tên, Toán, Ngữ văn, Tiếng Anh, Tin học. Nhập dữ liệu hoàn chỉnh cho 5 bạn học sinh trong lớp. (1,0 điểm)",
                    "c) Lưu bảng tính vào Desktop với tên tệp: 'BangDiem_8A1.xlsx'. (0,5 điểm)"
                ]
            },
            {
                "header": "Câu 2 (3,0 điểm) – Tìm kiếm thông tin trên máy tính và trả lời viết ra trên giấy:",
                "items": [
                    "a) [Thực hành trên máy tính] Mở trình duyệt web Google Chrome, tìm kiếm với cụm từ: \"Trí tuệ nhân tạo AI trong giáo dục\". (0,5 điểm)",
                    "b) [Trả lời viết ra trên giấy] Chọn 2 kết quả từ 2 nguồn trang web khác nhau, ghi lại tiêu đề bài viết và địa chỉ liên kết (URL): (1,0 điểm)",
                    "- Nguồn 1: ....................................................................................................................................",
                    "  Địa chỉ web 1: ................................................................................................................................",
                    "- Nguồn 2: ....................................................................................................................................",
                    "  Địa chỉ web 2: ................................................................................................................................",
                    "c) [Trả lời viết ra trên giấy] So sánh nội dung hai nguồn (điểm giống và khác nhau), đánh giá xem nguồn nào có độ tin cậy cao hơn và giải thích rõ lí do vì sao: (1,5 điểm)",
                    "- So sánh nội dung: ..........................................................................................................................",
                    "....................................................................................................................................................",
                    "- Đánh giá nguồn tin cậy hơn và lí do: .................................................................................................",
                    "...................................................................................................................................................."
                ]
            }
        ],
        "grading": [
            "Câu 1 (2,0 điểm) – Thực hành trên máy tính:",
            "a) Mở phần mềm Microsoft Excel thành công (0,5 điểm).",
            "b) Nhập đúng tiêu đề, đủ 6 cột và 5 dòng dữ liệu học sinh (1,0 điểm).",
            "c) Lưu đúng tên tệp BangDiem_8A1.xlsx vào Desktop (0,5 điểm).",
            "Câu 2 (3,0 điểm) – Tìm kiếm thông tin và trả lời viết ra trên giấy:",
            "a) Mở Google Chrome, tìm kiếm đúng cụm từ khóa (0,5 điểm).",
            "b) Ghi đúng tiêu đề và địa chỉ web của 2 nguồn thông tin khác nhau (1,0 điểm).",
            "c) Nêu được so sánh hợp lí, đánh giá chính xác nguồn uy tín hơn (ví dụ bài báo khoa học, trang .edu/.gov uy tín hơn mạng xã hội tự do) (1,5 điểm)."
        ]
    }
}

# ============================================================
# HELPER TẠO OXML ELEMENTS AN TOÀN
# ============================================================

def make_p(text="", bold=False, italic=False, align=None, space_before=0, space_after=0, indent=0):
    """Tạo paragraph element an toàn bằng OxmlElement, không bao giờ bị lỗi XML parse."""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    if align:
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), align)
        pPr.append(jc)
    if space_before or space_after:
        sp = OxmlElement('w:spacing')
        if space_before:
            sp.set(qn('w:before'), str(int(space_before * 20)))
        if space_after:
            sp.set(qn('w:after'), str(int(space_after * 20)))
        pPr.append(sp)
    if indent:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), str(int(indent * 20)))
        pPr.append(ind)
    p.append(pPr)
    
    if text:
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:cs'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        rPr.append(rFonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '26')
        rPr.append(sz)
        szCs = OxmlElement('w:szCs')
        szCs.set(qn('w:val'), '26')
        rPr.append(szCs)
        if bold:
            rPr.append(OxmlElement('w:b'))
            rPr.append(OxmlElement('w:bCs'))
        if italic:
            rPr.append(OxmlElement('w:i'))
            rPr.append(OxmlElement('w:iCs'))
        r.append(rPr)
        
        t = OxmlElement('w:t')
        t.text = text
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p.append(r)
    return p

def make_question_p(q_num, q_text):
    """Tạo paragraph câu hỏi: số câu in đậm, nội dung in thường."""
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '120')
    sp.set(qn('w:after'), '40')
    pPr.append(sp)
    p.append(pPr)
    
    # Run 1: "Câu X. "
    r1 = OxmlElement('w:r')
    r1Pr = OxmlElement('w:rPr')
    rFonts1 = OxmlElement('w:rFonts')
    rFonts1.set(qn('w:ascii'), 'Times New Roman')
    rFonts1.set(qn('w:hAnsi'), 'Times New Roman')
    r1Pr.append(rFonts1)
    sz1 = OxmlElement('w:sz')
    sz1.set(qn('w:val'), '26')
    r1Pr.append(sz1)
    r1Pr.append(OxmlElement('w:b'))
    r1.append(r1Pr)
    t1 = OxmlElement('w:t')
    t1.text = f"Câu {q_num}. "
    t1.set(qn('xml:space'), 'preserve')
    r1.append(t1)
    p.append(r1)
    
    # Run 2: q_text
    r2 = OxmlElement('w:r')
    r2Pr = OxmlElement('w:rPr')
    rFonts2 = OxmlElement('w:rFonts')
    rFonts2.set(qn('w:ascii'), 'Times New Roman')
    rFonts2.set(qn('w:hAnsi'), 'Times New Roman')
    r2Pr.append(rFonts2)
    sz2 = OxmlElement('w:sz')
    sz2.set(qn('w:val'), '26')
    r2Pr.append(sz2)
    r2.append(r2Pr)
    t2 = OxmlElement('w:t')
    t2.text = q_text
    t2.set(qn('xml:space'), 'preserve')
    r2.append(t2)
    p.append(r2)
    return p

def make_options_paragraphs(opt_dict):
    """Tạo 4 dòng lựa chọn A, B, C, D rõ ràng, thụt lề chuẩn."""
    paragraphs = []
    for letter in ['A', 'B', 'C', 'D']:
        p = make_p(f"{letter}. {opt_dict[letter]}", indent=18, space_before=1, space_after=2)
        paragraphs.append(p)
    return paragraphs

def set_font_run(run, name="Times New Roman", size=13, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), name)

def set_cell_text(cell, text, size=13, bold=False, italic=False, alignment=None):
    for p in cell.paragraphs:
        p.clear()
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_font_run(run, size=size, bold=bold, italic=italic)
    if alignment is not None:
        p.alignment = alignment

def generate_balanced_answers(seed):
    """Sinh mảng 20 đáp án đúng gồm đúng 5A, 5B, 5C, 5D, không quá 2 câu liên tiếp trùng nhau."""
    rng = random.Random(seed)
    items = ['A']*5 + ['B']*5 + ['C']*5 + ['D']*5
    for _ in range(2000):
        rng.shuffle(items)
        max_rep = 1
        curr_rep = 1
        for i in range(1, len(items)):
            if items[i] == items[i-1]:
                curr_rep += 1
                if curr_rep > max_rep:
                    max_rep = curr_rep
            else:
                curr_rep = 1
        if max_rep <= 2:
            return items
    return items

def extract_questions_and_options(doc):
    """Trích xuất 20 câu hỏi trắc nghiệm và các lựa chọn."""
    p1 = -1
    p2 = -1
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if re.search(r'PHẦN\s+I[\.\:\s]', t):
            p1 = i
        elif re.search(r'PHẦN\s+II[\.\:\s]', t):
            p2 = i
            break
            
    questions = []
    curr = None
    for i in range(p1+1, p2):
        t = doc.paragraphs[i].text.strip()
        if not t:
            continue
        m = re.match(r'^Câu\s+(\d+)\.\s*(.*)', t)
        if m:
            if curr:
                questions.append(curr)
            curr = {'num': int(m.group(1)), 'text': m.group(2).strip(), 'lines': []}
        elif curr is not None:
            curr['lines'].append(t)
    if curr:
        questions.append(curr)
        
    for q in questions:
        raw_text = ' \t '.join(q['lines'])
        pattern = r'(?:^|\s|\t)([A-D])\.\s*(.*?)(?=(?:[\s\t]+[A-D]\.|$))'
        matches = re.findall(pattern, raw_text, re.DOTALL)
        opts = {}
        for m in matches:
            opts[m[0]] = m[1].strip()
        q['options'] = opts
        
    return questions

def extract_current_answers(doc):
    """Lấy đáp án hiện tại từ Table 2 và Table 3."""
    ans_key = {}
    for ti in [2, 3]:
        if ti < len(doc.tables):
            t = doc.tables[ti]
            for col_idx in range(len(t.columns)):
                q_label = t.cell(0, col_idx).text.strip()
                ans_val = t.cell(1, col_idx).text.strip().upper()
                m = re.search(r'\d+', q_label)
                if m:
                    ans_key[int(m.group(0))] = ans_val
    return ans_key

def shuffle_options_for_question(q, curr_ans, target_ans, seed):
    """Đảo 4 phương án sao cho đáp án đúng chuyển sang target_ans."""
    correct_text = q['options'][curr_ans]
    wrong_texts = [q['options'][k] for k in ['A', 'B', 'C', 'D'] if k != curr_ans]
    
    rng = random.Random(seed + q['num'] * 17)
    rng.shuffle(wrong_texts)
    
    new_options = {}
    new_options[target_ans] = correct_text
    
    other_letters = [k for k in ['A', 'B', 'C', 'D'] if k != target_ans]
    for letter, wrong_text in zip(other_letters, wrong_texts):
        new_options[letter] = wrong_text
        
    return new_options

# ============================================================
# HÀM XỬ LÝ CHÍNH CHO 1 FILE ĐỀ THI
# ============================================================

def process_exam_file(grade, file_path):
    print(f"\n{'='*60}")
    print(f"XỬ LÝ ĐỀ KIỂM TRA TIN HỌC LỚP {grade}: {os.path.basename(file_path)}")
    print(f"{'='*60}")
    
    doc = Document(file_path)
    
    # 1. Trích xuất câu hỏi và đáp án hiện tại
    questions = extract_questions_and_options(doc)
    old_ans_key = extract_current_answers(doc)
    print(f"  Trích xuất: {len(questions)} câu hỏi, {len(old_ans_key)} đáp án cũ.")
    
    # 2. Sinh mảng đáp án mới chia đều 5A, 5B, 5C, 5D
    new_target_answers = generate_balanced_answers(2026 + grade * 13)
    counts = {k: new_target_answers.count(k) for k in 'ABCD'}
    print(f"  Phân bố đáp án mới: {counts}")
    print(f"  Chuỗi đáp án mới: {' '.join(new_target_answers)}")
    
    # 3. Đảo options cho từng câu hỏi
    shuffled_questions = []
    new_ans_key = {}
    for i, q in enumerate(questions):
        q_num = q['num']
        curr_ans = old_ans_key.get(q_num, 'B')
        target_ans = new_target_answers[i]
        new_opts = shuffle_options_for_question(q, curr_ans, target_ans, 2026 + grade * 100)
        
        shuffled_questions.append({
            'num': q_num,
            'text': q['text'],
            'options': new_opts,
            'target_ans': target_ans
        })
        new_ans_key[q_num] = target_ans
        
    # 4. Cập nhật Bảng đáp án Table 2 (Câu 1-10) và Table 3 (Câu 11-20)
    t2 = doc.tables[2]
    for col_idx in range(min(10, len(t2.columns))):
        q_num = col_idx + 1
        set_cell_text(t2.cell(1, col_idx), new_ans_key[q_num], bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        
    t3 = doc.tables[3]
    for col_idx in range(min(10, len(t3.columns))):
        q_num = col_idx + 11
        set_cell_text(t3.cell(1, col_idx), new_ans_key[q_num], bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    print("  ✅ Đã cập nhật Table 2 & Table 3 với đáp án chia đều mới.")
    
    # 5. Xây dựng lại nội dung paragraphs:
    body = doc.element.body
    tables_in_body = [child for child in body if child.tag.endswith('tbl')]
    t0_el, t1_el, t2_el, t3_el, t4_el = tables_in_body[0], tables_in_body[1], tables_in_body[2], tables_in_body[3], tables_in_body[4]
    
    idx_t1 = list(body).index(t1_el)
    idx_t2 = list(body).index(t2_el)
    
    # Xóa các paragraphs nằm giữa Table 1 và Table 2
    children_between_t1_t2 = list(body)[idx_t1 + 1 : idx_t2]
    for child in children_between_t1_t2:
        body.remove(child)
        
    # Xóa các paragraphs nằm giữa Table 3 và Table 4
    idx_t3_new = list(body).index(t3_el)
    idx_t4_new = list(body).index(t4_el)
    children_between_t3_t4 = list(body)[idx_t3_new + 1 : idx_t4_new]
    for child in children_between_t3_t4:
        body.remove(child)
        
    # Danh sách các XML elements cần chèn vào trước Table 2:
    new_elements_before_t2 = []
    
    # Dòng trống phân cách
    new_elements_before_t2.append(make_p("", space_after=4))
    
    # Tiêu đề Phần I
    new_elements_before_t2.append(make_p("PHẦN I. TRẮC NGHIỆM KHÁCH QUAN (20 câu – 5,0 điểm)", bold=True, space_before=6, space_after=2))
    new_elements_before_t2.append(make_p("Khoanh tròn vào chữ cái (A, B, C hoặc D) đứng trước câu trả lời đúng nhất:", italic=True, space_after=6))
    
    # 20 câu hỏi trắc nghiệm
    for q in shuffled_questions:
        qp = make_question_p(q['num'], q['text'])
        new_elements_before_t2.append(qp)
        opts_p = make_options_paragraphs(q['options'])
        new_elements_before_t2.extend(opts_p)
        
    # Tiêu đề Phần II
    p2_info = PART2_CONTENT[grade]
    new_elements_before_t2.append(make_p("", space_after=6))
    new_elements_before_t2.append(make_p(p2_info["title"], bold=True, space_before=8, space_after=4))
    
    # Các câu hỏi Phần II
    for task in p2_info["tasks"]:
        new_elements_before_t2.append(make_p(task["header"], bold=True, space_before=5, space_after=2))
        for item in task["items"]:
            new_elements_before_t2.append(make_p(item, space_before=1, space_after=2, indent=18))
            
    # Kết thúc đề
    new_elements_before_t2.append(make_p("", space_after=4))
    new_elements_before_t2.append(make_p("─── Hết ───", bold=True, align="center", space_before=6, space_after=2))
    new_elements_before_t2.append(make_p("Giám thị không giải thích gì thêm.", italic=True, align="center", space_after=12))
    
    # Tiêu đề Đáp án
    new_elements_before_t2.append(make_p("", space_after=8))
    new_elements_before_t2.append(make_p("ĐÁP ÁN VÀ HƯỚNG DẪN CHẤM", bold=True, align="center", space_before=8, space_after=4))
    new_elements_before_t2.append(make_p("I. TRẮC NGHIỆM (20 câu – 5,0 điểm. Mỗi câu trả lời đúng được 0,25 điểm)", bold=True, space_before=4, space_after=4))
    
    # Chèn các element này vào body ngay trước Table 2
    for el in new_elements_before_t2:
        t2_el.addprevious(el)
        
    # Bây giờ chèn nội dung vào giữa Table 3 và Table 4 (Hướng dẫn chấm Phần II):
    new_elements_before_t4 = []
    new_elements_before_t4.append(make_p("", space_after=4))
    new_elements_before_t4.append(make_p("II. THỰC HÀNH VÀ BÀI TẬP (2 câu – 5,0 điểm)", bold=True, space_before=6, space_after=4))
    for line in p2_info["grading"]:
        is_h = line.startswith("Câu ")
        new_elements_before_t4.append(make_p(line, bold=is_h, space_before=2, space_after=2, indent=0 if is_h else 18))
    new_elements_before_t4.append(make_p("", space_after=8))
    
    for el in new_elements_before_t4:
        t4_el.addprevious(el)
        
    # 6. Lưu file
    doc.save(file_path)
    print(f"  💾 Đã lưu thành công: {file_path}")
    return new_ans_key

# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 70)
    print("CHUẨN HÓA VÀ ĐẢO ĐÁP ÁN ĐỀ KIỂM TRA TIN HỌC (LỚP 3 - LỚP 8)")
    print("=" * 70)
    
    de_files = [
        (3, r'd:\UNIGO\KHBD_Tin_học\Lớp_3\Tuần_10\De_kiem_tra_DGDK1_Tin_hoc_Lop_3.docx'),
        (4, r'd:\UNIGO\KHBD_Tin_học\Lớp_4\Tuần_10\De_kiem_tra_DGDK1_Tin_hoc_Lop_4.docx'),
        (5, r'd:\UNIGO\KHBD_Tin_học\Lớp_5\Tuần_10\De_kiem_tra_DGDK1_Tin_hoc_Lop_5.docx'),
        (6, r'd:\UNIGO\KHBD_Tin_học\Lớp_6\Tuần_10\De_kiem_tra_DGDK1_Tin_hoc_Lop_6.docx'),
        (7, r'd:\UNIGO\KHBD_Tin_học\Lớp_7\Tuần_10\De_kiem_tra_DGDK1_Tin_hoc_Lop_7.docx'),
        (8, r'd:\UNIGO\KHBD_Tin_học\Lớp_8\Tuần_10\De_kiem_tra_DGDK1_Tin_hoc_Lop_8.docx')
    ]
    
    all_results = {}
    for grade, file_path in de_files:
        if not os.path.exists(file_path):
            print(f"❌ Không tìm thấy file: {file_path}")
            continue
        new_ans = process_exam_file(grade, file_path)
        all_results[grade] = new_ans
        
    print("\n" + "=" * 70)
    print("TỔNG HỢP KIỂM TRA ĐÁP ÁN SAU KHI XÁO CHO 6 LỚP")
    print("=" * 70)
    for grade, ans_dict in all_results.items():
        ans_list = [ans_dict[i] for i in range(1, 21)]
        cA = ans_list.count('A')
        cB = ans_list.count('B')
        cC = ans_list.count('C')
        cD = ans_list.count('D')
        ans_str = ' '.join(ans_list)
        print(f"Lớp {grade}: A={cA}, B={cB}, C={cC}, D={cD} -> {ans_str}")

if __name__ == '__main__':
    main()
