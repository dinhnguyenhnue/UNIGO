# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

# Read a KHBD target file
print('=== KHBD TARGET: Lớp 1 Tuần 03 ===')
tgt = r'D:\UNIGO\KHBD_Robotics\Lớp_1\Tuần_03\KHBD_Robotics_Lớp_1_Tiet02_Bai_2_Chu_cun_de_thuong.docx'
doc = Document(tgt)
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
for i, p in enumerate(doc.paragraphs[:60]):
    txt = p.text.strip()
    if txt:
        bold = any(r.bold for r in p.runs if r.bold)
        print(f'  P{i}: [{"B" if bold else " "}] {txt[:200]}')

for i, t in enumerate(doc.tables):
    print(f'\n  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows[:5]):
        cells = [c.text[:60].replace('\n', '|') for c in row.cells]
        print(f'    Row {j}: {cells}')

# Also check KHBD Lớp 5
print('\n\n=== KHBD TARGET: Lớp 5 Tuần 03 ===')
tgt5 = r'D:\UNIGO\KHBD_Robotics\Lớp_5\Tuần_03\KHBD_Robotics_Lớp_5_Tiet01_Bai_1_ong_co_la_gi.docx'
doc5 = Document(tgt5)
print(f'Paragraphs: {len(doc5.paragraphs)}, Tables: {len(doc5.tables)}')
for i, p in enumerate(doc5.paragraphs[:60]):
    txt = p.text.strip()
    if txt:
        bold = any(r.bold for r in p.runs if r.bold)
        print(f'  P{i}: [{"B" if bold else " "}] {txt[:200]}')

for i, t in enumerate(doc5.tables):
    print(f'\n  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows[:5]):
        cells = [c.text[:60].replace('\n', '|') for c in row.cells]
        print(f'    Row {j}: {cells}')

# Also check KHBD Lớp 6
print('\n\n=== KHBD TARGET: Lớp 6 Tuần 03 ===')
tgt6 = r'D:\UNIGO\KHBD_Robotics\Lớp_6\Tuần_03\KHBD_Robotics_Lớp_6_Tiet01_Bai_1_ong_co_la_gi.docx'
doc6 = Document(tgt6)
print(f'Paragraphs: {len(doc6.paragraphs)}, Tables: {len(doc6.tables)}')
for i, p in enumerate(doc6.paragraphs[:60]):
    txt = p.text.strip()
    if txt:
        bold = any(r.bold for r in p.runs if r.bold)
        print(f'  P{i}: [{"B" if bold else " "}] {txt[:200]}')

for i, t in enumerate(doc6.tables):
    print(f'\n  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows[:5]):
        cells = [c.text[:60].replace('\n', '|') for c in row.cells]
        print(f'    Row {j}: {cells}')
