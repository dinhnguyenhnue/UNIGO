# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

base = r'D:\UNIGO\Phân phối chương trình\Robotics'

print('=== INITIATE FILES - Full titles ===')
init_dir = os.path.join(base, 'Giáo án-Initiate')
for f in sorted(os.listdir(init_dir)):
    fpath = os.path.join(init_dir, f)
    doc = Document(fpath)
    title = ''
    for p in doc.paragraphs[:15]:
        txt = p.text.strip()
        if 'BÀI' in txt and ':' in txt:
            title = txt
            break
    print(f'  {f} -> {title}')

print('\n=== KINDER FILES - Full titles ===')
kind_dir = os.path.join(base, 'Giáo án-Kinder')
for f in sorted(os.listdir(kind_dir)):
    fpath = os.path.join(kind_dir, f)
    doc = Document(fpath)
    title = ''
    for p in doc.paragraphs[:15]:
        txt = p.text.strip()
        if 'BÀI' in txt and ':' in txt:
            title = txt
            break
    print(f'  {f} -> {title}')

print('\n=== EXCEL FILES - Full titles ===')
exc_dir = os.path.join(base, 'Giáo án-Excel')
for f in sorted(os.listdir(exc_dir)):
    fpath = os.path.join(exc_dir, f)
    doc = Document(fpath)
    title = ''
    for p in doc.paragraphs[:15]:
        txt = p.text.strip()
        if 'BÀI' in txt and ':' in txt:
            title = txt
            break
    print(f'  {f} -> {title}')

# Now check Kế hoạch dạy học to understand the mapping
print('\n\n=== Checking KHDH files ===')
khdh_base = r'D:\UNIGO\Hệ thống mẫu văn bản'
for root, dirs, files in os.walk(khdh_base):
    for f in files:
        if 'Robotics' in f and 'Kế hoạch dạy học' in f and f.endswith('.docx'):
            print(f'  Found: {os.path.join(root, f)}')
