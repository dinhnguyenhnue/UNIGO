# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

# Read Initiate source files - check ALL file names and their lesson titles
base = r'D:\UNIGO\Phân phối chương trình\Robotics'

print('=== INITIATE FILES - Lesson titles ===')
init_dir = os.path.join(base, 'Giáo án-Initiate')
for f in sorted(os.listdir(init_dir)):
    fpath = os.path.join(init_dir, f)
    doc = Document(fpath)
    # Find the BÀI line
    title = ''
    for p in doc.paragraphs[:10]:
        txt = p.text.strip()
        if 'BÀI' in txt.upper() or 'TÊN BÀI' in txt.upper():
            title = txt
            break
    print(f'  {f} -> {title}')

print('\n=== KINDER FILES - Lesson titles ===')
kind_dir = os.path.join(base, 'Giáo án-Kinder')
for f in sorted(os.listdir(kind_dir)):
    fpath = os.path.join(kind_dir, f)
    doc = Document(fpath)
    title = ''
    for p in doc.paragraphs[:10]:
        txt = p.text.strip()
        if 'BÀI' in txt.upper() or 'TÊN BÀI' in txt.upper():
            title = txt
            break
    print(f'  {f} -> {title}')

print('\n=== EXCEL FILES - Lesson titles ===')
exc_dir = os.path.join(base, 'Giáo án-Excel')
for f in sorted(os.listdir(exc_dir)):
    fpath = os.path.join(exc_dir, f)
    doc = Document(fpath)
    title = ''
    for p in doc.paragraphs[:10]:
        txt = p.text.strip()
        if 'BÀI' in txt.upper() or 'TÊN BÀI' in txt.upper():
            title = txt
            break
    print(f'  {f} -> {title}')
