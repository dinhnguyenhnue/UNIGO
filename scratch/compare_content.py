# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

# Check KHBD Lớp 5 Tuần 03 Tiết 01 (Bài 1 - Động cơ là gì) vs Excel source
print('=== KHBD Lớp 5 Tuần 03 - Bài 1 (Current content) ===')
doc = Document(r'D:\UNIGO\KHBD_Robotics\Lớp_5\Tuần_03\KHBD_Robotics_Lớp_5_Tiet01_Bai_1_ong_co_la_gi.docx')
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        print(f'  P{i}: {txt[:180]}')

print('\n\n=== Excel Source robohub-exce.01-01 (What should be there) ===')
doc2 = Document(r'D:\UNIGO\Phân phối chương trình\Robotics\Giáo án-Excel\robohub-exce.01-01.docx')
for i, p in enumerate(doc2.paragraphs):
    txt = p.text.strip()
    if txt:
        print(f'  P{i}: {txt[:180]}')

# Now check Lớp 1 mapping - Initiate I 03-01 = Bài 1
print('\n\n=== KHBD Lớp 1 Tuần 02 - Bài 1 (Current content) ===')
doc3 = Document(r'D:\UNIGO\KHBD_Robotics\Lớp_1\Tuần_02\KHBD_Robotics_Lớp_1_Tiet01_Bai_1_Tap_the_duc_nao.docx')
for i, p in enumerate(doc3.paragraphs[:40]):
    txt = p.text.strip()
    if txt:
        print(f'  P{i}: {txt[:180]}')

print('\n\n=== Initiate Source I 03-01 đu xà (What should be there) ===')
doc4 = Document(r'D:\UNIGO\Phân phối chương trình\Robotics\Giáo án-Initiate\I 03 - 01 - đu xà.docx')
for i, p in enumerate(doc4.paragraphs[:40]):
    txt = p.text.strip()
    if txt:
        print(f'  P{i}: {txt[:180]}')
