# -*- coding: utf-8 -*-
import sys, os, glob
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

base = r'D:\UNIGO\Phân phối chương trình\Robotics'

# List all source files in each folder
for folder in ['Giáo án-Initiate', 'Giáo án-Kinder', 'Giáo án-Excel']:
    fpath = os.path.join(base, folder)
    print(f'\n=== {folder} ===')
    files = sorted(os.listdir(fpath))
    for f in files:
        print(f'  {f}')

# Read first Initiate file to understand structure
print('\n\n=== READING FIRST INITIATE FILE ===')
init_dir = os.path.join(base, 'Giáo án-Initiate')
first_file = sorted(os.listdir(init_dir))[0]
fpath = os.path.join(init_dir, first_file)
print(f'File: {first_file}')
doc = Document(fpath)
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
for i, p in enumerate(doc.paragraphs[:50]):
    txt = p.text.strip()
    if txt:
        print(f'  P{i}: [{p.style.name}] {txt[:200]}')
for i, t in enumerate(doc.tables):
    print(f'\n  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows[:8]):
        cells = [c.text[:60].replace('\n', '|') for c in row.cells]
        print(f'    Row {j}: {cells}')

# Read first Kinder file
print('\n\n=== READING FIRST KINDER FILE ===')
kind_dir = os.path.join(base, 'Giáo án-Kinder')
first_file = sorted(os.listdir(kind_dir))[0]
fpath = os.path.join(kind_dir, first_file)
print(f'File: {first_file}')
doc = Document(fpath)
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
for i, p in enumerate(doc.paragraphs[:50]):
    txt = p.text.strip()
    if txt:
        print(f'  P{i}: [{p.style.name}] {txt[:200]}')
for i, t in enumerate(doc.tables):
    print(f'\n  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows[:8]):
        cells = [c.text[:60].replace('\n', '|') for c in row.cells]
        print(f'    Row {j}: {cells}')

# Read first Excel file
print('\n\n=== READING FIRST EXCEL FILE ===')
exc_dir = os.path.join(base, 'Giáo án-Excel')
first_file = sorted(os.listdir(exc_dir))[0]
fpath = os.path.join(exc_dir, first_file)
print(f'File: {first_file}')
doc = Document(fpath)
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
for i, p in enumerate(doc.paragraphs[:50]):
    txt = p.text.strip()
    if txt:
        print(f'  P{i}: [{p.style.name}] {txt[:200]}')
for i, t in enumerate(doc.tables):
    print(f'\n  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols')
    for j, row in enumerate(t.rows[:8]):
        cells = [c.text[:60].replace('\n', '|') for c in row.cells]
        print(f'    Row {j}: {cells}')
